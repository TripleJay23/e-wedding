"""Opt-in disposable PostgreSQL M4 end-to-end import acceptance."""

from __future__ import annotations

import asyncio
import os
from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4

import httpx
import pytest
from openpyxl import Workbook
from sqlalchemy import create_engine, delete, func, inspect, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.cli import disposable_alembic
from app.core.security import AuthenticationSecurity
from app.core.settings import Settings
from app.db.runtime import DatabaseRuntime
from app.main import create_app
from app.models.administrator import Administrator
from app.models.attendance import GateScanAttempt, GuestAttendanceRecord
from app.models.attendance_corrections import GuestAttendanceCorrection
from app.models.cards import GeneratedGuestCard
from app.models.event import Event
from app.models.guest import Guest
from app.models.imports import (
    ImportBatch,
    ImportBatchRow,
    ImportBatchStatus,
    ImportRevisionKind,
    ImportRowRevision,
    ImportRowStatus,
)
from app.models.qr import GuestQrCredential
from app.repositories.administrators import AdministratorRepository
from app.services.admin_provisioning import AdministratorProvisioningService

DISPOSABLE_URL = os.getenv("WEDDING_DISPOSABLE_DATABASE_URL")
pytestmark = [
    pytest.mark.postgresql,
    pytest.mark.skipif(
        DISPOSABLE_URL is None,
        reason="WEDDING_DISPOSABLE_DATABASE_URL is not explicitly configured",
    ),
]


def synthetic_workbook() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Synthetic Guest List"
    worksheet.merge_cells("A1:E1")
    worksheet["A1"] = "Synthetic PostgreSQL Wedding Guests"
    worksheet.append([])
    worksheet.append(
        [
            "Unmanaged Notes",
            "Full Name",
            "Allowed Companions",
            "Invitation Type",
            "Email",
            "Phone No.",
            "Table Number",
        ]
    )
    worksheet.append(
        [None, "Synthetic Valid One", "2.00", "Family", "one@example.test", None, "02"]
    )
    worksheet.append([None, "Synthetic Correctable", 0, None, "correct@example.test", None, 2.0])
    worksheet.append([None, "Synthetic Valid Two", 0, "Individual", None, "+255700000022", "A1"])
    worksheet.append([None, "Synthetic Duplicate", 1, "Family", "ONE@example.test", None, 12])
    worksheet.append(["Grand Total", None, "=COUNTA(B4:B7)", None, None, None, None])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def contextual_footer_workbook() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Synthetic Contextual Footer"
    worksheet.merge_cells("A1:E1")
    worksheet["A1"] = "Synthetic Wedding Guest Register"
    worksheet.merge_cells("A2:E2")
    worksheet["A2"] = "Synthetic Import Preview"
    worksheet.append([])
    worksheet.append(
        [
            "Unmanaged Notes",
            "Full Name",
            "Invitation Type",
            "Allowed Companions",
            "Category",
        ]
    )
    worksheet.append([None, "Synthetic Context One", "Family", 2, "Friends"])
    worksheet.append([None, "Synthetic Context Two", "Individual", 0, "Colleagues"])
    worksheet.append([None, "Synthetic Context Three", "Family", 1, None])
    worksheet.append(["Closing constellation reckoning", None, None, "=COUNTA(B5:B7)", None])
    worksheet.append(["Secondary constellation reckoning", None, None, "=COUNTA(B5:B7)", None])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def overlength_email_workbook() -> bytes:
    email_limit = getattr(Guest.__table__.columns.email.type, "length", None)
    assert isinstance(email_limit, int)
    suffix = "@example.test"
    overlength_email = ("a" * (email_limit + 1 - len(suffix))) + suffix

    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Synthetic Persistence Boundary"
    worksheet.append(["Full Name", "Invitation Type", "Email"])
    worksheet.append(["Guest Alpha", "Family", overlength_email])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def partial_retry_workbook(*, corrected: bool) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Synthetic Partial Retry"
    worksheet.append(["Full Name", "Invitation Type", "Email", "Allowed Companions"])
    worksheet.append(["Synthetic Persisted", "Family", "persisted@example.test", 1])
    worksheet.append(
        [
            "Synthetic Preview Only",
            "Individual" if corrected else None,
            "preview-only@example.test",
            0,
        ]
    )
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def concurrent_retry_workbook() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Synthetic Concurrent Retry"
    worksheet.append(["Full Name", "Invitation Type", "Email", "Allowed Companions"])
    worksheet.append(["Synthetic Concurrent One", "Family", "concurrent-one@example.test", 2])
    worksheet.append(["Synthetic Concurrent Two", "Individual", "concurrent-two@example.test", 0])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def selectable_workbook(
    row_count: int,
    prefix: str,
    *,
    invalid_last_row: bool = False,
) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = f"Synthetic {prefix}"
    worksheet.append(
        [
            "Full Name",
            "Phone",
            "Email",
            "Invitation Type",
            "Allowed Companions",
            "Table Number",
            "Category",
        ]
    )
    for index in range(1, row_count + 1):
        worksheet.append(
            [
                f"Synthetic {prefix} Guest {index:03d}",
                f"+2557{index:08d}",
                f"{prefix.casefold()}-{index:03d}@example.test",
                None if invalid_last_row and index == row_count else "Family",
                index % 3,
                f"T-{index:03d}",
                "Friends",
            ]
        )
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def configure_test_environment(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> Settings:
    assert DISPOSABLE_URL is not None
    private_root = tmp_path / "private"
    monkeypatch.setenv("APP_ENV", "test")
    monkeypatch.setenv("DATABASE_URL", DISPOSABLE_URL)
    monkeypatch.setenv("JWT_SECRET", "synthetic-jwt-signing-secret-at-least-32-bytes")
    monkeypatch.setenv("JWT_ISSUER", "synthetic-issuer")
    monkeypatch.setenv("JWT_AUDIENCE", "synthetic-audience")
    monkeypatch.setenv("PRIVATE_STORAGE_ROOT", str(private_root))
    monkeypatch.setenv("IMPORT_EVIDENCE_STORAGE_ROOT", str(private_root / "import-evidence"))
    return Settings()


async def exercise_overlength_persistence_boundary(
    settings: Settings,
    runtime: DatabaseRuntime,
    username: str,
) -> None:
    app = create_app(
        settings,
        database_runtime=runtime,
        authentication_security=AuthenticationSecurity.from_settings(settings),
    )
    transport = httpx.ASGITransport(app=app)
    async with (
        app.router.lifespan_context(app),
        httpx.AsyncClient(transport=transport, base_url="http://synthetic.test") as client,
    ):
        login = await client.post(
            "/api/v1/auth/login",
            json={"username": username, "password": "synthetic administrator password"},
        )
        assert login.status_code == 200
        headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
        created_event = await client.post(
            "/api/v1/events",
            headers=headers,
            json={
                "event_name": "Synthetic Persistence Boundary",
                "bride_name": "Guest Alpha",
                "groom_name": "Guest Beta",
                "event_datetime": "2027-06-12T14:30:00+03:00",
                "venue": "Synthetic Hall",
                "max_companions_per_invitation": 2,
            },
        )
        assert created_event.status_code == 201
        event_id = created_event.json()["id"]
        upload = await client.post(
            f"/api/v1/events/{event_id}/imports",
            headers=headers,
            files={
                "workbook": (
                    "synthetic-overlength.xlsx",
                    overlength_email_workbook(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert upload.status_code == 201
        batch_id = upload.json()["id"]
        rows = await client.get(
            f"/api/v1/events/{event_id}/imports/{batch_id}/rows",
            headers=headers,
        )
        assert rows.status_code == 200
        assert len(rows.json()["items"]) == 1

        progression = await client.post(
            f"/api/v1/events/{event_id}/imports/{batch_id}/progress",
            headers=headers,
        )
        assert progression.status_code == 200
        assert progression.json()["imported_rows"] == 0
        assert rows.json()["items"][0]["status"] != ImportRowStatus.VALID.value
        assert any(
            item["field"] == "email" for item in rows.json()["items"][0]["validation_errors"]
        )
        guests = await client.get(f"/api/v1/events/{event_id}/guests", headers=headers)
        assert guests.status_code == 200
        assert guests.json()["items"] == []


async def exercise_import_api(
    settings: Settings,
    runtime: DatabaseRuntime,
    username: str,
) -> dict[str, str]:
    app = create_app(
        settings,
        database_runtime=runtime,
        authentication_security=AuthenticationSecurity.from_settings(settings),
    )
    transport = httpx.ASGITransport(app=app)
    async with (
        app.router.lifespan_context(app),
        httpx.AsyncClient(transport=transport, base_url="http://synthetic.test") as client,
    ):
        login = await client.post(
            "/api/v1/auth/login",
            json={"username": username, "password": "synthetic administrator password"},
        )
        assert login.status_code == 200
        headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
        event_payload = {
            "event_name": "Synthetic M4 PostgreSQL Wedding",
            "bride_name": "Amina",
            "groom_name": "Baraka",
            "event_datetime": "2027-06-12T14:30:00+03:00",
            "venue": "Synthetic Hall",
            "max_companions_per_invitation": 2,
        }
        created_event = await client.post("/api/v1/events", headers=headers, json=event_payload)
        assert created_event.status_code == 201, created_event.text
        event_id = created_event.json()["id"]

        upload = await client.post(
            f"/api/v1/events/{event_id}/imports",
            headers=headers,
            files={
                "workbook": (
                    "synthetic.xlsx",
                    synthetic_workbook(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert upload.status_code == 201, upload.text
        upload_body = upload.json()
        batch_id = upload_body["id"]
        reconciliation = upload_body["reconciliation"]
        assert upload_body["detected_sheet_name"] == "Synthetic Guest List"
        assert upload_body["header_row_number"] == 3
        assert reconciliation == {
            "actual_guest_rows_discovered": 4,
            "ignored_non_data_rows": 3,
            "currently_valid_staged_rows": 2,
            "currently_invalid_staged_rows": 1,
            "duplicates": 1,
            "pending_revalidation": 0,
            "rows_already_imported": 0,
            "newly_importable_rows": 2,
            "remaining_unresolved_rows": 2,
            "guests_created_from_batch": 0,
            "overall_batch_state": "VALIDATED",
        }
        assert "storage_key" not in upload_body
        assert "private" not in upload.text.casefold()

        staged_rows = await client.get(
            f"/api/v1/events/{event_id}/imports/{batch_id}/rows?limit=100&offset=0",
            headers=headers,
        )
        assert staged_rows.status_code == 200
        staged_items = staged_rows.json()["items"]
        assert [item["source_row_number"] for item in staged_items] == [4, 5, 6, 7]
        assert [item["current_values"]["table_number"] for item in staged_items] == [
            "02",
            "2",
            "A1",
            "12",
        ]
        assert [item["current_values"]["allowed_companions"] for item in staged_items] == [
            2,
            0,
            0,
            1,
        ]
        assert all(item["original_values"] == item["current_values"] for item in staged_items)

        first_preview_snapshot = [
            {
                "id": item["id"],
                "status": item["status"],
                "original_values": item["original_values"],
                "validation_errors": item["validation_errors"],
            }
            for item in staged_items
        ]
        retry_batch_ids: list[str] = []
        for retry_number in (2, 3):
            retry = await client.post(
                f"/api/v1/events/{event_id}/imports",
                headers=headers,
                files={
                    "workbook": (
                        f"synthetic-retry-{retry_number}.xlsx",
                        synthetic_workbook(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert retry.status_code == 201, retry.text
            retry_body = retry.json()
            retry_batch_ids.append(retry_body["id"])
            assert retry_body["reconciliation"] == reconciliation
            retry_rows = await client.get(
                f"/api/v1/events/{event_id}/imports/{retry_body['id']}/rows?limit=100&offset=0",
                headers=headers,
            )
            assert retry_rows.status_code == 200
            retry_items = retry_rows.json()["items"]
            duplicate_codes = [
                error_item["code"]
                for item in retry_items
                for error_item in item["validation_errors"]
                if error_item["code"].startswith("DUPLICATE")
            ]
            assert duplicate_codes == ["DUPLICATE_IN_WORKBOOK"]

        unchanged_first_preview = await client.get(
            f"/api/v1/events/{event_id}/imports/{batch_id}/rows?limit=100&offset=0",
            headers=headers,
        )
        assert unchanged_first_preview.status_code == 200
        assert [
            {
                "id": item["id"],
                "status": item["status"],
                "original_values": item["original_values"],
                "validation_errors": item["validation_errors"],
            }
            for item in unchanged_first_preview.json()["items"]
        ] == first_preview_snapshot

        guests_before_progress = await client.get(
            f"/api/v1/events/{event_id}/guests?limit=100&offset=0",
            headers=headers,
        )
        assert guests_before_progress.status_code == 200
        assert guests_before_progress.json()["items"] == []

        contextual_event = await client.post(
            "/api/v1/events",
            headers=headers,
            json={**event_payload, "event_name": "Synthetic Contextual Footer Wedding"},
        )
        assert contextual_event.status_code == 201
        contextual_event_id = contextual_event.json()["id"]
        contextual_upload = await client.post(
            f"/api/v1/events/{contextual_event_id}/imports",
            headers=headers,
            files={
                "workbook": (
                    "synthetic-contextual-footer.xlsx",
                    contextual_footer_workbook(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert contextual_upload.status_code == 201, contextual_upload.text
        contextual_body = contextual_upload.json()
        contextual_batch_id = contextual_body["id"]
        assert contextual_body["reconciliation"] == {
            "actual_guest_rows_discovered": 3,
            "ignored_non_data_rows": 5,
            "currently_valid_staged_rows": 3,
            "currently_invalid_staged_rows": 0,
            "duplicates": 0,
            "pending_revalidation": 0,
            "rows_already_imported": 0,
            "newly_importable_rows": 3,
            "remaining_unresolved_rows": 0,
            "guests_created_from_batch": 0,
            "overall_batch_state": "VALIDATED",
        }
        contextual_rows = await client.get(
            f"/api/v1/events/{contextual_event_id}/imports/{contextual_batch_id}/rows",
            headers=headers,
        )
        assert contextual_rows.status_code == 200
        assert [item["source_row_number"] for item in contextual_rows.json()["items"]] == [
            5,
            6,
            7,
        ]
        assert all(item["validation_errors"] == [] for item in contextual_rows.json()["items"])
        contextual_guests = await client.get(
            f"/api/v1/events/{contextual_event_id}/guests?limit=100&offset=0",
            headers=headers,
        )
        assert contextual_guests.status_code == 200
        assert contextual_guests.json()["items"] == []

        first_progress = await client.post(
            f"/api/v1/events/{event_id}/imports/{batch_id}/progress",
            headers=headers,
        )
        assert first_progress.status_code == 200, first_progress.text
        assert first_progress.json()["imported_rows"] == 2
        assert first_progress.json()["reconciliation"]["rows_already_imported"] == 2

        invalid_rows = await client.get(
            f"/api/v1/events/{event_id}/imports/{batch_id}/rows?status=INVALID",
            headers=headers,
        )
        assert invalid_rows.status_code == 200
        invalid_row = invalid_rows.json()["items"][0]
        original_values = invalid_row["original_values"]
        row_id = invalid_row["id"]
        corrected = await client.patch(
            f"/api/v1/events/{event_id}/imports/{batch_id}/rows/{row_id}",
            headers=headers,
            json={
                "expected_version": invalid_row["row_version"],
                "values": {"invitation_type": "Family"},
            },
        )
        assert corrected.status_code == 200, corrected.text
        assert corrected.json()["status"] == "PENDING_REVALIDATION"
        assert corrected.json()["original_values"] == original_values

        revalidated = await client.post(
            f"/api/v1/events/{event_id}/imports/{batch_id}/rows/{row_id}/revalidate",
            headers=headers,
        )
        assert revalidated.status_code == 200, revalidated.text
        assert revalidated.json()["status"] == "VALID"

        second_progress = await client.post(
            f"/api/v1/events/{event_id}/imports/{batch_id}/progress",
            headers=headers,
        )
        repeated_progress = await client.post(
            f"/api/v1/events/{event_id}/imports/{batch_id}/progress",
            headers=headers,
        )
        assert second_progress.json()["imported_rows"] == 1
        assert repeated_progress.json()["imported_rows"] == 0
        assert repeated_progress.json()["reconciliation"]["rows_already_imported"] == 3
        assert repeated_progress.json()["reconciliation"]["duplicates"] == 1

        guests = await client.get(
            f"/api/v1/events/{event_id}/guests?limit=100&offset=0",
            headers=headers,
        )
        assert guests.status_code == 200
        guest_items = guests.json()["items"]
        assert len(guest_items) == 3
        first_guest = next(
            item for item in guest_items if item["full_name"] == "Synthetic Valid One"
        )
        assert first_guest["allowed_companions"] == 2
        assert first_guest["table_number"] == "02"

        reupload = await client.post(
            f"/api/v1/events/{event_id}/imports",
            headers=headers,
            files={
                "workbook": (
                    "synthetic.xlsx",
                    synthetic_workbook(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert reupload.status_code == 201, reupload.text
        assert reupload.json()["reconciliation"]["duplicates"] == 3
        assert reupload.json()["reconciliation"]["currently_invalid_staged_rows"] == 1

        partial_event = await client.post(
            "/api/v1/events",
            headers=headers,
            json={**event_payload, "event_name": "Synthetic Partial Retry Wedding"},
        )
        assert partial_event.status_code == 201
        partial_event_id = partial_event.json()["id"]
        partial_preview = await client.post(
            f"/api/v1/events/{partial_event_id}/imports",
            headers=headers,
            files={
                "workbook": (
                    "synthetic-partial.xlsx",
                    partial_retry_workbook(corrected=False),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert partial_preview.status_code == 201, partial_preview.text
        partial_batch_id = partial_preview.json()["id"]
        assert partial_preview.json()["reconciliation"]["currently_valid_staged_rows"] == 1
        assert partial_preview.json()["reconciliation"]["currently_invalid_staged_rows"] == 1
        partial_progress = await client.post(
            f"/api/v1/events/{partial_event_id}/imports/{partial_batch_id}/progress",
            headers=headers,
        )
        assert partial_progress.status_code == 200, partial_progress.text
        assert partial_progress.json()["imported_rows"] == 1

        partial_retry = await client.post(
            f"/api/v1/events/{partial_event_id}/imports",
            headers=headers,
            files={
                "workbook": (
                    "synthetic-partial-retry.xlsx",
                    partial_retry_workbook(corrected=True),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        assert partial_retry.status_code == 201, partial_retry.text
        partial_retry_reconciliation = partial_retry.json()["reconciliation"]
        assert partial_retry_reconciliation["currently_valid_staged_rows"] == 1
        assert partial_retry_reconciliation["currently_invalid_staged_rows"] == 0
        assert partial_retry_reconciliation["duplicates"] == 1
        partial_retry_rows = await client.get(
            f"/api/v1/events/{partial_event_id}/imports/{partial_retry.json()['id']}/rows",
            headers=headers,
        )
        assert partial_retry_rows.status_code == 200
        partial_retry_items = partial_retry_rows.json()["items"]
        assert [item["status"] for item in partial_retry_items] == ["DUPLICATE", "VALID"]
        assert partial_retry_items[0]["validation_errors"][0]["code"] == (
            "DUPLICATE_EXISTING_GUEST"
        )
        assert partial_retry_items[1]["validation_errors"] == []
        partial_historical_rows = await client.get(
            f"/api/v1/events/{partial_event_id}/imports/{partial_batch_id}/rows",
            headers=headers,
        )
        assert partial_historical_rows.status_code == 200
        assert [item["status"] for item in partial_historical_rows.json()["items"]] == [
            "IMPORTED",
            "INVALID",
        ]
        partial_guests = await client.get(
            f"/api/v1/events/{partial_event_id}/guests?limit=100&offset=0",
            headers=headers,
        )
        assert partial_guests.status_code == 200
        assert len(partial_guests.json()["items"]) == 1

        concurrent_event = await client.post(
            "/api/v1/events",
            headers=headers,
            json={**event_payload, "event_name": "Synthetic Concurrent Retry Wedding"},
        )
        assert concurrent_event.status_code == 201
        concurrent_event_id = concurrent_event.json()["id"]
        concurrent_batch_ids: list[str] = []
        for upload_number in (1, 2, 3):
            concurrent_preview = await client.post(
                f"/api/v1/events/{concurrent_event_id}/imports",
                headers=headers,
                files={
                    "workbook": (
                        f"synthetic-concurrent-{upload_number}.xlsx",
                        concurrent_retry_workbook(),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert concurrent_preview.status_code == 201, concurrent_preview.text
            concurrent_reconciliation = concurrent_preview.json()["reconciliation"]
            assert concurrent_reconciliation["currently_valid_staged_rows"] == 2
            assert concurrent_reconciliation["duplicates"] == 0
            concurrent_batch_ids.append(concurrent_preview.json()["id"])

        concurrent_guests_before_progress = await client.get(
            f"/api/v1/events/{concurrent_event_id}/guests?limit=100&offset=0",
            headers=headers,
        )
        assert concurrent_guests_before_progress.status_code == 200
        assert concurrent_guests_before_progress.json()["items"] == []

        concurrent_results = await asyncio.gather(
            *(
                client.post(
                    f"/api/v1/events/{concurrent_event_id}/imports/{concurrent_batch_id}/progress",
                    headers=headers,
                )
                for concurrent_batch_id in concurrent_batch_ids[:2]
            )
        )
        assert all(result.status_code == 200 for result in concurrent_results)
        assert sum(result.json()["imported_rows"] for result in concurrent_results) == 2
        assert sum(result.json()["duplicate_rows"] for result in concurrent_results) == 2
        concurrent_guests = await client.get(
            f"/api/v1/events/{concurrent_event_id}/guests?limit=100&offset=0",
            headers=headers,
        )
        assert concurrent_guests.status_code == 200
        assert len(concurrent_guests.json()["items"]) == 2
        untouched_third_preview = await client.get(
            f"/api/v1/events/{concurrent_event_id}/imports/{concurrent_batch_ids[2]}/rows",
            headers=headers,
        )
        assert untouched_third_preview.status_code == 200
        assert [item["status"] for item in untouched_third_preview.json()["items"]] == [
            "VALID",
            "VALID",
        ]

        second_event = await client.post(
            "/api/v1/events",
            headers=headers,
            json={**event_payload, "event_name": "Synthetic Cross-Scope Wedding"},
        )
        assert second_event.status_code == 201
        cross_scope = await client.get(
            f"/api/v1/events/{second_event.json()['id']}/imports/{batch_id}",
            headers=headers,
        )
        unauthenticated = await client.get(f"/api/v1/events/{event_id}/imports/{batch_id}")
        assert cross_scope.status_code == 404
        assert unauthenticated.status_code == 401

        return {
            "event_id": event_id,
            "second_event_id": second_event.json()["id"],
            "batch_id": batch_id,
            "row_id": row_id,
            "contextual_event_id": contextual_event_id,
            "contextual_batch_id": contextual_batch_id,
        }


async def exercise_partial_selection_api(
    settings: Settings,
    runtime: DatabaseRuntime,
    username: str,
) -> dict[str, object]:
    app = create_app(
        settings,
        database_runtime=runtime,
        authentication_security=AuthenticationSecurity.from_settings(settings),
    )
    transport = httpx.ASGITransport(app=app)
    async with (
        app.router.lifespan_context(app),
        httpx.AsyncClient(transport=transport, base_url="http://synthetic.test") as client,
    ):
        login = await client.post(
            "/api/v1/auth/login",
            json={"username": username, "password": "synthetic administrator password"},
        )
        assert login.status_code == 200
        headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
        base_event = {
            "bride_name": "Synthetic",
            "groom_name": "Selection",
            "event_datetime": "2027-06-12T14:30:00+03:00",
            "venue": "Synthetic Hall",
            "max_companions_per_invitation": 2,
        }

        async def create_event(name: str) -> str:
            response = await client.post(
                "/api/v1/events",
                headers=headers,
                json={**base_event, "event_name": name},
            )
            assert response.status_code == 201, response.text
            return str(response.json()["id"])

        async def upload(
            event_id: str,
            row_count: int,
            prefix: str,
            *,
            invalid_last_row: bool = False,
        ) -> tuple[str, list[dict[str, object]]]:
            response = await client.post(
                f"/api/v1/events/{event_id}/imports",
                headers=headers,
                files={
                    "workbook": (
                        f"synthetic-{prefix.casefold()}.xlsx",
                        selectable_workbook(
                            row_count,
                            prefix,
                            invalid_last_row=invalid_last_row,
                        ),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert response.status_code == 201, response.text
            batch_id = str(response.json()["id"])
            rows = await client.get(
                f"/api/v1/events/{event_id}/imports/{batch_id}/rows?limit=100",
                headers=headers,
            )
            assert rows.status_code == 200
            return batch_id, rows.json()["items"]

        twenty_event_id = await create_event("Synthetic Twenty Row Selection")
        twenty_batch_id, twenty_rows = await upload(twenty_event_id, 20, "Twenty")
        first_five_ids = [str(row["id"]) for row in twenty_rows[:5]]
        first_five = await client.post(
            f"/api/v1/events/{twenty_event_id}/imports/{twenty_batch_id}/progress",
            headers=headers,
            json={"staged_row_ids": first_five_ids},
        )
        assert first_five.status_code == 200, first_five.text
        assert first_five.json()["imported_rows"] == 5
        assert first_five.json()["reconciliation"]["rows_already_imported"] == 5
        assert first_five.json()["reconciliation"]["currently_valid_staged_rows"] == 15
        assert first_five.json()["reconciliation"]["overall_batch_state"] == ("PARTIALLY_IMPORTED")

        next_three_ids = [str(row["id"]) for row in twenty_rows[5:8]]
        next_three = await client.post(
            f"/api/v1/events/{twenty_event_id}/imports/{twenty_batch_id}/progress",
            headers=headers,
            json={"staged_row_ids": next_three_ids},
        )
        assert next_three.status_code == 200, next_three.text
        assert next_three.json()["imported_rows"] == 3
        assert next_three.json()["reconciliation"]["rows_already_imported"] == 8
        assert next_three.json()["reconciliation"]["currently_valid_staged_rows"] == 12

        continuation = await client.post(
            f"/api/v1/events/{twenty_event_id}/imports/{twenty_batch_id}/progress",
            headers=headers,
        )
        assert continuation.status_code == 200, continuation.text
        assert continuation.json()["imported_rows"] == 12
        assert continuation.json()["duplicate_rows"] == 0
        assert continuation.json()["reconciliation"]["rows_already_imported"] == 20
        assert continuation.json()["reconciliation"]["overall_batch_state"] == "RECONCILED"

        fifty_event_id = await create_event("Synthetic Fifty Row Selection")
        fifty_batch_id, fifty_rows = await upload(fifty_event_id, 50, "Fifty")
        selected_ten_ids = [str(row["id"]) for row in fifty_rows[:10]]
        selected_ten = await client.post(
            f"/api/v1/events/{fifty_event_id}/imports/{fifty_batch_id}/progress",
            headers=headers,
            json={"staged_row_ids": selected_ten_ids},
        )
        assert selected_ten.status_code == 200, selected_ten.text
        assert selected_ten.json()["imported_rows"] == 10
        assert selected_ten.json()["reconciliation"]["rows_already_imported"] == 10
        assert selected_ten.json()["reconciliation"]["currently_valid_staged_rows"] == 40
        repeated_ten = await client.post(
            f"/api/v1/events/{fifty_event_id}/imports/{fifty_batch_id}/progress",
            headers=headers,
            json={"staged_row_ids": selected_ten_ids},
        )
        assert repeated_ten.status_code == 200
        assert repeated_ten.json()["imported_rows"] == 0

        foreign_batch_id, foreign_rows = await upload(fifty_event_id, 2, "Foreign")
        before_atomic_refusals = await client.get(
            f"/api/v1/events/{fifty_event_id}/guests?limit=100",
            headers=headers,
        )
        assert before_atomic_refusals.status_code == 200
        assert len(before_atomic_refusals.json()["items"]) == 10
        valid_unselected_id = str(fifty_rows[10]["id"])
        wrong_batch = await client.post(
            f"/api/v1/events/{fifty_event_id}/imports/{fifty_batch_id}/progress",
            headers=headers,
            json={"staged_row_ids": [valid_unselected_id, str(foreign_rows[0]["id"])]},
        )
        nonexistent = await client.post(
            f"/api/v1/events/{fifty_event_id}/imports/{fifty_batch_id}/progress",
            headers=headers,
            json={"staged_row_ids": [valid_unselected_id, str(uuid4())]},
        )
        assert wrong_batch.status_code == nonexistent.status_code == 404

        invalid_batch_id, invalid_rows = await upload(
            fifty_event_id,
            2,
            "Invalid",
            invalid_last_row=True,
        )
        invalid_selection = await client.post(
            f"/api/v1/events/{fifty_event_id}/imports/{invalid_batch_id}/progress",
            headers=headers,
            json={"staged_row_ids": [str(row["id"]) for row in invalid_rows]},
        )
        assert invalid_selection.status_code == 409
        assert "non-progressable" in invalid_selection.text

        empty_selection = await client.post(
            f"/api/v1/events/{fifty_event_id}/imports/{fifty_batch_id}/progress",
            headers=headers,
            json={"staged_row_ids": []},
        )
        duplicate_selection = await client.post(
            f"/api/v1/events/{fifty_event_id}/imports/{fifty_batch_id}/progress",
            headers=headers,
            json={"staged_row_ids": [valid_unselected_id, valid_unselected_id]},
        )
        unauthenticated = await client.post(
            f"/api/v1/events/{fifty_event_id}/imports/{fifty_batch_id}/progress",
            json={"staged_row_ids": [valid_unselected_id]},
        )
        assert empty_selection.status_code == duplicate_selection.status_code == 422
        assert unauthenticated.status_code == 401
        after_atomic_refusals = await client.get(
            f"/api/v1/events/{fifty_event_id}/guests?limit=100",
            headers=headers,
        )
        assert after_atomic_refusals.status_code == 200
        assert len(after_atomic_refusals.json()["items"]) == 10
        privacy_text = " ".join(
            response.text
            for response in (wrong_batch, nonexistent, invalid_selection, empty_selection)
        )
        assert "@example.test" not in privacy_text
        assert "Synthetic Fifty" not in privacy_text

        concurrency_event_id = await create_event("Synthetic Concurrent Selection")
        concurrency_batch_id, concurrency_rows = await upload(
            concurrency_event_id,
            20,
            "ConcurrentSelection",
        )
        concurrency_ids = [str(row["id"]) for row in concurrency_rows]
        overlapping = await asyncio.gather(
            client.post(
                f"/api/v1/events/{concurrency_event_id}/imports/{concurrency_batch_id}/progress",
                headers=headers,
                json={"staged_row_ids": concurrency_ids[:10]},
            ),
            client.post(
                f"/api/v1/events/{concurrency_event_id}/imports/{concurrency_batch_id}/progress",
                headers=headers,
                json={"staged_row_ids": concurrency_ids[5:15]},
            ),
        )
        assert all(response.status_code == 200 for response in overlapping)
        assert sum(response.json()["imported_rows"] for response in overlapping) == 15
        non_overlapping = await asyncio.gather(
            client.post(
                f"/api/v1/events/{concurrency_event_id}/imports/{concurrency_batch_id}/progress",
                headers=headers,
                json={"staged_row_ids": concurrency_ids[15:18]},
            ),
            client.post(
                f"/api/v1/events/{concurrency_event_id}/imports/{concurrency_batch_id}/progress",
                headers=headers,
                json={"staged_row_ids": concurrency_ids[18:]},
            ),
        )
        assert all(response.status_code == 200 for response in non_overlapping)
        assert sum(response.json()["imported_rows"] for response in non_overlapping) == 5
        concurrent_guests = await client.get(
            f"/api/v1/events/{concurrency_event_id}/guests?limit=100",
            headers=headers,
        )
        assert concurrent_guests.status_code == 200
        assert len(concurrent_guests.json()["items"]) == 20

        return {
            "event_ids": [twenty_event_id, fifty_event_id, concurrency_event_id],
            "twenty_batch_id": twenty_batch_id,
            "fifty_batch_id": fifty_batch_id,
            "selected_ten_ids": selected_ten_ids,
            "concurrency_batch_id": concurrency_batch_id,
            "foreign_batch_id": foreign_batch_id,
            "invalid_batch_id": invalid_batch_id,
        }


def test_partial_selection_postgresql_exact_sets_concurrency_and_side_effects(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    settings = configure_test_environment(monkeypatch, tmp_path)
    assert disposable_alembic.main(["upgrade", "head"]) == 0
    assert disposable_alembic.main(["check"]) == 0

    assert DISPOSABLE_URL is not None
    engine = create_engine(DISPOSABLE_URL, pool_pre_ping=True)
    runtime = DatabaseRuntime(settings)
    administrator_id: UUID | None = None
    identifiers: dict[str, object] = {}
    username = f"s4-partial-admin-{uuid4().hex[:8]}"
    try:
        security = AuthenticationSecurity.from_settings(settings)
        with Session(engine, expire_on_commit=False) as session:
            administrator = AdministratorProvisioningService(
                AdministratorRepository(session),
                security.passwords,
            ).provision(username, "synthetic administrator password")
            administrator_id = administrator.id

        identifiers = asyncio.run(exercise_partial_selection_api(settings, runtime, username))
        event_ids = [UUID(value) for value in identifiers["event_ids"]]  # type: ignore[arg-type]
        twenty_event_id, fifty_event_id, concurrency_event_id = event_ids
        twenty_batch_id = UUID(str(identifiers["twenty_batch_id"]))
        fifty_batch_id = UUID(str(identifiers["fifty_batch_id"]))
        concurrency_batch_id = UUID(str(identifiers["concurrency_batch_id"]))
        selected_ten_ids = {
            UUID(value)
            for value in identifiers["selected_ten_ids"]  # type: ignore[union-attr]
        }

        with Session(engine) as session:
            assert (
                session.scalar(
                    select(func.count()).select_from(Guest).where(Guest.event_id == twenty_event_id)
                )
                == 20
            )
            twenty_rows = list(
                session.execute(
                    select(ImportBatchRow).where(ImportBatchRow.batch_id == twenty_batch_id)
                ).scalars()
            )
            assert len(twenty_rows) == 20
            assert all(row.status is ImportRowStatus.IMPORTED for row in twenty_rows)

            fifty_rows = list(
                session.execute(
                    select(ImportBatchRow)
                    .where(ImportBatchRow.batch_id == fifty_batch_id)
                    .order_by(ImportBatchRow.source_row_number)
                ).scalars()
            )
            assert len(fifty_rows) == 50
            assert sum(row.status is ImportRowStatus.IMPORTED for row in fifty_rows) == 10
            assert sum(row.status is ImportRowStatus.VALID for row in fifty_rows) == 40
            assert {row.id for row in fifty_rows if row.guest_id is not None} == selected_ten_ids
            assert all(row.row_version == 1 for row in fifty_rows)
            assert (
                session.scalar(
                    select(func.count())
                    .select_from(ImportRowRevision)
                    .where(ImportRowRevision.row_id.in_([row.id for row in fifty_rows]))
                )
                == 0
            )
            selected_guests = list(
                session.execute(
                    select(Guest).where(Guest.event_id == fifty_event_id).order_by(Guest.full_name)
                ).scalars()
            )
            assert len(selected_guests) == 10
            assert selected_guests[0].phone_number == "+255700000001"
            assert selected_guests[0].email == "fifty-001@example.test"
            assert selected_guests[0].invitation_type == "Family"
            assert selected_guests[0].allowed_companions == 1
            assert selected_guests[0].table_number == "T-001"
            assert selected_guests[0].category == "Friends"

            concurrency_rows = list(
                session.execute(
                    select(ImportBatchRow).where(ImportBatchRow.batch_id == concurrency_batch_id)
                ).scalars()
            )
            assert len(concurrency_rows) == 20
            assert all(row.status is ImportRowStatus.IMPORTED for row in concurrency_rows)
            assert len({row.guest_id for row in concurrency_rows}) == 20
            assert (
                session.scalar(
                    select(func.count())
                    .select_from(Guest)
                    .where(Guest.event_id == concurrency_event_id)
                )
                == 20
            )

            for model in (
                GuestQrCredential,
                GeneratedGuestCard,
                GuestAttendanceRecord,
                GateScanAttempt,
                GuestAttendanceCorrection,
            ):
                assert (
                    session.scalar(
                        select(func.count()).select_from(model).where(model.event_id.in_(event_ids))
                    )
                    == 0
                )

        assert "@example.test" not in caplog.text
        assert "Synthetic Fifty" not in caplog.text
    finally:
        runtime.close()
        if administrator_id is not None:
            with Session(engine) as session:
                events = list(
                    session.execute(
                        select(Event).where(Event.created_by == administrator_id)
                    ).scalars()
                )
                event_ids = [event.id for event in events]
                if event_ids:
                    session.execute(delete(ImportBatch).where(ImportBatch.event_id.in_(event_ids)))
                for event in events:
                    session.delete(event)
                session.flush()
                administrator = session.get(Administrator, administrator_id)
                if administrator is not None:
                    session.delete(administrator)
                session.commit()
        engine.dispose()


def test_m4_schema_api_correction_reimport_and_immutability(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    settings = configure_test_environment(monkeypatch, tmp_path)
    assert disposable_alembic.main(["upgrade", "head"]) == 0
    assert disposable_alembic.main(["check"]) == 0

    assert DISPOSABLE_URL is not None
    engine = create_engine(DISPOSABLE_URL, pool_pre_ping=True)
    runtime = DatabaseRuntime(settings)
    administrator_id: UUID | None = None
    identifiers: dict[str, str] = {}
    username = f"m4-admin-{uuid4().hex[:8]}"
    try:
        security = AuthenticationSecurity.from_settings(settings)
        with Session(engine, expire_on_commit=False) as session:
            administrator = AdministratorProvisioningService(
                AdministratorRepository(session),
                security.passwords,
            ).provision(username, "synthetic administrator password")
            administrator_id = administrator.id

        identifiers = asyncio.run(exercise_import_api(settings, runtime, username))

        with Session(engine) as session:
            contextual_batch = session.get(
                ImportBatch,
                UUID(identifiers["contextual_batch_id"]),
            )
            assert contextual_batch is not None
            assert contextual_batch.event_id == UUID(identifiers["contextual_event_id"])
            assert contextual_batch.candidate_row_count == 3
            assert contextual_batch.ignored_row_count == 5
            ignored_rows = contextual_batch.structure_snapshot["ignored_rows"]
            assert [item["reason"] for item in ignored_rows].count("PRESENTATION") == 2
            assert [item["reason"] for item in ignored_rows].count("BLANK") == 1
            assert [item["reason"] for item in ignored_rows].count("FOOTER") == 2
            contextual_rows = list(
                session.execute(
                    select(ImportBatchRow)
                    .where(ImportBatchRow.batch_id == contextual_batch.id)
                    .order_by(ImportBatchRow.source_row_number)
                ).scalars()
            )
            assert [row.source_row_number for row in contextual_rows] == [5, 6, 7]
            assert all(row.status is ImportRowStatus.VALID for row in contextual_rows)
            assert all(row.validation_errors == [] for row in contextual_rows)

        inspector = inspect(engine)
        assert {"import_batches", "import_batch_rows", "import_row_revisions"} <= set(
            inspector.get_table_names()
        )
        assert {index["name"] for index in inspector.get_indexes("import_batch_rows")} >= {
            "ix_import_batch_rows_batch_status_source"
        }
        assert {
            constraint["name"]
            for constraint in inspector.get_check_constraints("import_batch_rows")
        } >= {"ck_import_batch_rows_guest_link", "ck_import_batch_rows_status"}

        with Session(engine) as session:
            row = session.get(ImportBatchRow, UUID(identifiers["row_id"]))
            assert row is not None
            row.source_row_number += 1
            with pytest.raises(IntegrityError):
                session.commit()
            session.rollback()

            revision = (
                session.execute(
                    select(ImportRowRevision)
                    .where(ImportRowRevision.row_id == UUID(identifiers["row_id"]))
                    .order_by(ImportRowRevision.revision)
                )
                .scalars()
                .first()
            )
            assert revision is not None
            revision.changed_values = {"forbidden": "mutation"}
            with pytest.raises(IntegrityError):
                session.commit()
            session.rollback()

            revision = session.get(ImportRowRevision, revision.id)
            assert revision is not None
            session.delete(revision)
            with pytest.raises(IntegrityError):
                session.commit()
            session.rollback()

        evidence_files = list(settings.IMPORT_EVIDENCE_STORAGE_ROOT.rglob("*.xlsx"))
        assert len(evidence_files) == 10
        assert all("synthetic.xlsx" not in path.name for path in evidence_files)
    finally:
        runtime.close()
        if administrator_id is not None:
            with Session(engine) as session:
                events = list(
                    session.execute(
                        select(Event).where(Event.created_by == administrator_id)
                    ).scalars()
                )
                event_ids = [event.id for event in events]
                if event_ids:
                    session.execute(delete(ImportBatch).where(ImportBatch.event_id.in_(event_ids)))
                for event in events:
                    session.delete(event)
                session.flush()
                administrator = session.get(Administrator, administrator_id)
                if administrator is not None:
                    session.delete(administrator)
                session.commit()
        engine.dispose()


async def exercise_valid_row_correction_api(
    settings: Settings,
    runtime: DatabaseRuntime,
    username: str,
) -> dict[str, object]:
    app = create_app(
        settings,
        database_runtime=runtime,
        authentication_security=AuthenticationSecurity.from_settings(settings),
    )
    transport = httpx.ASGITransport(app=app)
    async with (
        app.router.lifespan_context(app),
        httpx.AsyncClient(transport=transport, base_url="http://synthetic.test") as client,
    ):
        login = await client.post(
            "/api/v1/auth/login",
            json={"username": username, "password": "synthetic administrator password"},
        )
        assert login.status_code == 200
        headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
        event_payload = {
            "bride_name": "Synthetic",
            "groom_name": "Correction",
            "event_datetime": "2027-06-12T14:30:00+03:00",
            "venue": "Synthetic Hall",
            "max_companions_per_invitation": 2,
        }

        async def create_event(name: str) -> str:
            response = await client.post(
                "/api/v1/events",
                headers=headers,
                json={**event_payload, "event_name": name},
            )
            assert response.status_code == 201, response.text
            return str(response.json()["id"])

        async def upload(
            event_id: str, count: int, prefix: str
        ) -> tuple[str, list[dict[str, object]]]:
            response = await client.post(
                f"/api/v1/events/{event_id}/imports",
                headers=headers,
                files={
                    "workbook": (
                        f"synthetic-{prefix.casefold()}.xlsx",
                        selectable_workbook(count, prefix),
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    )
                },
            )
            assert response.status_code == 201, response.text
            batch_id = str(response.json()["id"])
            staged = await client.get(
                f"/api/v1/events/{event_id}/imports/{batch_id}/rows?limit=100",
                headers=headers,
            )
            assert staged.status_code == 200
            return batch_id, staged.json()["items"]

        main_event = await create_event("Synthetic Valid Row Correction")
        main_batch, main_rows = await upload(main_event, 20, "ValidCorrection")
        initial_values = dict(main_rows[5]["current_values"])  # type: ignore[arg-type]
        first_five_ids = [str(row["id"]) for row in main_rows[:5]]
        progressed = await client.post(
            f"/api/v1/events/{main_event}/imports/{main_batch}/progress",
            headers=headers,
            json={"staged_row_ids": first_five_ids},
        )
        assert progressed.status_code == 200
        assert progressed.json()["imported_rows"] == 5
        assert progressed.json()["reconciliation"]["currently_valid_staged_rows"] == 15
        target_id = str(main_rows[5]["id"])

        first_correction = await client.patch(
            f"/api/v1/events/{main_event}/imports/{main_batch}/rows/{target_id}",
            headers=headers,
            json={"expected_version": 1, "values": {"allowed_companions": 3}},
        )
        assert first_correction.status_code == 200, first_correction.text
        assert first_correction.json()["status"] == "PENDING_REVALIDATION"
        assert first_correction.json()["row_version"] == 2
        assert first_correction.json()["original_values"] == initial_values
        assert first_correction.json()["guest_id"] is None

        premature_progress = await client.post(
            f"/api/v1/events/{main_event}/imports/{main_batch}/progress",
            headers=headers,
            json={"staged_row_ids": [target_id]},
        )
        assert premature_progress.status_code == 409
        guests_before_revalidation = await client.get(
            f"/api/v1/events/{main_event}/guests?limit=100",
            headers=headers,
        )
        assert guests_before_revalidation.status_code == 200
        assert len(guests_before_revalidation.json()["items"]) == 5

        invalid = await client.post(
            f"/api/v1/events/{main_event}/imports/{main_batch}/rows/{target_id}/revalidate",
            headers=headers,
        )
        assert invalid.status_code == 200, invalid.text
        assert invalid.json()["status"] == "INVALID"
        assert invalid.json()["row_version"] == 3
        assert [item["code"] for item in invalid.json()["validation_errors"]] == [
            "COMPANION_LIMIT_EXCEEDED"
        ]

        stale = await client.patch(
            f"/api/v1/events/{main_event}/imports/{main_batch}/rows/{target_id}",
            headers=headers,
            json={
                "expected_version": 1,
                "values": {"allowed_companions": initial_values["allowed_companions"]},
            },
        )
        assert stale.status_code == 409
        assert "stale" in stale.text

        missing_auth = await client.patch(
            f"/api/v1/events/{main_event}/imports/{main_batch}/rows/{target_id}",
            json={
                "expected_version": 3,
                "values": {"allowed_companions": initial_values["allowed_companions"]},
            },
        )
        invalid_auth = await client.post(
            f"/api/v1/events/{main_event}/imports/{main_batch}/rows/{target_id}/revalidate",
            headers={"Authorization": "Bearer invalid-token"},
        )
        assert missing_auth.status_code == invalid_auth.status_code == 401

        second_event = await create_event("Synthetic Wrong Scope Correction")
        wrong_event = await client.patch(
            f"/api/v1/events/{second_event}/imports/{main_batch}/rows/{target_id}",
            headers=headers,
            json={
                "expected_version": 3,
                "values": {"allowed_companions": initial_values["allowed_companions"]},
            },
        )
        wrong_batch = await client.patch(
            f"/api/v1/events/{main_event}/imports/{uuid4()}/rows/{target_id}",
            headers=headers,
            json={
                "expected_version": 3,
                "values": {"allowed_companions": initial_values["allowed_companions"]},
            },
        )
        wrong_row = await client.patch(
            f"/api/v1/events/{main_event}/imports/{main_batch}/rows/{uuid4()}",
            headers=headers,
            json={
                "expected_version": 3,
                "values": {"allowed_companions": initial_values["allowed_companions"]},
            },
        )
        assert wrong_event.status_code == wrong_batch.status_code == wrong_row.status_code == 404

        restored = await client.patch(
            f"/api/v1/events/{main_event}/imports/{main_batch}/rows/{target_id}",
            headers=headers,
            json={
                "expected_version": 3,
                "values": {"allowed_companions": initial_values["allowed_companions"]},
            },
        )
        assert restored.status_code == 200, restored.text
        assert restored.json()["status"] == "PENDING_REVALIDATION"
        assert restored.json()["row_version"] == 4
        final_target = await client.post(
            f"/api/v1/events/{main_event}/imports/{main_batch}/rows/{target_id}/revalidate",
            headers=headers,
        )
        assert final_target.status_code == 200, final_target.text
        assert final_target.json()["status"] == "VALID"
        assert final_target.json()["validation_errors"] == []
        assert final_target.json()["row_version"] == 5
        assert final_target.json()["current_values"] == initial_values
        assert final_target.json()["original_values"] == initial_values
        assert final_target.json()["guest_id"] is None

        imported_refusal = await client.patch(
            f"/api/v1/events/{main_event}/imports/{main_batch}/rows/{main_rows[0]['id']}",
            headers=headers,
            json={"expected_version": 1, "values": {"category": "Protected Attempt"}},
        )
        assert imported_refusal.status_code == 409
        assert "imported or guest-linked" in imported_refusal.text

        valid_to_valid_id = str(main_rows[6]["id"])
        valid_to_valid = await client.patch(
            f"/api/v1/events/{main_event}/imports/{main_batch}/rows/{valid_to_valid_id}",
            headers=headers,
            json={"expected_version": 1, "values": {"category": "Colleagues"}},
        )
        assert valid_to_valid.status_code == 200
        assert valid_to_valid.json()["status"] == "PENDING_REVALIDATION"
        blocked_valid_to_valid = await client.post(
            f"/api/v1/events/{main_event}/imports/{main_batch}/progress",
            headers=headers,
            json={"staged_row_ids": [valid_to_valid_id]},
        )
        assert blocked_valid_to_valid.status_code == 409
        revalidated_valid_to_valid = await client.post(
            f"/api/v1/events/{main_event}/imports/{main_batch}/rows/{valid_to_valid_id}/revalidate",
            headers=headers,
        )
        assert revalidated_valid_to_valid.status_code == 200
        assert revalidated_valid_to_valid.json()["status"] == "VALID"

        persisted_event = await create_event("Synthetic Persisted Duplicate Correction")
        persisted_batch, persisted_rows = await upload(persisted_event, 2, "PersistedDuplicate")
        persisted_progress = await client.post(
            f"/api/v1/events/{persisted_event}/imports/{persisted_batch}/progress",
            headers=headers,
            json={"staged_row_ids": [str(persisted_rows[0]["id"])]},
        )
        assert persisted_progress.status_code == 200
        persisted_duplicate = await client.patch(
            f"/api/v1/events/{persisted_event}/imports/{persisted_batch}/rows/"
            f"{persisted_rows[1]['id']}",
            headers=headers,
            json={
                "expected_version": 1,
                "values": {"email": persisted_rows[0]["current_values"]["email"]},  # type: ignore[index]
            },
        )
        assert persisted_duplicate.status_code == 200
        persisted_duplicate_result = await client.post(
            f"/api/v1/events/{persisted_event}/imports/{persisted_batch}/rows/"
            f"{persisted_rows[1]['id']}/revalidate",
            headers=headers,
        )
        assert persisted_duplicate_result.status_code == 200
        assert persisted_duplicate_result.json()["status"] == "DUPLICATE"
        assert persisted_duplicate_result.json()["validation_errors"][0]["code"] == (
            "DUPLICATE_STAGED_OR_EXISTING"
        )

        within_event = await create_event("Synthetic Within Batch Duplicate Correction")
        within_batch, within_rows = await upload(within_event, 3, "WithinDuplicate")
        within_duplicate = await client.patch(
            f"/api/v1/events/{within_event}/imports/{within_batch}/rows/{within_rows[2]['id']}",
            headers=headers,
            json={
                "expected_version": 1,
                "values": {"email": within_rows[1]["current_values"]["email"]},  # type: ignore[index]
            },
        )
        assert within_duplicate.status_code == 200
        within_duplicate_result = await client.post(
            f"/api/v1/events/{within_event}/imports/{within_batch}/rows/"
            f"{within_rows[2]['id']}/revalidate",
            headers=headers,
        )
        assert within_duplicate_result.status_code == 200
        assert within_duplicate_result.json()["status"] == "DUPLICATE"
        assert within_duplicate_result.json()["validation_errors"][0]["code"] == (
            "DUPLICATE_STAGED_OR_EXISTING"
        )

        preview_event = await create_event("Synthetic Historical Preview Correction")
        _, preview_rows = await upload(preview_event, 2, "PreviewHistory")
        _, preview_retry_rows = await upload(preview_event, 2, "PreviewHistory")
        assert all(row["status"] == "VALID" for row in preview_rows)
        assert all(row["status"] == "VALID" for row in preview_retry_rows)

        concurrent_event = await create_event("Synthetic Concurrent Correction")
        concurrent_batch, concurrent_rows = await upload(
            concurrent_event, 2, "ConcurrentCorrection"
        )
        concurrent_id = str(concurrent_rows[0]["id"])
        simultaneous_corrections = await asyncio.gather(
            client.patch(
                f"/api/v1/events/{concurrent_event}/imports/{concurrent_batch}/rows/"
                f"{concurrent_id}",
                headers=headers,
                json={"expected_version": 1, "values": {"allowed_companions": 3}},
            ),
            client.patch(
                f"/api/v1/events/{concurrent_event}/imports/{concurrent_batch}/rows/"
                f"{concurrent_id}",
                headers=headers,
                json={"expected_version": 1, "values": {"category": "Colleagues"}},
            ),
        )
        assert sorted(response.status_code for response in simultaneous_corrections) == [200, 409]
        rejected_correction = next(
            response for response in simultaneous_corrections if response.status_code == 409
        )
        assert "stale" in rejected_correction.text
        concurrent_state = await client.get(
            f"/api/v1/events/{concurrent_event}/imports/{concurrent_batch}/rows?limit=100",
            headers=headers,
        )
        assert concurrent_state.status_code == 200
        assert concurrent_state.json()["items"][0]["status"] == "PENDING_REVALIDATION"
        assert concurrent_state.json()["items"][0]["row_version"] == 2

        race_event = await create_event("Synthetic Correction Progression Race")
        race_batch, race_rows = await upload(race_event, 1, "CorrectionRace")
        race_id = str(race_rows[0]["id"])
        race_original = dict(race_rows[0]["current_values"])  # type: ignore[arg-type]
        race_correction, race_progression = await asyncio.gather(
            client.patch(
                f"/api/v1/events/{race_event}/imports/{race_batch}/rows/{race_id}",
                headers=headers,
                json={"expected_version": 1, "values": {"allowed_companions": 3}},
            ),
            client.post(
                f"/api/v1/events/{race_event}/imports/{race_batch}/progress",
                headers=headers,
                json={"staged_row_ids": [race_id]},
            ),
        )
        assert sorted([race_correction.status_code, race_progression.status_code]) == [200, 409]
        race_state_response = await client.get(
            f"/api/v1/events/{race_event}/imports/{race_batch}/rows?limit=100",
            headers=headers,
        )
        race_guests = await client.get(
            f"/api/v1/events/{race_event}/guests?limit=100",
            headers=headers,
        )
        race_state = race_state_response.json()["items"][0]
        if race_correction.status_code == 200:
            assert race_progression.status_code == 409
            assert race_state["status"] == "PENDING_REVALIDATION"
            assert race_state["current_values"]["allowed_companions"] == 3
            assert race_state["guest_id"] is None
            assert race_guests.json()["items"] == []
        else:
            assert race_progression.status_code == 200
            assert race_progression.json()["imported_rows"] == 1
            assert race_state["status"] == "IMPORTED"
            assert race_state["current_values"] == race_original
            assert race_state["guest_id"] is not None
            assert len(race_guests.json()["items"]) == 1

        final_main_rows = await client.get(
            f"/api/v1/events/{main_event}/imports/{main_batch}/rows?limit=100",
            headers=headers,
        )
        final_main_guests = await client.get(
            f"/api/v1/events/{main_event}/guests?limit=100",
            headers=headers,
        )
        assert final_main_rows.status_code == final_main_guests.status_code == 200
        assert sum(row["status"] == "IMPORTED" for row in final_main_rows.json()["items"]) == 5
        assert sum(row["status"] == "VALID" for row in final_main_rows.json()["items"]) == 15
        assert len(final_main_guests.json()["items"]) == 5
        assert (
            next(row for row in final_main_rows.json()["items"] if row["id"] == target_id)[
                "guest_id"
            ]
            is None
        )

        return {
            "event_ids": [
                main_event,
                second_event,
                persisted_event,
                within_event,
                preview_event,
                concurrent_event,
                race_event,
            ],
            "main_event": main_event,
            "main_batch": main_batch,
            "target_id": target_id,
            "concurrent_id": concurrent_id,
            "race_id": race_id,
        }


def test_valid_unlinked_correction_postgresql_state_history_duplicates_and_races(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    settings = configure_test_environment(monkeypatch, tmp_path)
    assert disposable_alembic.main(["upgrade", "head"]) == 0
    assert disposable_alembic.main(["check"]) == 0

    assert DISPOSABLE_URL is not None
    engine = create_engine(DISPOSABLE_URL, pool_pre_ping=True)
    runtime = DatabaseRuntime(settings)
    administrator_id: UUID | None = None
    identifiers: dict[str, object] = {}
    username = f"valid-correction-admin-{uuid4().hex[:8]}"
    try:
        security = AuthenticationSecurity.from_settings(settings)
        with Session(engine, expire_on_commit=False) as session:
            administrator = AdministratorProvisioningService(
                AdministratorRepository(session),
                security.passwords,
            ).provision(username, "synthetic administrator password")
            administrator_id = administrator.id

        identifiers = asyncio.run(exercise_valid_row_correction_api(settings, runtime, username))
        event_ids = [UUID(str(value)) for value in identifiers["event_ids"]]  # type: ignore[union-attr]
        with Session(engine) as session:
            target = session.get(ImportBatchRow, UUID(str(identifiers["target_id"])))
            assert target is not None
            assert target.status is ImportRowStatus.VALID
            assert target.row_version == 5
            assert target.guest_id is None
            assert target.current_values == target.original_values
            target_revisions = list(
                session.execute(
                    select(ImportRowRevision)
                    .where(ImportRowRevision.row_id == target.id)
                    .order_by(ImportRowRevision.revision)
                ).scalars()
            )
            assert [revision.revision for revision in target_revisions] == [2, 3, 4, 5]
            assert [revision.kind for revision in target_revisions] == [
                ImportRevisionKind.CORRECTION,
                ImportRevisionKind.REVALIDATION,
                ImportRevisionKind.CORRECTION,
                ImportRevisionKind.REVALIDATION,
            ]

            main_batch = session.get(ImportBatch, UUID(str(identifiers["main_batch"])))
            assert main_batch is not None
            assert main_batch.status is ImportBatchStatus.PARTIALLY_IMPORTED
            main_rows = list(
                session.execute(
                    select(ImportBatchRow).where(ImportBatchRow.batch_id == main_batch.id)
                ).scalars()
            )
            assert len(main_rows) == 20
            assert sum(row.status is ImportRowStatus.IMPORTED for row in main_rows) == 5
            assert sum(row.status is ImportRowStatus.VALID for row in main_rows) == 15
            assert (
                session.scalar(
                    select(func.count())
                    .select_from(Guest)
                    .where(Guest.event_id == UUID(str(identifiers["main_event"])))
                )
                == 5
            )

            concurrent_revisions = session.scalar(
                select(func.count())
                .select_from(ImportRowRevision)
                .where(ImportRowRevision.row_id == UUID(str(identifiers["concurrent_id"])))
            )
            assert concurrent_revisions == 1

            for model in (
                GuestQrCredential,
                GeneratedGuestCard,
                GuestAttendanceRecord,
                GateScanAttempt,
                GuestAttendanceCorrection,
            ):
                assert (
                    session.scalar(
                        select(func.count()).select_from(model).where(model.event_id.in_(event_ids))
                    )
                    == 0
                )

        assert "@example.test" not in caplog.text
        assert "synthetic administrator password" not in caplog.text
    finally:
        runtime.close()
        if administrator_id is not None:
            with Session(engine) as session:
                events = list(
                    session.execute(
                        select(Event).where(Event.created_by == administrator_id)
                    ).scalars()
                )
                event_ids = [event.id for event in events]
                if event_ids:
                    session.execute(delete(ImportBatch).where(ImportBatch.event_id.in_(event_ids)))
                for event in events:
                    session.delete(event)
                session.flush()
                administrator = session.get(Administrator, administrator_id)
                if administrator is not None:
                    session.delete(administrator)
                session.commit()
        engine.dispose()


def test_overlength_import_value_is_rejected_before_guest_persistence(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    settings = configure_test_environment(monkeypatch, tmp_path)
    assert disposable_alembic.main(["upgrade", "head"]) == 0
    assert disposable_alembic.main(["check"]) == 0

    assert DISPOSABLE_URL is not None
    engine = create_engine(DISPOSABLE_URL, pool_pre_ping=True)
    runtime = DatabaseRuntime(settings)
    administrator_id: UUID | None = None
    username = f"length-boundary-admin-{uuid4().hex[:8]}"
    try:
        security = AuthenticationSecurity.from_settings(settings)
        with Session(engine, expire_on_commit=False) as session:
            administrator = AdministratorProvisioningService(
                AdministratorRepository(session),
                security.passwords,
            ).provision(username, "synthetic administrator password")
            administrator_id = administrator.id

        asyncio.run(exercise_overlength_persistence_boundary(settings, runtime, username))
    finally:
        runtime.close()
        if administrator_id is not None:
            with Session(engine) as session:
                events = list(
                    session.execute(
                        select(Event).where(Event.created_by == administrator_id)
                    ).scalars()
                )
                event_ids = [event.id for event in events]
                if event_ids:
                    session.execute(delete(ImportBatch).where(ImportBatch.event_id.in_(event_ids)))
                for event in events:
                    session.delete(event)
                session.flush()
                administrator = session.get(Administrator, administrator_id)
                if administrator is not None:
                    session.delete(administrator)
                session.commit()
        engine.dispose()
