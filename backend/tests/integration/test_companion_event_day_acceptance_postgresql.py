"""C2.7 integrated companion event-day acceptance on disposable PostgreSQL."""

from __future__ import annotations

import asyncio
import csv
import io
import json
import math
import os
from datetime import UTC, datetime
from pathlib import Path
from statistics import median
from time import perf_counter
from uuid import UUID, uuid4

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, select, text
from sqlalchemy import event as sqlalchemy_event
from sqlalchemy.orm import Session

import app.services.qr as qr_service_module
from app.cli import disposable_alembic
from app.core.security import AuthenticationSecurity
from app.core.settings import Settings
from app.db.runtime import DatabaseRuntime
from app.main import create_app
from app.models.attendance import (
    GateArrivalMode,
    GateScanAttempt,
    GateScanOutcome,
    GuestAttendanceRecord,
    GuestCompanionEntitlement,
)
from app.models.attendance_corrections import (
    AttendanceCorrectionDecision,
    AttendanceCorrectionReason,
    EffectiveAttendanceStatus,
    GuestAttendanceCorrection,
)
from app.models.event import Event, EventStatus
from app.models.guest import Guest, GuestStatus
from app.models.qr import GuestQrCredential, QrCredentialStatus
from app.qr.credentials import generate_qr_credential, hash_qr_credential
from app.reports.exports import encode_csv
from app.repositories.administrators import AdministratorRepository
from app.repositories.attendance import AttendanceRepository
from app.repositories.attendance_corrections import AttendanceCorrectionRepository
from app.repositories.reports import ReportRepository
from app.services.admin_provisioning import AdministratorProvisioningService
from app.services.attendance import AttendanceService, GateAdmissionCommand
from app.services.attendance_corrections import AttendanceCorrectionService
from app.services.reports import ReportService

DISPOSABLE_URL = os.getenv("WEDDING_DISPOSABLE_DATABASE_URL")
pytestmark = [
    pytest.mark.postgresql,
    pytest.mark.skipif(
        DISPOSABLE_URL is None,
        reason="WEDDING_DISPOSABLE_DATABASE_URL is not explicitly configured",
    ),
]

_ADMIN_PASSWORD = "synthetic administrator password"
_PRIVATE_EMAIL = "private-c27@example.test"
_PRIVATE_PHONE = "+255000000000"
_FORBIDDEN_RESPONSE_TERMS = (
    "phone_number",
    "email",
    "token_hash",
    "credential_fingerprint",
    "reason_note",
    "checked_in_by",
    "issued_by",
    "image_storage_key",
)


def configure_test_environment(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Settings:
    assert DISPOSABLE_URL is not None
    private_root = tmp_path / "private"
    monkeypatch.setenv("APP_ENV", "test")
    monkeypatch.setenv("DATABASE_URL", DISPOSABLE_URL)
    monkeypatch.setenv("JWT_SECRET", "synthetic-jwt-signing-secret-at-least-32-bytes")
    monkeypatch.setenv("JWT_ISSUER", "synthetic-issuer")
    monkeypatch.setenv("JWT_AUDIENCE", "synthetic-audience")
    monkeypatch.setenv("PRIVATE_STORAGE_ROOT", str(private_root))
    monkeypatch.setenv("IMPORT_EVIDENCE_STORAGE_ROOT", str(private_root / "imports"))
    monkeypatch.setenv("QR_IMAGE_STORAGE_ROOT", str(private_root / "qr"))
    monkeypatch.setenv("MAX_REPORT_EXPORT_ROWS", "5000")
    return Settings()


def make_event(administrator_id: UUID, label: str, *, status: EventStatus) -> Event:
    return Event(
        id=uuid4(),
        event_name=f"Synthetic C2.7 {label}",
        bride_name="Amina",
        groom_name="Baraka",
        event_datetime=datetime(2027, 9, 12, 14, 30, tzinfo=UTC),
        venue="Synthetic Hall",
        status=status,
        max_companions_per_invitation=2,
        created_by=administrator_id,
    )


def add_guest_credential(
    session: Session,
    event: Event,
    administrator_id: UUID,
    label: str,
    *,
    allowed: int = 2,
    guest_status: GuestStatus = GuestStatus.ACTIVE,
    credential_status: QrCredentialStatus = QrCredentialStatus.ACTIVE,
) -> tuple[Guest, str]:
    guest = Guest(
        id=uuid4(),
        event_id=event.id,
        full_name=f"Synthetic C2.7 {label}",
        phone_number=_PRIVATE_PHONE,
        email=_PRIVATE_EMAIL,
        invitation_type="Family",
        allowed_companions=allowed,
        table_number="T-7",
        category="Synthetic",
        status=guest_status,
    )
    raw = generate_qr_credential()
    credential_id = uuid4()
    credential = GuestQrCredential(
        id=credential_id,
        event_id=event.id,
        guest_id=guest.id,
        token_hash=hash_qr_credential(raw),
        image_storage_key=f"{event.id.hex}/{guest.id.hex}/{credential_id.hex}.png",
        image_sha256="a" * 64,
        status=credential_status,
        replaces_credential_id=None,
        issued_by=administrator_id,
        issued_at=datetime.now(UTC),
        deactivated_at=(
            datetime.now(UTC) if credential_status is QrCredentialStatus.INACTIVE else None
        ),
        superseded_at=(
            datetime.now(UTC) if credential_status is QrCredentialStatus.SUPERSEDED else None
        ),
    )
    session.add_all((guest, credential))
    return guest, raw


def admission_command(
    *,
    event_id: UUID,
    raw: str,
    administrator_id: UUID,
    mode: GateArrivalMode,
    entering: int,
    declared: int | None = None,
    request_id: UUID | None = None,
) -> GateAdmissionCommand:
    return GateAdmissionCommand(
        request_id=request_id or uuid4(),
        event_id=event_id,
        presented_credential=raw,
        arrival_mode=mode,
        declared_companions=declared,
        companions_entering_now=entering,
        administrator_id=administrator_id,
    )


def admit_once(engine, command: GateAdmissionCommand):
    with Session(engine, expire_on_commit=False, autoflush=False) as session:
        return AttendanceService(AttendanceRepository(session)).admit(command)


def inspect_once(engine, event_id: UUID, raw: str):
    with Session(engine, expire_on_commit=False, autoflush=False) as session:
        return AttendanceService(AttendanceRepository(session)).inspect(
            event_id=event_id,
            presented_credential=raw,
        )


def correct_once(
    engine,
    *,
    event_id: UUID,
    guest_id: UUID,
    administrator_id: UUID,
    revision: int,
    decision: AttendanceCorrectionDecision,
    expected_status: EffectiveAttendanceStatus | None = None,
    declared: int | None = None,
    admitted: int | None = None,
):
    with Session(engine, expire_on_commit=False, autoflush=False) as session:
        return AttendanceCorrectionService(
            AttendanceCorrectionRepository(session)
        ).record_correction(
            event_id=event_id,
            guest_id=guest_id,
            expected_correction_revision=revision,
            expected_effective_status=expected_status,
            decision=decision,
            reason_code=AttendanceCorrectionReason.OPERATOR_ERROR,
            reason_note=None,
            declared_companions=declared,
            companions_admitted=admitted,
            administrator_id=administrator_id,
        )


def report_once(engine, event_id: UUID):
    with Session(engine, expire_on_commit=False, autoflush=False) as session:
        return ReportService(ReportRepository(session), maximum_rows=5_000).build_snapshot(event_id)


def observe_sql(engine, operation) -> tuple[object, int]:
    statements: list[str] = []

    def record_statement(*args) -> None:
        statements.append(str(args[2]))

    sqlalchemy_event.listen(engine, "before_cursor_execute", record_statement)
    try:
        result = operation()
    finally:
        sqlalchemy_event.remove(engine, "before_cursor_execute", record_statement)
    return result, len(statements)


def timing_summary(values: list[float]) -> dict[str, float | int]:
    ordered = sorted(values)
    p95_index = max(0, math.ceil(len(ordered) * 0.95) - 1)
    return {
        "samples": len(values),
        "p50": round(median(ordered), 3),
        "p95": round(ordered[p95_index], 3),
        "max": round(max(ordered), 3),
    }


def measure(operation, *, samples: int) -> dict[str, float | int]:
    durations: list[float] = []
    for _ in range(samples):
        started = perf_counter()
        operation()
        durations.append((perf_counter() - started) * 1_000)
    return timing_summary(durations)


def assert_summary_invariants(summary: dict) -> None:
    totals = summary["totals"]
    assert totals["declared_companions"] == (
        totals["effective_admitted_companions"] + totals["remaining_declared_companion_slots"]
    )
    assert totals["companion_capacity"] == (
        totals["declared_companions"]
        + totals["released_companion_allowance"]
        + totals["uncommitted_companion_allowance"]
    )
    assert totals["effective_total_people_attended"] == (
        totals["named_holders_effectively_attended"] + totals["effective_admitted_companions"]
    )


def assert_private_response(response: httpx.Response, raw: str) -> None:
    body = response.text.casefold()
    assert raw.casefold() not in body
    assert _PRIVATE_PHONE.casefold() not in body
    assert _PRIVATE_EMAIL.casefold() not in body
    assert all(term not in body for term in _FORBIDDEN_RESPONSE_TERMS)


class EventDayApi:
    def __init__(self, client: httpx.AsyncClient, headers: dict[str, str]) -> None:
        self.client = client
        self.headers = headers

    async def inspect_response(self, event_id: UUID, raw: str) -> httpx.Response:
        response = await self.client.post(
            "/api/v1/attendance/inspect",
            headers=self.headers,
            json={"event_id": str(event_id), "credential": raw},
        )
        assert response.status_code == 200
        return response

    async def inspect(self, event_id: UUID, raw: str) -> dict:
        return (await self.inspect_response(event_id, raw)).json()

    async def admit_response(
        self,
        event_id: UUID,
        raw: str,
        *,
        mode: str,
        entering: int,
        declared: int | None = None,
        request_id: UUID | None = None,
    ) -> httpx.Response:
        payload: dict[str, object] = {
            "request_id": str(request_id or uuid4()),
            "event_id": str(event_id),
            "credential": raw,
            "arrival_mode": mode,
            "companions_entering_now": entering,
        }
        if declared is not None:
            payload["declared_companions"] = declared
        response = await self.client.post(
            "/api/v1/attendance/admit",
            headers=self.headers,
            json=payload,
        )
        assert response.status_code == 200
        return response

    async def admit(self, *args, **kwargs) -> dict:
        return (await self.admit_response(*args, **kwargs)).json()

    async def correct(
        self,
        event_id: UUID,
        guest_id: UUID,
        *,
        revision: int,
        decision: str,
        expected_status: str | None = None,
        declared: int | None = None,
        admitted: int | None = None,
    ) -> dict:
        payload: dict[str, object] = {
            "expected_correction_revision": revision,
            "decision": decision,
            "reason_code": "OPERATOR_ERROR",
        }
        if expected_status is not None:
            payload["expected_effective_status"] = expected_status
        if declared is not None:
            payload["declared_companions"] = declared
        if admitted is not None:
            payload["companions_admitted"] = admitted
        response = await self.client.post(
            f"/api/v1/events/{event_id}/attendance/guests/{guest_id}/corrections",
            headers=self.headers,
            json=payload,
        )
        assert response.status_code == 201, response.text
        return response.json()

    async def summary(self, event_id: UUID) -> dict:
        response = await self.client.get(
            f"/api/v1/events/{event_id}/reports/summary",
            headers=self.headers,
        )
        assert response.status_code == 200
        result = response.json()
        assert_summary_invariants(result)
        return result

    async def roster(self, event_id: UUID) -> dict:
        response = await self.client.get(
            f"/api/v1/events/{event_id}/reports/attendance?limit=100",
            headers=self.headers,
        )
        assert response.status_code == 200
        result = response.json()
        assert_summary_invariants(result["summary"])
        return result


async def exercise_event_day_api(
    settings: Settings,
    engine,
    *,
    username: str,
    administrator_id: UUID,
    main_event: Event,
    other_event: Event,
    guests: dict[str, tuple[Guest, str]],
    lifecycle: dict[EventStatus, tuple[Event, Guest, str]],
    regenerated_raw: str,
) -> dict[str, object]:
    app = create_app(
        settings,
        database_runtime=DatabaseRuntime(settings),
        authentication_security=AuthenticationSecurity.from_settings(settings),
    )
    transport = httpx.ASGITransport(app=app)
    async with (
        app.router.lifespan_context(app),
        httpx.AsyncClient(transport=transport, base_url="http://synthetic.test") as client,
    ):
        login = await client.post(
            "/api/v1/auth/login",
            json={"username": username, "password": _ADMIN_PASSWORD},
        )
        assert login.status_code == 200
        api = EventDayApi(client, {"Authorization": f"Bearer {login.json()['access_token']}"})

        happy_guest, happy_raw = guests["happy_two"]
        assert (await api.inspect(main_event.id, happy_raw))["inspection_state"] == (
            "READY_FOR_PRIMARY"
        )
        primary_request = uuid4()
        primary = await api.admit(
            main_event.id,
            happy_raw,
            mode="PRIMARY",
            declared=2,
            entering=0,
            request_id=primary_request,
        )
        primary_replay = await api.admit(
            main_event.id,
            happy_raw,
            mode="PRIMARY",
            declared=2,
            entering=0,
            request_id=primary_request,
        )
        assert primary_replay["scan_attempt_id"] == primary["scan_attempt_id"]
        assert (await api.inspect(main_event.id, happy_raw))["remaining_companion_slots"] == 2
        await api.admit(main_event.id, happy_raw, mode="COMPANION", entering=1)
        assert (await api.inspect(main_event.id, happy_raw))["remaining_companion_slots"] == 1
        await api.admit(main_event.id, happy_raw, mode="COMPANION", entering=1)
        exhausted = await api.inspect(main_event.id, happy_raw)
        assert exhausted["inspection_state"] == "ENTITLEMENT_EXHAUSTED"
        denied = await api.admit(main_event.id, happy_raw, mode="COMPANION", entering=1)
        assert denied["outcome"] == "COMPANION_ENTITLEMENT_EXHAUSTED"

        for key, declared in (("happy_zero_allowed", 0), ("happy_one_allowed", 1)):
            _, raw = guests[key]
            await api.admit(
                main_event.id,
                raw,
                mode="PRIMARY",
                declared=declared,
                entering=0,
            )
            if declared:
                await api.admit(main_event.id, raw, mode="COMPANION", entering=1)
            assert (await api.inspect(main_event.id, raw))["inspection_state"] == (
                "ENTITLEMENT_EXHAUSTED"
            )
        await api.summary(main_event.id)

        zero_guest, zero_raw = guests["declaration_zero"]
        await api.admit(
            main_event.id,
            zero_raw,
            mode="PRIMARY",
            declared=0,
            entering=0,
        )
        zero_inspection = await api.inspect(main_event.id, zero_raw)
        assert (
            zero_inspection["declared_companions"],
            zero_inspection["companions_admitted"],
            zero_inspection["remaining_companion_slots"],
        ) == (0, 0, 0)
        zero_denial = await api.admit(main_event.id, zero_raw, mode="COMPANION", entering=1)
        assert zero_denial["outcome"] == "COMPANION_ENTITLEMENT_EXHAUSTED"

        recovery_guest, recovery_raw = guests["correction_recovery"]
        await api.admit(
            main_event.id,
            recovery_raw,
            mode="PRIMARY",
            declared=0,
            entering=0,
        )
        await api.correct(
            main_event.id,
            recovery_guest.id,
            revision=0,
            decision="SET_COMPANION_STATE",
            declared=2,
            admitted=0,
        )
        assert (await api.inspect(main_event.id, recovery_raw))["remaining_companion_slots"] == 2
        await api.admit(main_event.id, recovery_raw, mode="COMPANION", entering=1)
        recovery_row = {
            item["guest_id"]: item for item in (await api.roster(main_event.id))["items"]
        }[str(recovery_guest.id)]
        assert (
            recovery_row["declared_companions"],
            recovery_row["effective_admitted_companions"],
            recovery_row["remaining_companion_slots"],
        ) == (2, 1, 1)

        raw_guest, raw_credential = guests["raw_effective"]
        await api.admit(
            main_event.id,
            raw_credential,
            mode="PRIMARY",
            declared=2,
            entering=2,
        )
        await api.correct(
            main_event.id,
            raw_guest.id,
            revision=0,
            decision="SET_COMPANION_STATE",
            declared=2,
            admitted=1,
        )
        raw_row = {item["guest_id"]: item for item in (await api.roster(main_event.id))["items"]}[
            str(raw_guest.id)
        ]
        assert raw_row["raw_companions_admitted"] == 2
        assert raw_row["effective_admitted_companions"] == 1
        assert (await api.inspect(main_event.id, raw_credential))["remaining_companion_slots"] == 1
        await api.correct(
            main_event.id,
            raw_guest.id,
            revision=1,
            decision="SET_COMPANION_STATE",
            declared=2,
            admitted=2,
        )

        not_guest, not_raw = guests["not_attended"]
        await api.admit(
            main_event.id,
            not_raw,
            mode="PRIMARY",
            declared=2,
            entering=1,
        )
        await api.correct(
            main_event.id,
            not_guest.id,
            revision=0,
            decision="SET_NOT_ATTENDED",
            expected_status="ATTENDED",
        )
        assert (await api.inspect(main_event.id, not_raw))["inspection_state"] == "ADMISSION_LOCKED"
        locked = await api.admit(main_event.id, not_raw, mode="COMPANION", entering=1)
        assert locked["outcome"] == "PRIMARY_REQUIRED"
        not_row = {item["guest_id"]: item for item in (await api.roster(main_event.id))["items"]}[
            str(not_guest.id)
        ]
        assert not_row["effective_status"] == "NOT_ATTENDED"
        assert not_row["effective_admitted_companions"] == 1
        assert not_row["effective_party_size"] == 1
        await api.correct(
            main_event.id,
            not_guest.id,
            revision=1,
            decision="RESTORE_GATE_DERIVED",
            expected_status="NOT_ATTENDED",
        )
        assert (await api.inspect(main_event.id, not_raw))["remaining_companion_slots"] == 1
        await api.admit(main_event.id, not_raw, mode="COMPANION", entering=1)

        origin_guest, origin_raw = guests["correction_origin"]
        await api.correct(
            main_event.id,
            origin_guest.id,
            revision=0,
            decision="SET_ATTENDED",
            expected_status="NOT_ATTENDED",
            admitted=0,
        )
        await api.correct(
            main_event.id,
            origin_guest.id,
            revision=1,
            decision="SET_COMPANION_STATE",
            declared=2,
            admitted=0,
        )
        assert (await api.inspect(main_event.id, origin_raw))["inspection_state"] == (
            "READY_FOR_COMPANION"
        )
        await api.admit(
            main_event.id,
            origin_raw,
            mode="COMPANION",
            entering=1,
        )
        repeated_origin = await api.admit(
            main_event.id,
            origin_raw,
            mode="PRIMARY",
            declared=2,
            entering=0,
        )
        assert repeated_origin["outcome"] == "ALREADY_CHECKED_IN"
        await api.correct(
            main_event.id,
            origin_guest.id,
            revision=2,
            decision="SET_COMPANION_STATE",
            declared=1,
            admitted=1,
        )

        qr_guest, old_qr_raw = guests["qr_regeneration"]
        await api.admit(
            main_event.id,
            old_qr_raw,
            mode="PRIMARY",
            declared=2,
            entering=1,
        )
        before_qr = {item["guest_id"]: item for item in (await api.roster(main_event.id))["items"]}[
            str(qr_guest.id)
        ]
        regenerated = await client.post(
            f"/api/v1/guests/{qr_guest.id}/qr-code/regenerate",
            headers=api.headers,
        )
        assert regenerated.status_code == 201
        assert (await api.inspect(main_event.id, old_qr_raw))["inspection_state"] == (
            "SUPERSEDED_CREDENTIAL"
        )
        new_qr_state = await api.inspect(main_event.id, regenerated_raw)
        assert (
            new_qr_state["declared_companions"],
            new_qr_state["companions_admitted"],
            new_qr_state["remaining_companion_slots"],
        ) == (2, 1, 1)
        after_qr = {item["guest_id"]: item for item in (await api.roster(main_event.id))["items"]}[
            str(qr_guest.id)
        ]
        assert before_qr == after_qr

        _, stale_primary_raw = guests["stale_primary"]
        assert (await api.inspect(main_event.id, stale_primary_raw))["inspection_state"] == (
            "READY_FOR_PRIMARY"
        )
        await api.admit(
            main_event.id,
            stale_primary_raw,
            mode="PRIMARY",
            declared=1,
            entering=0,
        )
        stale_primary = await api.admit(
            main_event.id,
            stale_primary_raw,
            mode="PRIMARY",
            declared=2,
            entering=0,
        )
        assert stale_primary["outcome"] == "ALREADY_CHECKED_IN"
        assert stale_primary["declared_companions"] == 1

        _, stale_slot_raw = guests["stale_slot"]
        await api.admit(
            main_event.id,
            stale_slot_raw,
            mode="PRIMARY",
            declared=2,
            entering=1,
        )
        assert (await api.inspect(main_event.id, stale_slot_raw))["remaining_companion_slots"] == 1
        await api.admit(main_event.id, stale_slot_raw, mode="COMPANION", entering=1)
        stale_slot = await api.admit(
            main_event.id,
            stale_slot_raw,
            mode="COMPANION",
            entering=1,
        )
        assert stale_slot["outcome"] == "COMPANION_ENTITLEMENT_EXHAUSTED"
        assert stale_slot["companions_admitted"] == 2

        _, first_raw = guests["concurrent_first"]
        first_responses = await asyncio.gather(
            api.admit_response(
                main_event.id,
                first_raw,
                mode="PRIMARY",
                declared=0,
                entering=0,
            ),
            api.admit_response(
                main_event.id,
                first_raw,
                mode="PRIMARY",
                declared=2,
                entering=0,
            ),
        )
        assert sorted(response.json()["outcome"] for response in first_responses) == [
            "ADMITTED",
            "ALREADY_CHECKED_IN",
        ]

        _, concurrent_slot_raw = guests["concurrent_slot"]
        await api.admit(
            main_event.id,
            concurrent_slot_raw,
            mode="PRIMARY",
            declared=2,
            entering=1,
        )
        slot_responses = await asyncio.gather(
            api.admit_response(
                main_event.id,
                concurrent_slot_raw,
                mode="COMPANION",
                entering=1,
            ),
            api.admit_response(
                main_event.id,
                concurrent_slot_raw,
                mode="COMPANION",
                entering=1,
            ),
        )
        assert sorted(response.json()["outcome"] for response in slot_responses) == [
            "ADMITTED",
            "COMPANION_ENTITLEMENT_EXHAUSTED",
        ]

        _, same_request_raw = guests["same_request"]
        await api.admit(
            main_event.id,
            same_request_raw,
            mode="PRIMARY",
            declared=2,
            entering=0,
        )
        shared_request_id = uuid4()
        replay_responses = await asyncio.gather(
            *(
                api.admit_response(
                    main_event.id,
                    same_request_raw,
                    mode="COMPANION",
                    entering=1,
                    request_id=shared_request_id,
                )
                for _ in range(2)
            )
        )
        replay_payloads = [response.json() for response in replay_responses]
        assert [payload["outcome"] for payload in replay_payloads] == ["ADMITTED", "ADMITTED"]
        assert replay_payloads[0]["scan_attempt_id"] == replay_payloads[1]["scan_attempt_id"]

        race_guest, race_raw = guests["correction_race"]
        await api.admit(
            main_event.id,
            race_raw,
            mode="PRIMARY",
            declared=2,
            entering=1,
        )
        race_correction, race_admission = await asyncio.gather(
            api.correct(
                main_event.id,
                race_guest.id,
                revision=0,
                decision="SET_COMPANION_STATE",
                declared=1,
                admitted=1,
            ),
            api.admit(
                main_event.id,
                race_raw,
                mode="COMPANION",
                entering=1,
            ),
        )
        expected_previous = (2, 2) if race_admission["outcome"] == "ADMITTED" else (2, 1)
        assert (
            race_correction["correction"]["previous_declared_companions"],
            race_correction["correction"]["previous_companions_admitted"],
        ) == expected_previous
        assert (
            race_correction["effective_state"]["declared_companions"],
            race_correction["effective_state"]["companions_admitted"],
        ) == (1, 1)

        parallel_tasks = []
        for key in ("parallel_a", "parallel_b"):
            _, raw = guests[key]
            parallel_tasks.append(
                api.admit_response(
                    main_event.id,
                    raw,
                    mode="PRIMARY",
                    declared=1,
                    entering=0,
                )
            )
        parallel = await asyncio.gather(*parallel_tasks)
        assert [response.json()["outcome"] for response in parallel] == ["ADMITTED", "ADMITTED"]
        await api.summary(main_event.id)

        for status, (event, _guest, raw) in lifecycle.items():
            inspection_response = await api.inspect_response(event.id, raw)
            admission_response = await api.admit_response(
                event.id,
                raw,
                mode="PRIMARY",
                declared=0,
                entering=0,
            )
            expected = "READY_FOR_PRIMARY" if status is EventStatus.ACTIVE else "EVENT_NOT_ACTIVE"
            admission_expected = "ADMITTED" if status is EventStatus.ACTIVE else "EVENT_NOT_ACTIVE"
            assert inspection_response.json()["inspection_state"] == expected
            assert admission_response.json()["outcome"] == admission_expected

        invalid_cases = {
            "UNKNOWN_CREDENTIAL": (main_event.id, generate_qr_credential()),
            "WRONG_EVENT": (other_event.id, guests["privacy_active"][1]),
            "INACTIVE_CREDENTIAL": (main_event.id, guests["privacy_inactive"][1]),
            "SUPERSEDED_CREDENTIAL": (main_event.id, guests["privacy_superseded"][1]),
            "INELIGIBLE_GUEST": (main_event.id, guests["privacy_ineligible"][1]),
        }
        for expected, (event_id, raw) in invalid_cases.items():
            inspection_response = await api.inspect_response(event_id, raw)
            admission_response = await api.admit_response(
                event_id,
                raw,
                mode="PRIMARY",
                declared=0,
                entering=0,
            )
            assert inspection_response.json()["inspection_state"] == expected
            assert admission_response.json()["outcome"] == expected
            assert_private_response(inspection_response, raw)
            assert_private_response(admission_response, raw)

        summary_response = await client.get(
            f"/api/v1/events/{main_event.id}/reports/summary",
            headers=api.headers,
        )
        roster_response = await client.get(
            f"/api/v1/events/{main_event.id}/reports/attendance?limit=100",
            headers=api.headers,
        )
        json_response = await client.get(
            f"/api/v1/events/{main_event.id}/reports/export/json",
            headers=api.headers,
        )
        csv_response = await client.get(
            f"/api/v1/events/{main_event.id}/reports/export/csv",
            headers=api.headers,
        )
        xlsx_response = await client.get(
            f"/api/v1/events/{main_event.id}/reports/export/xlsx",
            headers=api.headers,
        )
        assert all(
            response.status_code == 200
            for response in (
                summary_response,
                roster_response,
                json_response,
                csv_response,
                xlsx_response,
            )
        )
        assert_summary_invariants(summary_response.json())
        assert_summary_invariants(roster_response.json()["summary"])
        aggregate_body = summary_response.text.casefold()
        assert "guest_id" not in aggregate_body
        assert "full_name" not in aggregate_body
        assert _PRIVATE_EMAIL.casefold() not in aggregate_body
        assert _PRIVATE_PHONE.casefold() not in aggregate_body

        json_report = json_response.json()
        assert_summary_invariants(json_report["summary"])
        roster_by_id = {item["guest_id"]: item for item in json_report["attendance_roster"]}
        unestablished_guest, _ = guests["unestablished"]
        zero_row = roster_by_id[str(zero_guest.id)]
        unestablished_row = roster_by_id[str(unestablished_guest.id)]
        assert (
            zero_row["declared_companions"],
            zero_row["effective_admitted_companions"],
            zero_row["remaining_companion_slots"],
        ) == (0, 0, 0)
        assert (
            unestablished_row["declared_companions"],
            unestablished_row["remaining_companion_slots"],
        ) == (None, None)
        json_keys = set(json_report["attendance_roster"][0])
        assert {"guest_id", "full_name", "invitation_type", "table_number", "category"} <= json_keys
        assert not json_keys.intersection(_FORBIDDEN_RESPONSE_TERMS)
        assert _PRIVATE_EMAIL not in json_response.text
        assert _PRIVATE_PHONE not in json_response.text
        assert old_qr_raw not in json_response.text
        assert regenerated_raw not in json_response.text

        csv_rows = list(csv.reader(io.StringIO(csv_response.content.decode("utf-8"))))
        csv_header_index = csv_rows.index(["Attendance Roster"]) + 1
        csv_header = csv_rows[csv_header_index]
        assert {"Full Name", "Invitation Type", "Table Number", "Category"} <= set(csv_header)
        assert not any(
            term.casefold() in " ".join(csv_header).casefold() for term in _FORBIDDEN_RESPONSE_TERMS
        )
        csv_by_name = {row[0]: row for row in csv_rows[csv_header_index + 1 :]}
        declared_index = csv_header.index("Declared Companions")
        remaining_index = csv_header.index("Remaining Companion Slots")
        assert csv_by_name[zero_guest.full_name][declared_index] == "0"
        assert csv_by_name[zero_guest.full_name][remaining_index] == "0"
        assert csv_by_name[unestablished_guest.full_name][declared_index] == ""
        assert csv_by_name[unestablished_guest.full_name][remaining_index] == ""

        workbook = load_workbook(io.BytesIO(xlsx_response.content), read_only=True)
        try:
            rows = list(workbook["Attendance Roster"].iter_rows(values_only=True))
            xlsx_header = list(rows[0])
            xlsx_by_name = {str(row[0]): row for row in rows[1:]}
            assert xlsx_header == csv_header
            assert xlsx_by_name[zero_guest.full_name][declared_index] == 0
            assert xlsx_by_name[zero_guest.full_name][remaining_index] == 0
            assert xlsx_by_name[unestablished_guest.full_name][declared_index] is None
            assert xlsx_by_name[unestablished_guest.full_name][remaining_index] is None
            all_values = " ".join(
                str(value) for row in rows for value in row if value is not None
            ).casefold()
            assert _PRIVATE_EMAIL.casefold() not in all_values
            assert _PRIVATE_PHONE.casefold() not in all_values
            assert old_qr_raw.casefold() not in all_values
            assert regenerated_raw.casefold() not in all_values
        finally:
            workbook.close()

        with Session(engine) as session:
            assert (
                session.scalar(
                    select(func.count())
                    .select_from(GuestAttendanceRecord)
                    .where(GuestAttendanceRecord.guest_id == origin_guest.id)
                )
                == 0
            )
            assert (
                session.scalar(
                    select(func.count())
                    .select_from(GuestCompanionEntitlement)
                    .where(GuestCompanionEntitlement.guest_id == origin_guest.id)
                )
                == 1
            )
            for status, (event, guest, _raw) in lifecycle.items():
                if status is EventStatus.ACTIVE:
                    continue
                assert (
                    session.scalar(
                        select(func.count())
                        .select_from(GuestAttendanceRecord)
                        .where(GuestAttendanceRecord.event_id == event.id)
                    )
                    == 0
                )
                assert (
                    session.scalar(
                        select(func.count())
                        .select_from(GuestCompanionEntitlement)
                        .where(GuestCompanionEntitlement.guest_id == guest.id)
                    )
                    == 0
                )

        return {
            "happy_guest_id": happy_guest.id,
            "happy_raw": happy_raw,
            "origin_guest_id": origin_guest.id,
            "origin_raw": origin_raw,
            "regenerated_guest_id": qr_guest.id,
            "regenerated_raw": regenerated_raw,
        }


def test_c27_integrated_companion_event_day_acceptance(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    settings = configure_test_environment(monkeypatch, tmp_path)
    assert DISPOSABLE_URL is not None
    assert disposable_alembic.main(["downgrade", "base"]) == 0
    assert disposable_alembic.main(["upgrade", "head"]) == 0
    assert disposable_alembic.main(["check"]) == 0

    engine = create_engine(DISPOSABLE_URL, pool_pre_ping=True)
    cleanup_required = False
    try:
        username = f"c27-admin-{uuid4().hex[:8]}"
        with Session(engine, expire_on_commit=False) as session:
            administrator = AdministratorProvisioningService(
                AdministratorRepository(session),
                AuthenticationSecurity.from_settings(settings).passwords,
            ).provision(username, _ADMIN_PASSWORD)
            main_event = make_event(administrator.id, "Event Day", status=EventStatus.ACTIVE)
            other_event = make_event(administrator.id, "Other Event", status=EventStatus.ACTIVE)
            load_event = make_event(administrator.id, "Load Event", status=EventStatus.ACTIVE)
            session.add_all((main_event, other_event, load_event))
            session.flush()

            guests: dict[str, tuple[Guest, str]] = {}
            for key, allowed in (
                ("happy_two", 2),
                ("happy_zero_allowed", 0),
                ("happy_one_allowed", 1),
                ("declaration_zero", 2),
                ("correction_recovery", 2),
                ("raw_effective", 2),
                ("not_attended", 2),
                ("correction_origin", 2),
                ("qr_regeneration", 2),
                ("stale_primary", 2),
                ("stale_slot", 2),
                ("concurrent_first", 2),
                ("concurrent_slot", 2),
                ("same_request", 2),
                ("correction_race", 2),
                ("parallel_a", 2),
                ("parallel_b", 2),
                ("privacy_active", 2),
                ("unestablished", 2),
            ):
                guests[key] = add_guest_credential(
                    session,
                    main_event,
                    administrator.id,
                    key,
                    allowed=allowed,
                )
            guests["privacy_inactive"] = add_guest_credential(
                session,
                main_event,
                administrator.id,
                "privacy_inactive",
                credential_status=QrCredentialStatus.INACTIVE,
            )
            guests["privacy_superseded"] = add_guest_credential(
                session,
                main_event,
                administrator.id,
                "privacy_superseded",
                credential_status=QrCredentialStatus.SUPERSEDED,
            )
            guests["privacy_ineligible"] = add_guest_credential(
                session,
                main_event,
                administrator.id,
                "privacy_ineligible",
                guest_status=GuestStatus.INACTIVE,
            )

            lifecycle: dict[EventStatus, tuple[Event, Guest, str]] = {}
            for status in EventStatus:
                event = make_event(administrator.id, f"Lifecycle {status.value}", status=status)
                session.add(event)
                session.flush()
                guest, raw = add_guest_credential(
                    session,
                    event,
                    administrator.id,
                    f"lifecycle_{status.value.casefold()}",
                    allowed=0,
                )
                lifecycle[status] = (event, guest, raw)
            session.commit()
            cleanup_required = True

        regenerated_raw = generate_qr_credential()
        monkeypatch.setattr(qr_service_module, "generate_qr_credential", lambda: regenerated_raw)
        api_evidence = asyncio.run(
            exercise_event_day_api(
                settings,
                engine,
                username=username,
                administrator_id=administrator.id,
                main_event=main_event,
                other_event=other_event,
                guests=guests,
                lifecycle=lifecycle,
                regenerated_raw=regenerated_raw,
            )
        )

        load_credentials: list[tuple[Guest, str]] = []
        with Session(engine, expire_on_commit=False) as session:
            for index in range(160):
                load_credentials.append(
                    add_guest_credential(
                        session,
                        load_event,
                        administrator.id,
                        f"load_{index:03d}",
                    )
                )
            session.commit()

        primary_timings: list[float] = []
        companion_timings: list[float] = []
        for index, (_guest, raw) in enumerate(load_credentials[:64]):
            command = admission_command(
                event_id=load_event.id,
                raw=raw,
                administrator_id=administrator.id,
                mode=GateArrivalMode.PRIMARY,
                declared=2,
                entering=0,
            )
            started = perf_counter()
            result = admit_once(engine, command)
            elapsed = (perf_counter() - started) * 1_000
            assert result.outcome is GateScanOutcome.ADMITTED
            if index < 10:
                primary_timings.append(elapsed)

        companion_commands: list[GateAdmissionCommand] = []
        for _guest, raw in load_credentials[:10]:
            command = admission_command(
                event_id=load_event.id,
                raw=raw,
                administrator_id=administrator.id,
                mode=GateArrivalMode.COMPANION,
                entering=1,
            )
            started = perf_counter()
            result = admit_once(engine, command)
            companion_timings.append((perf_counter() - started) * 1_000)
            assert result.outcome is GateScanOutcome.ADMITTED
            companion_commands.append(command)

        replay_timings: list[float] = []
        for _ in range(10):
            started = perf_counter()
            replayed = admit_once(engine, companion_commands[0])
            replay_timings.append((perf_counter() - started) * 1_000)
            assert replayed.outcome is GateScanOutcome.ADMITTED

        correction_timings: list[float] = []
        for guest, _raw in load_credentials[:8]:
            started = perf_counter()
            correction = correct_once(
                engine,
                event_id=load_event.id,
                guest_id=guest.id,
                administrator_id=administrator.id,
                revision=0,
                decision=AttendanceCorrectionDecision.SET_NOT_ATTENDED,
                expected_status=EffectiveAttendanceStatus.ATTENDED,
            )
            correction_timings.append((perf_counter() - started) * 1_000)
            assert correction.effective_state.primary_attended is False

        query_guest, query_raw = load_credentials[64]
        _, inspect_queries = observe_sql(
            engine,
            lambda: inspect_once(engine, load_event.id, query_raw),
        )
        query_primary_command = admission_command(
            event_id=load_event.id,
            raw=query_raw,
            administrator_id=administrator.id,
            mode=GateArrivalMode.PRIMARY,
            declared=2,
            entering=0,
        )
        _, primary_queries = observe_sql(
            engine,
            lambda: admit_once(engine, query_primary_command),
        )
        query_companion_command = admission_command(
            event_id=load_event.id,
            raw=query_raw,
            administrator_id=administrator.id,
            mode=GateArrivalMode.COMPANION,
            entering=1,
        )
        _, companion_queries = observe_sql(
            engine,
            lambda: admit_once(engine, query_companion_command),
        )
        _, replay_queries = observe_sql(
            engine,
            lambda: admit_once(engine, query_companion_command),
        )
        _, correction_queries = observe_sql(
            engine,
            lambda: correct_once(
                engine,
                event_id=load_event.id,
                guest_id=query_guest.id,
                administrator_id=administrator.id,
                revision=0,
                decision=AttendanceCorrectionDecision.SET_COMPANION_STATE,
                declared=1,
                admitted=1,
            ),
        )

        legacy_guest, legacy_raw = load_credentials[65]
        admit_once(
            engine,
            admission_command(
                event_id=load_event.id,
                raw=legacy_raw,
                administrator_id=administrator.id,
                mode=GateArrivalMode.PRIMARY,
                declared=0,
                entering=0,
            ),
        )
        _, legacy_correction_queries = observe_sql(
            engine,
            lambda: correct_once(
                engine,
                event_id=load_event.id,
                guest_id=legacy_guest.id,
                administrator_id=administrator.id,
                revision=0,
                decision=AttendanceCorrectionDecision.SET_NOT_ATTENDED,
                expected_status=EffectiveAttendanceStatus.ATTENDED,
            ),
        )
        load_report, aggregate_report_queries = observe_sql(
            engine,
            lambda: report_once(engine, load_event.id),
        )
        _, guest_report_queries = observe_sql(
            engine,
            lambda: report_once(engine, load_event.id).attendance_roster,
        )
        assert inspect_queries == 4
        assert primary_queries == 9
        assert companion_queries == 8
        assert replay_queries == 4
        assert correction_queries <= 8
        assert legacy_correction_queries <= 6
        assert aggregate_report_queries == 1
        assert guest_report_queries == 1

        inspection_latency = measure(
            lambda: inspect_once(engine, load_event.id, load_credentials[20][1]),
            samples=10,
        )
        report_latency = measure(lambda: report_once(engine, load_event.id), samples=10)
        guest_export_latency = measure(
            lambda: encode_csv(report_once(engine, load_event.id)),
            samples=10,
        )

        with Session(engine) as session:
            load_counts = {
                "guests": session.scalar(
                    select(func.count()).select_from(Guest).where(Guest.event_id == load_event.id)
                ),
                "entitlements": session.scalar(
                    select(func.count())
                    .select_from(GuestCompanionEntitlement)
                    .where(GuestCompanionEntitlement.event_id == load_event.id)
                ),
                "admissions": session.scalar(
                    select(func.count())
                    .select_from(GateScanAttempt)
                    .where(
                        GateScanAttempt.event_id == load_event.id,
                        GateScanAttempt.outcome == GateScanOutcome.ADMITTED,
                    )
                ),
                "corrections": session.scalar(
                    select(func.count())
                    .select_from(GuestAttendanceCorrection)
                    .where(GuestAttendanceCorrection.event_id == load_event.id)
                ),
                "report_queries": aggregate_report_queries,
            }
        assert load_counts == {
            "guests": 160,
            "entitlements": 66,
            "admissions": 77,
            "corrections": 10,
            "report_queries": 1,
        }
        assert load_report.totals.declared_companions == (
            load_report.totals.effective_admitted_companions
            + load_report.totals.remaining_declared_companion_slots
        )
        assert load_report.totals.companion_capacity == (
            load_report.totals.declared_companions
            + load_report.totals.released_companion_allowance
            + load_report.totals.uncommitted_companion_allowance
        )

        query_counts = {
            "inspect": inspect_queries,
            "primary": primary_queries,
            "companion": companion_queries,
            "replay": replay_queries,
            "correction": correction_queries,
            "legacy_correction": legacy_correction_queries,
            "aggregate_report": aggregate_report_queries,
            "guest_report": guest_report_queries,
        }
        latencies = {
            "inspect": inspection_latency,
            "primary": timing_summary(primary_timings),
            "companion": timing_summary(companion_timings),
            "replay": timing_summary(replay_timings),
            "correction": timing_summary(correction_timings),
            "aggregate_report": report_latency,
            "guest_csv_export": guest_export_latency,
        }

        engine.dispose()
        engine = create_engine(DISPOSABLE_URL, pool_pre_ping=True)
        reconnected_inspection = inspect_once(
            engine,
            main_event.id,
            str(api_evidence["regenerated_raw"]),
        )
        assert reconnected_inspection.declared_companions == 2
        assert reconnected_inspection.companions_admitted == 1
        reconnected_report = report_once(engine, main_event.id)
        assert reconnected_report.totals.named_holders_invited > 0
        with Session(engine) as session:
            assert (
                session.scalar(
                    select(func.count())
                    .select_from(GuestAttendanceCorrection)
                    .where(GuestAttendanceCorrection.guest_id == api_evidence["origin_guest_id"])
                )
                == 3
            )
            origin_entitlement = session.scalar(
                select(GuestCompanionEntitlement).where(
                    GuestCompanionEntitlement.guest_id == api_evidence["origin_guest_id"]
                )
            )
            assert origin_entitlement is not None
            assert origin_entitlement.primary_attendance_id is None
            assert origin_entitlement.establishing_correction_id is not None

        print(
            "C27_ACCEPTANCE_METRICS="
            + json.dumps(
                {
                    "synthetic_event": load_counts,
                    "query_counts": query_counts,
                    "latency_ms": latencies,
                    "durability": "reconnected",
                },
                sort_keys=True,
            )
        )
    finally:
        if cleanup_required:
            with engine.begin() as connection:
                connection.execute(text("TRUNCATE TABLE users CASCADE"))
        engine.dispose()
        assert disposable_alembic.main(["check"]) == 0
