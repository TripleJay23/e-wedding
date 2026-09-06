"""Opt-in PostgreSQL M9 reconciliation, privacy, query, and timing verification."""

from __future__ import annotations

import asyncio
import csv
import io
import math
import os
from datetime import UTC, datetime, timedelta
from pathlib import Path
from statistics import median
from time import perf_counter
from uuid import UUID, uuid4

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine, delete, func, select, text, update
from sqlalchemy import event as sqlalchemy_event
from sqlalchemy.orm import Session

from app.cli import disposable_alembic
from app.core.security import AuthenticationSecurity
from app.core.settings import Settings
from app.db.runtime import DatabaseRuntime
from app.main import create_app
from app.models.administrator import Administrator
from app.models.attendance import GateScanAttempt, GuestAttendanceRecord
from app.models.attendance_corrections import (
    AttendanceCorrectionDecision,
    AttendanceCorrectionReason,
    EffectiveAttendanceStatus,
    GuestAttendanceCorrection,
)
from app.models.event import Event, EventStatus
from app.models.guest import Guest, GuestStatus
from app.models.qr import GuestQrCredential, QrCredentialStatus
from app.qr.credentials import generate_qr_credential, hash_qr_credential
from app.reports.exports import encode_csv, encode_xlsx
from app.repositories.administrators import AdministratorRepository
from app.repositories.attendance import AttendanceRepository
from app.repositories.attendance_corrections import AttendanceCorrectionRepository
from app.repositories.reports import ReportRepository
from app.services.admin_provisioning import AdministratorProvisioningService
from app.services.attendance import AttendanceService
from app.services.attendance_corrections import AttendanceCorrectionService
from app.services.reports import ReportService

DISPOSABLE_URL = os.getenv("WEDDING_DISPOSABLE_DATABASE_URL")
pytestmark = [
    pytest.mark.postgresql,
    pytest.mark.skipif(
        DISPOSABLE_URL is None,
        reason="WEDDING_DISPOSABLE_DATABASE_URL is not explicitly configured",
    ),
]


def configure_test_environment(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Settings:
    assert DISPOSABLE_URL is not None
    private_root = tmp_path / "private"
    monkeypatch.setenv("APP_ENV", "test")
    monkeypatch.setenv("DATABASE_URL", DISPOSABLE_URL)
    monkeypatch.setenv("JWT_SECRET", "synthetic-jwt-signing-secret-at-least-32-bytes")
    monkeypatch.setenv("JWT_ISSUER", "synthetic-issuer")
    monkeypatch.setenv("JWT_AUDIENCE", "synthetic-audience")
    monkeypatch.setenv("PRIVATE_STORAGE_ROOT", str(private_root))
    monkeypatch.setenv("IMPORT_EVIDENCE_STORAGE_ROOT", str(private_root / "import-evidence"))
    monkeypatch.setenv("QR_IMAGE_STORAGE_ROOT", str(private_root / "qr-codes"))
    monkeypatch.setenv("MAX_REPORT_EXPORT_ROWS", "5000")
    return Settings()


def make_event(administrator_id: UUID, name: str) -> Event:
    return Event(
        id=uuid4(),
        event_name=name,
        bride_name="Jordan",
        groom_name="Jessica",
        event_datetime=datetime(2027, 7, 12, 14, 30, tzinfo=UTC),
        venue="Synthetic Hall",
        status=EventStatus.ACTIVE,
        max_companions_per_invitation=2,
        created_by=administrator_id,
    )


def add_guest(
    session: Session,
    event: Event,
    administrator_id: UUID,
    name: str,
    *,
    allowed: int,
    invitation_type: str = "Family",
    category: str | None = None,
    table: str | None = None,
    guest_status: GuestStatus = GuestStatus.ACTIVE,
    credential_status: QrCredentialStatus | None = None,
) -> tuple[Guest, str | None]:
    guest = Guest(
        id=uuid4(),
        event_id=event.id,
        full_name=name,
        phone_number="+255000000000",
        email="private@example.test",
        invitation_type=invitation_type,
        allowed_companions=allowed,
        table_number=table,
        category=category,
        status=guest_status,
    )
    session.add(guest)
    if credential_status is None:
        return guest, None
    raw = generate_qr_credential()
    credential_id = uuid4()
    session.add(
        GuestQrCredential(
            id=credential_id,
            event_id=event.id,
            guest_id=guest.id,
            token_hash=hash_qr_credential(raw),
            image_storage_key=f"{event.id.hex}/{guest.id.hex}/{credential_id.hex}.png",
            image_sha256="a" * 64,
            status=credential_status,
            replaces_credential_id=None,
            issued_by=administrator_id,
            issued_at=datetime.now(UTC),
            deactivated_at=(
                datetime.now(UTC) if credential_status is QrCredentialStatus.INACTIVE else None
            ),
            superseded_at=(
                datetime.now(UTC) if credential_status is QrCredentialStatus.SUPERSEDED else None
            ),
        )
    )
    return guest, raw


def scan(
    session: Session,
    *,
    event_id: UUID,
    raw: str,
    companions: int,
    administrator_id: UUID,
) -> None:
    AttendanceService(AttendanceRepository(session)).scan(
        event_id=event_id,
        presented_credential=raw,
        companions_admitted=companions,
        administrator_id=administrator_id,
    )


def correct(
    session: Session,
    *,
    event_id: UUID,
    guest_id: UUID,
    administrator_id: UUID,
    revision: int,
    expected_status: EffectiveAttendanceStatus,
    decision: AttendanceCorrectionDecision,
    companions: int | None = None,
    declared: int | None = None,
) -> None:
    AttendanceCorrectionService(AttendanceCorrectionRepository(session)).record_correction(
        event_id=event_id,
        guest_id=guest_id,
        expected_correction_revision=revision,
        expected_effective_status=expected_status,
        decision=decision,
        reason_code=AttendanceCorrectionReason.POST_EVENT_RECONCILIATION,
        reason_note="Synthetic private reason that reports must omit",
        companions_admitted=companions,
        administrator_id=administrator_id,
        declared_companions=declared,
    )


def report_once(engine, event_id: UUID):
    with Session(engine, expire_on_commit=False, autoflush=False) as session:
        return ReportService(ReportRepository(session), maximum_rows=5_000).build_snapshot(event_id)


def observe_sql(engine, operation) -> int:
    statements: list[str] = []

    def record_statement(*args) -> None:
        statements.append(str(args[2]))

    sqlalchemy_event.listen(engine, "before_cursor_execute", record_statement)
    try:
        operation()
    finally:
        sqlalchemy_event.remove(engine, "before_cursor_execute", record_statement)
    return len(statements)


def timing_summary(values: list[float]) -> tuple[float, float]:
    ordered = sorted(values)
    p95_index = max(0, math.ceil(len(ordered) * 0.95) - 1)
    return median(ordered), ordered[p95_index]


def persisted_evidence_counts(session: Session) -> tuple[int, int, int]:
    return (
        session.scalar(select(func.count()).select_from(GuestAttendanceRecord)) or 0,
        session.scalar(select(func.count()).select_from(GateScanAttempt)) or 0,
        session.scalar(select(func.count()).select_from(GuestAttendanceCorrection)) or 0,
    )


async def exercise_api(
    settings: Settings,
    engine,
    *,
    username: str,
    administrator_id: UUID,
    event_id: UUID,
    other_event_id: UUID,
) -> dict[str, httpx.Response]:
    security = AuthenticationSecurity.from_settings(settings)
    app = create_app(
        settings,
        database_runtime=DatabaseRuntime(settings),
        authentication_security=security,
    )
    transport = httpx.ASGITransport(app=app)
    async with (
        app.router.lifespan_context(app),
        httpx.AsyncClient(transport=transport, base_url="http://synthetic.test") as client,
    ):
        login = await client.post(
            "/api/v1/auth/login",
            json={"username": username, "password": "synthetic administrator password"},
        )
        assert login.status_code == 200
        headers = {"Authorization": f"Bearer {login.json()['access_token']}"}
        paths = {
            "summary": f"/api/v1/events/{event_id}/reports/summary",
            "roster": f"/api/v1/events/{event_id}/reports/attendance?limit=3&offset=1",
            "filtered": (
                f"/api/v1/events/{event_id}/reports/attendance"
                "?status=ATTENDED&source=MANUAL&category=Family&table_number=1"
                "&invitation_type=Family"
            ),
            "json": f"/api/v1/events/{event_id}/reports/export/json",
            "csv": f"/api/v1/events/{event_id}/reports/export/csv",
            "xlsx": f"/api/v1/events/{event_id}/reports/export/xlsx",
            "other": f"/api/v1/events/{other_event_id}/reports/summary",
            "missing_event": f"/api/v1/events/{uuid4()}/reports/summary",
            "bad_filter": (f"/api/v1/events/{event_id}/reports/attendance?category={'S' * 80}"),
            "bad_format": f"/api/v1/events/{event_id}/reports/export/pdf",
        }
        responses = {name: await client.get(path, headers=headers) for name, path in paths.items()}
        protected_paths = tuple(
            paths[name] for name in ("summary", "roster", "json", "csv", "xlsx")
        )
        for index, path in enumerate(protected_paths):
            responses[f"missing_auth_{index}"] = await client.get(path)
            responses[f"invalid_auth_{index}"] = await client.get(
                path,
                headers={"Authorization": "Bearer invalid-token"},
            )

        expired = security.access_tokens.issue(
            administrator_id,
            now=datetime.now(UTC) - timedelta(days=1),
        ).value
        for index, path in enumerate(protected_paths):
            responses[f"expired_auth_{index}"] = await client.get(
                path,
                headers={"Authorization": f"Bearer {expired}"},
            )
        with Session(engine) as session:
            session.execute(
                update(Administrator)
                .where(Administrator.id == administrator_id)
                .values(is_active=False)
            )
            session.commit()
        for index, path in enumerate(protected_paths):
            responses[f"inactive_auth_{index}"] = await client.get(path, headers=headers)
        with Session(engine) as session:
            session.execute(
                update(Administrator)
                .where(Administrator.id == administrator_id)
                .values(is_active=True)
            )
            session.commit()
        return responses


def test_m9_c26_postgresql_reconciliation_exports_privacy_queries_and_timings(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    settings = configure_test_environment(monkeypatch, tmp_path)
    assert disposable_alembic.main(["upgrade", "head"]) == 0
    assert disposable_alembic.main(["check"]) == 0

    assert DISPOSABLE_URL is not None
    engine = create_engine(DISPOSABLE_URL, pool_pre_ping=True)
    administrator_id: UUID | None = None
    cleanup_required = False
    username = f"m9-admin-{uuid4().hex[:8]}"
    try:
        with Session(engine, expire_on_commit=False) as session:
            administrator = AdministratorProvisioningService(
                AdministratorRepository(session),
                AuthenticationSecurity.from_settings(settings).passwords,
            ).provision(username, "synthetic administrator password")
            administrator_id = administrator.id
            event = make_event(administrator.id, "Synthetic M9 Reconciliation Wedding")
            other_event = make_event(administrator.id, "Synthetic M9 Other Wedding")
            large_event = make_event(administrator.id, "Synthetic M9 Large Wedding")
            session.add_all((event, other_event, large_event))
            session.flush()

            _guest_a, raw_a = add_guest(
                session,
                event,
                administrator.id,
                '=HYPERLINK("https://invalid.example")',
                allowed=2,
                invitation_type="@individual",
                category="+formula",
                table="-1",
                credential_status=QrCredentialStatus.INACTIVE,
            )
            _guest_b, raw_b = add_guest(
                session,
                event,
                administrator.id,
                "Guest B",
                allowed=2,
                category="Family",
                table="1",
                credential_status=QrCredentialStatus.ACTIVE,
            )
            guest_c, raw_c = add_guest(
                session,
                event,
                administrator.id,
                "Guest C",
                allowed=2,
                category="Family",
                table="1",
                credential_status=QrCredentialStatus.ACTIVE,
            )
            _guest_d, raw_d = add_guest(
                session,
                event,
                administrator.id,
                "Guest D",
                allowed=1,
                credential_status=QrCredentialStatus.ACTIVE,
            )
            guest_e, raw_e = add_guest(
                session,
                event,
                administrator.id,
                "Guest E",
                allowed=2,
                category="Family",
                table="1",
                credential_status=QrCredentialStatus.ACTIVE,
            )
            guest_f, raw_f = add_guest(
                session,
                event,
                administrator.id,
                "Guest F",
                allowed=2,
                credential_status=QrCredentialStatus.ACTIVE,
            )
            guest_g, raw_g = add_guest(
                session,
                event,
                administrator.id,
                "Guest G",
                allowed=2,
                credential_status=QrCredentialStatus.ACTIVE,
            )
            _, raw_superseded = add_guest(
                session,
                event,
                administrator.id,
                "Excluded Superseded Guest",
                allowed=0,
                guest_status=GuestStatus.INACTIVE,
                credential_status=QrCredentialStatus.SUPERSEDED,
            )
            _, raw_ineligible = add_guest(
                session,
                event,
                administrator.id,
                "Excluded Inactive Guest",
                allowed=0,
                guest_status=GuestStatus.INACTIVE,
                credential_status=QrCredentialStatus.ACTIVE,
            )
            _, other_raw = add_guest(
                session,
                other_event,
                administrator.id,
                "Other Event Guest",
                allowed=0,
                credential_status=QrCredentialStatus.ACTIVE,
            )
            for index in range(120):
                add_guest(
                    session,
                    large_event,
                    administrator.id,
                    f"Large Synthetic Guest {index:03d}",
                    allowed=index % 3,
                    category=None if index % 2 else "Family",
                    table=None if index % 4 else str(index // 4),
                )
            session.commit()
            cleanup_required = True

            assert all(
                raw is not None
                for raw in (
                    raw_a,
                    raw_b,
                    raw_c,
                    raw_d,
                    raw_e,
                    raw_f,
                    raw_g,
                    raw_superseded,
                    raw_ineligible,
                    other_raw,
                )
            )
            scan(
                session,
                event_id=event.id,
                raw=raw_b or "",
                companions=0,
                administrator_id=administrator.id,
            )
            scan(
                session,
                event_id=event.id,
                raw=raw_c or "",
                companions=1,
                administrator_id=administrator.id,
            )
            scan(
                session,
                event_id=event.id,
                raw=raw_d or "",
                companions=2,
                administrator_id=administrator.id,
            )
            scan(
                session,
                event_id=event.id,
                raw=raw_d or "",
                companions=1,
                administrator_id=administrator.id,
            )
            scan(
                session,
                event_id=event.id,
                raw=raw_f or "",
                companions=2,
                administrator_id=administrator.id,
            )
            scan(
                session,
                event_id=event.id,
                raw=raw_g or "",
                companions=1,
                administrator_id=administrator.id,
            )
            scan(
                session,
                event_id=event.id,
                raw=raw_b or "",
                companions=0,
                administrator_id=administrator.id,
            )
            scan(
                session,
                event_id=event.id,
                raw=generate_qr_credential(),
                companions=0,
                administrator_id=administrator.id,
            )
            scan(
                session,
                event_id=event.id,
                raw=other_raw or "",
                companions=0,
                administrator_id=administrator.id,
            )
            scan(
                session,
                event_id=event.id,
                raw=raw_a or "",
                companions=0,
                administrator_id=administrator.id,
            )
            scan(
                session,
                event_id=event.id,
                raw=raw_superseded or "",
                companions=0,
                administrator_id=administrator.id,
            )
            scan(
                session,
                event_id=event.id,
                raw=raw_ineligible or "",
                companions=0,
                administrator_id=administrator.id,
            )
            correct(
                session,
                event_id=event.id,
                guest_id=guest_c.id,
                administrator_id=administrator.id,
                revision=0,
                expected_status=EffectiveAttendanceStatus.ATTENDED,
                decision=AttendanceCorrectionDecision.SET_COMPANION_STATE,
                companions=1,
                declared=2,
            )
            correct(
                session,
                event_id=event.id,
                guest_id=guest_e.id,
                administrator_id=administrator.id,
                revision=0,
                expected_status=EffectiveAttendanceStatus.NOT_ATTENDED,
                decision=AttendanceCorrectionDecision.SET_ATTENDED,
                companions=1,
            )
            correct(
                session,
                event_id=event.id,
                guest_id=guest_f.id,
                administrator_id=administrator.id,
                revision=0,
                expected_status=EffectiveAttendanceStatus.ATTENDED,
                decision=AttendanceCorrectionDecision.SET_COMPANION_STATE,
                companions=1,
                declared=2,
            )
            correct(
                session,
                event_id=event.id,
                guest_id=guest_g.id,
                administrator_id=administrator.id,
                revision=0,
                expected_status=EffectiveAttendanceStatus.ATTENDED,
                decision=AttendanceCorrectionDecision.SET_NOT_ATTENDED,
            )

        with Session(engine) as session:
            before = persisted_evidence_counts(session)
        responses = asyncio.run(
            exercise_api(
                settings,
                engine,
                username=username,
                administrator_id=administrator_id,
                event_id=event.id,
                other_event_id=other_event.id,
            )
        )
        with Session(engine) as session:
            after = persisted_evidence_counts(session)
        assert before == after

        assert responses["summary"].status_code == 200
        summary = responses["summary"].json()
        aggregate_payload = responses["summary"].text
        assert all(
            forbidden not in aggregate_payload
            for forbidden in ('"guest_id"', '"full_name"', '"table_number"', '"category"')
        )
        assert summary["totals"] == {
            "named_holders_invited": 7,
            "named_holders_effectively_attended": 5,
            "named_holders_effectively_not_attended": 2,
            "companion_capacity": 13,
            "declared_companions": 7,
            "effective_admitted_companions": 5,
            "raw_companions_admitted": 6,
            "remaining_declared_companion_slots": 2,
            "released_companion_allowance": 4,
            "uncommitted_companion_allowance": 2,
            "maximum_total_people_capacity": 20,
            "effective_total_people_attended": 10,
            "remaining_people_capacity": 10,
            "named_holder_attendance_percentage": 71.43,
        }
        assert summary["raw_gate_metrics"] == {
            "real_gate_admissions": 5,
            "duplicate_scan_attempts": 1,
            "unknown_credential_attempts": 1,
            "wrong_event_attempts": 1,
            "inactive_credential_attempts": 1,
            "superseded_credential_attempts": 1,
            "ineligible_guest_attempts": 1,
            "companion_limit_exceeded_attempts": 1,
            "actual_companions_admitted_at_gate": 5,
        }
        assert responses["roster"].status_code == 200
        assert responses["roster"].json()["total"] == 7
        assert len(responses["roster"].json()["items"]) == 3
        assert responses["filtered"].status_code == 200
        assert [row["full_name"] for row in responses["filtered"].json()["items"]] == ["Guest E"]
        assert responses["other"].json()["totals"]["named_holders_invited"] == 1
        assert responses["missing_event"].status_code == 404
        assert responses["bad_filter"].status_code == 422
        assert "S" * 80 not in responses["bad_filter"].text
        assert responses["bad_format"].status_code == 422
        assert all(
            response.status_code == 401
            for name, response in responses.items()
            if name.startswith(("missing_auth_", "invalid_auth_"))
        )
        assert all(
            response.status_code == 401
            for name, response in responses.items()
            if name.startswith(("expired_auth_", "inactive_auth_"))
        )

        for name in ("summary", "roster", "filtered", "json", "csv", "xlsx"):
            response = responses[name]
            assert response.headers["cache-control"] == "private, no-store"
            for forbidden in (
                "private@example.test",
                "+255000000000",
                "reason_note",
                "Synthetic private reason",
                "token_hash",
                "image_storage_key",
                "password_hash",
            ):
                assert forbidden not in response.text

        json_report = responses["json"].json()
        json_roster = {row["full_name"]: row for row in json_report["attendance_roster"]}
        guest_a = json_roster['=HYPERLINK("https://invalid.example")']
        assert guest_a["declared_companions"] is None
        assert guest_a["remaining_companion_slots"] is None
        assert guest_a["released_companion_allowance"] is None
        assert json_roster["Guest B"]["declared_companions"] == 0
        assert json_roster["Guest B"]["effective_admitted_companions"] == 0
        assert json_roster["Guest B"]["remaining_companion_slots"] == 0
        assert json_roster["Guest C"]["remaining_companion_slots"] == 1
        assert json_roster["Guest D"]["remaining_companion_slots"] == 0
        assert json_roster["Guest E"]["raw_gate_attended"] is False
        assert json_roster["Guest E"]["raw_companions_admitted"] == 1
        assert json_roster["Guest F"]["raw_companions_admitted"] == 2
        assert json_roster["Guest F"]["effective_admitted_companions"] == 1
        assert json_roster["Guest G"]["effective_status"] == "NOT_ATTENDED"
        assert json_roster["Guest G"]["effective_admitted_companions"] == 1
        assert json_roster["Guest G"]["effective_party_size"] == 1
        assert all(
            0
            <= row["effective_admitted_companions"]
            <= row["declared_companions"]
            <= row["allowed_companions"]
            for row in json_roster.values()
            if row["declared_companions"] is not None
        )
        csv_rows = list(csv.reader(io.StringIO(responses["csv"].content.decode("utf-8"))))
        workbook = load_workbook(io.BytesIO(responses["xlsx"].content), read_only=True)
        try:
            summary_rows = list(workbook["Summary"].iter_rows(values_only=True))
            roster_rows = list(workbook["Attendance Roster"].iter_rows())
            csv_attended = next(
                int(row[2])
                for row in csv_rows
                if row[:2] == ["totals", "Named Holders Effectively Attended"]
            )
            xlsx_attended = next(
                row[2]
                for row in summary_rows
                if row[:2] == ("totals", "Named Holders Effectively Attended")
            )
            assert (
                json_report["summary"]["totals"]["named_holders_effectively_attended"]
                == csv_attended
                == xlsx_attended
                == 5
            )
            csv_header_index = csv_rows.index(["Attendance Roster"]) + 1
            csv_header = csv_rows[csv_header_index]
            csv_by_name = {row[0]: row for row in csv_rows[csv_header_index + 1 :]}
            xlsx_header = [cell.value for cell in roster_rows[0]]
            xlsx_by_name = {row[0].value: row for row in roster_rows[1:]}
            declared_index = csv_header.index("Declared Companions")
            remaining_index = csv_header.index("Remaining Companion Slots")
            released_index = csv_header.index("Released Companion Allowance")
            guest_a_csv = next(
                row for row in csv_rows[csv_header_index + 1 :] if "HYPERLINK" in row[0]
            )
            guest_a_xlsx = next(row for row in roster_rows[1:] if "HYPERLINK" in str(row[0].value))
            assert csv_header == xlsx_header
            assert guest_a_csv[declared_index] == ""
            assert guest_a_xlsx[declared_index].value is None
            assert csv_by_name["Guest B"][declared_index] == "0"
            assert csv_by_name["Guest B"][remaining_index] == "0"
            assert xlsx_by_name["Guest B"][declared_index].value == 0
            assert xlsx_by_name["Guest B"][remaining_index].value == 0
            assert guest_a_csv[released_index] == ""
            dangerous_csv = guest_a_csv
            dangerous_xlsx = guest_a_xlsx
            assert dangerous_csv[:4] == [
                '\'=HYPERLINK("https://invalid.example")',
                "'@individual",
                "'+formula",
                "'-1",
            ]
            assert all(cell.data_type != "f" for cell in dangerous_xlsx)
            assert workbook.vba_archive is None
            assert workbook._external_links == []
        finally:
            workbook.close()

        query_counts = {
            "small_event": observe_sql(engine, lambda: report_once(engine, event.id)),
            "large_event": observe_sql(engine, lambda: report_once(engine, large_event.id)),
        }
        assert query_counts == {"small_event": 1, "large_event": 1}

        observations: dict[str, tuple[float, float]] = {}
        for name, operation in {
            "summary": lambda: report_once(engine, event.id).totals,
            "roster_100": lambda: report_once(engine, large_event.id).attendance_roster[:100],
            "csv": lambda: encode_csv(report_once(engine, large_event.id)),
            "xlsx": lambda: encode_xlsx(report_once(engine, large_event.id)),
        }.items():
            durations: list[float] = []
            for _ in range(10):
                started = perf_counter()
                operation()
                durations.append((perf_counter() - started) * 1_000)
            observations[name] = timing_summary(durations)

        large_report = report_once(engine, large_event.id)
        csv_size = len(encode_csv(large_report))
        xlsx_size = len(encode_xlsx(large_report))
        print(
            "M9_QUERY_OBSERVATION "
            + " ".join(f"{case}_sql={count}" for case, count in query_counts.items())
        )
        print(
            "M9_TIMING_OBSERVATION "
            + " ".join(
                f"{case}_median_ms={values[0]:.3f} {case}_p95_ms={values[1]:.3f}"
                for case, values in observations.items()
            )
        )
        print(f"M9_SIZE_OBSERVATION csv_bytes={csv_size} xlsx_bytes={xlsx_size}")
    finally:
        if cleanup_required:
            with engine.begin() as connection:
                connection.execute(
                    text(
                        "TRUNCATE TABLE gate_scan_attempts, "
                        "guest_companion_entitlements, guest_attendance_corrections, "
                        "guest_attendance_records CASCADE"
                    )
                )
        engine.dispose()
        if cleanup_required:
            assert disposable_alembic.main(["downgrade", "20260820_0006"]) == 0
            cleanup_engine = create_engine(DISPOSABLE_URL, pool_pre_ping=True)
            try:
                if administrator_id is not None:
                    with Session(cleanup_engine) as session:
                        session.execute(delete(Event).where(Event.created_by == administrator_id))
                        administrator = session.get(Administrator, administrator_id)
                        if administrator is not None:
                            session.delete(administrator)
                        session.commit()
            finally:
                cleanup_engine.dispose()
            assert disposable_alembic.main(["upgrade", "head"]) == 0
