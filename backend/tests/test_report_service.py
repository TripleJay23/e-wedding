"""Canonical M9 calculations, filters, and independent export validation."""

import csv
import io
import json
from datetime import UTC, datetime, timedelta
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from app.models.attendance_corrections import (
    AttendanceCorrectionDecision,
    EffectiveAttendanceSource,
    EffectiveAttendanceStatus,
)
from app.models.event import EventStatus
from app.reports.exports import encode_csv, encode_xlsx
from app.repositories.attendance_corrections import (
    AttendanceEvidence,
    AttendanceStateSnapshot,
    CorrectionHead,
)
from app.repositories.reports import (
    ReportDatasetEvidence,
    ReportEventEvidence,
    ReportGuestEvidence,
    ReportRawGateEvidence,
)
from app.schemas.reports import JsonReportResponse, ReportFiltersResponse, ReportSummaryResponse
from app.services.reports import ReportFilters, ReportService, ReportSizeLimitError

NOW = datetime(2026, 8, 28, 10, 30, tzinfo=UTC)


class FakeReportRepository:
    def __init__(self, dataset: ReportDatasetEvidence | None) -> None:
        self.dataset = dataset
        self.calls: list[tuple[object, int]] = []
        self.ended = 0

    def fetch_event_dataset(self, event_id, *, maximum_rows):
        self.calls.append((event_id, maximum_rows))
        return self.dataset

    def end_read(self) -> None:
        self.ended += 1


def guest_evidence(
    name: str,
    *,
    allowed: int,
    gate_companions: int | None = None,
    correction: AttendanceCorrectionDecision | None = None,
    correction_companions: int | None = None,
    declared: int | None = None,
    admitted: int | None = None,
    raw_companions: int | None = None,
    category: str | None = None,
    table: str | None = None,
    invitation_type: str = "Family",
) -> ReportGuestEvidence:
    guest_id = uuid4()
    attendance = None
    if gate_companions is not None:
        attendance = AttendanceEvidence(
            id=uuid4(),
            companions_admitted=gate_companions,
            checked_in_at=NOW,
        )
    head = None
    if correction is not None:
        resulting_status = (
            EffectiveAttendanceStatus.NOT_ATTENDED
            if correction is AttendanceCorrectionDecision.SET_NOT_ATTENDED
            else EffectiveAttendanceStatus.ATTENDED
        )
        resulting_source = (
            EffectiveAttendanceSource.GATE
            if attendance is not None
            else EffectiveAttendanceSource.MANUAL
        )
        head = CorrectionHead(
            id=uuid4(),
            revision_number=1,
            decision=correction,
            companions_admitted=correction_companions,
            created_at=NOW + timedelta(minutes=5),
            resulting_effective_status=resulting_status,
            resulting_effective_source=resulting_source,
            resulting_declared_companions=(
                declared if correction is AttendanceCorrectionDecision.SET_COMPANION_STATE else None
            ),
            resulting_companions_admitted=(
                admitted if correction is AttendanceCorrectionDecision.SET_COMPANION_STATE else None
            ),
        )
    return ReportGuestEvidence(
        id=guest_id,
        full_name=name,
        invitation_type=invitation_type,
        category=category,
        table_number=table,
        allowed_companions=allowed,
        state_snapshot=AttendanceStateSnapshot(
            guest_id=guest_id,
            attendance=attendance,
            latest_correction=head,
            allowed_companions=allowed,
            declared_companions=declared,
            companions_admitted=admitted,
        ),
        raw_gate_attended=attendance is not None,
        raw_gate_checked_in_at=None if attendance is None else attendance.checked_in_at,
        raw_companions_admitted=(
            raw_companions
            if raw_companions is not None
            else (gate_companions or 0)
            + (
                correction_companions or 0
                if correction is AttendanceCorrectionDecision.SET_ATTENDED
                else 0
            )
        ),
    )


def reconciliation_dataset() -> ReportDatasetEvidence:
    return ReportDatasetEvidence(
        event=ReportEventEvidence(
            id=uuid4(),
            event_name="Synthetic M9 Wedding",
            bride_name="Jordan",
            groom_name="Jessica",
            event_datetime=NOW,
            venue="Synthetic Hall",
            status=EventStatus.COMPLETED,
        ),
        raw_gate=ReportRawGateEvidence(
            real_gate_admissions=4,
            duplicate_scan_attempts=2,
            unknown_credential_attempts=1,
            wrong_event_attempts=1,
            inactive_credential_attempts=1,
            superseded_credential_attempts=1,
            ineligible_guest_attempts=1,
            companion_limit_exceeded_attempts=1,
            actual_companions_admitted_at_gate=4,
        ),
        guests=(
            guest_evidence(
                "Guest A",
                allowed=0,
                gate_companions=0,
                declared=0,
                admitted=0,
                category="Friends",
            ),
            guest_evidence(
                "Guest B",
                allowed=2,
                gate_companions=2,
                declared=2,
                admitted=2,
                category="Family",
                table="1",
            ),
            guest_evidence(
                "Guest C",
                allowed=1,
                correction=AttendanceCorrectionDecision.SET_ATTENDED,
                correction_companions=1,
                declared=1,
                admitted=1,
                category="Family",
                table="1",
            ),
            guest_evidence(
                "Guest D",
                allowed=1,
                gate_companions=1,
                correction=AttendanceCorrectionDecision.SET_NOT_ATTENDED,
                declared=1,
                admitted=1,
                category=None,
            ),
            guest_evidence(
                "Guest E",
                allowed=1,
                gate_companions=1,
                correction=AttendanceCorrectionDecision.RESTORE_GATE_DERIVED,
                declared=1,
                admitted=1,
                category="Friends",
            ),
            guest_evidence("Guest F", allowed=0),
            guest_evidence("Guest G", allowed=2),
        ),
    )


def build_report(*, filters: ReportFilters | None = None):
    dataset = reconciliation_dataset()
    repository = FakeReportRepository(dataset)
    report = ReportService(repository, maximum_rows=5_000).build_snapshot(
        dataset.event.id,
        filters=filters,
    )
    assert repository.calls == [(dataset.event.id, 5_000)]
    assert repository.ended == 1
    return report


def companion_reconciliation_dataset() -> ReportDatasetEvidence:
    return ReportDatasetEvidence(
        event=ReportEventEvidence(
            id=uuid4(),
            event_name="Synthetic C2.6 Wedding",
            bride_name="Amina",
            groom_name="Baraka",
            event_datetime=NOW,
            venue="Synthetic Hall",
            status=EventStatus.ACTIVE,
        ),
        raw_gate=ReportRawGateEvidence(
            real_gate_admissions=6,
            duplicate_scan_attempts=0,
            unknown_credential_attempts=0,
            wrong_event_attempts=0,
            inactive_credential_attempts=0,
            superseded_credential_attempts=0,
            ineligible_guest_attempts=0,
            companion_limit_exceeded_attempts=0,
            actual_companions_admitted_at_gate=5,
        ),
        guests=(
            guest_evidence("Guest A", allowed=2),
            guest_evidence(
                "Guest B",
                allowed=2,
                gate_companions=0,
                declared=0,
                admitted=0,
            ),
            guest_evidence(
                "Guest C",
                allowed=2,
                gate_companions=1,
                declared=2,
                admitted=1,
            ),
            guest_evidence(
                "Guest D",
                allowed=1,
                gate_companions=1,
                declared=1,
                admitted=1,
            ),
            guest_evidence(
                "Guest E",
                allowed=2,
                correction=AttendanceCorrectionDecision.SET_ATTENDED,
                correction_companions=1,
                declared=1,
                admitted=1,
            ),
            guest_evidence(
                "Guest F",
                allowed=2,
                gate_companions=2,
                correction=AttendanceCorrectionDecision.SET_COMPANION_STATE,
                declared=2,
                admitted=1,
                raw_companions=2,
            ),
            guest_evidence(
                "Guest G",
                allowed=2,
                gate_companions=1,
                correction=AttendanceCorrectionDecision.SET_NOT_ATTENDED,
                declared=1,
                admitted=1,
            ),
        ),
    )


def test_c26_companion_metrics_nullability_and_reconciliation() -> None:
    dataset = companion_reconciliation_dataset()
    report = ReportService(FakeReportRepository(dataset), maximum_rows=5_000).build_snapshot(
        dataset.event.id
    )
    totals = report.totals

    assert totals.named_holders_invited == 7
    assert totals.named_holders_effectively_attended == 5
    assert totals.named_holders_effectively_not_attended == 2
    assert totals.companion_capacity == 13
    assert totals.declared_companions == 7
    assert totals.effective_admitted_companions == 5
    assert totals.raw_companions_admitted == 6
    assert totals.remaining_declared_companion_slots == 2
    assert totals.released_companion_allowance == 4
    assert totals.uncommitted_companion_allowance == 2
    assert totals.effective_total_people_attended == 10
    assert totals.remaining_people_capacity == 10
    assert totals.declared_companions == (
        totals.effective_admitted_companions + totals.remaining_declared_companion_slots
    )
    assert totals.companion_capacity == (
        totals.declared_companions
        + totals.released_companion_allowance
        + totals.uncommitted_companion_allowance
    )
    assert totals.effective_total_people_attended == (
        totals.named_holders_effectively_attended + totals.effective_admitted_companions
    )

    rows = {row.full_name: row for row in report.attendance_roster}
    assert rows["Guest A"].declared_companions is None
    assert rows["Guest A"].effective_admitted_companions == 0
    assert rows["Guest A"].remaining_companion_slots is None
    assert rows["Guest A"].released_companion_allowance is None
    assert rows["Guest B"].declared_companions == 0
    assert rows["Guest B"].effective_admitted_companions == 0
    assert rows["Guest B"].remaining_companion_slots == 0
    assert rows["Guest B"].released_companion_allowance == 2
    assert rows["Guest C"].remaining_companion_slots == 1
    assert rows["Guest D"].remaining_companion_slots == 0
    assert rows["Guest E"].raw_gate_attended is False
    assert rows["Guest E"].raw_companions_admitted == 1
    assert rows["Guest F"].raw_companions_admitted == 2
    assert rows["Guest F"].effective_admitted_companions == 1
    assert rows["Guest G"].effective_status is EffectiveAttendanceStatus.NOT_ATTENDED
    assert rows["Guest G"].effective_admitted_companions == 1
    assert rows["Guest G"].effective_party_size == 1

    for row in report.attendance_roster:
        if row.declared_companions is not None:
            assert 0 <= row.effective_admitted_companions <= row.declared_companions
            assert row.declared_companions <= row.allowed_companions


def test_hand_calculated_raw_and_effective_truth_reconcile() -> None:
    report = build_report()

    assert report.totals.named_holders_invited == 7
    assert report.totals.named_holders_effectively_attended == 4
    assert report.totals.named_holders_effectively_not_attended == 3
    assert report.totals.companion_capacity == 7
    assert report.totals.effective_admitted_companions == 5
    assert report.totals.declared_companions == 5
    assert report.totals.raw_companions_admitted == 5
    assert report.totals.remaining_declared_companion_slots == 0
    assert report.totals.released_companion_allowance == 0
    assert report.totals.uncommitted_companion_allowance == 2
    assert report.totals.maximum_total_people_capacity == 14
    assert report.totals.effective_total_people_attended == 9
    assert report.totals.remaining_people_capacity == 5
    assert report.totals.named_holder_attendance_percentage == 57.14
    assert report.raw_gate_metrics.real_gate_admissions == 4
    assert report.raw_gate_metrics.actual_companions_admitted_at_gate == 4

    rows = {row.full_name: row for row in report.attendance_roster}
    assert rows["Guest C"].effective_source is EffectiveAttendanceSource.MANUAL
    assert rows["Guest D"].raw_gate_attended is True
    assert rows["Guest D"].effective_status is EffectiveAttendanceStatus.NOT_ATTENDED
    assert rows["Guest E"].effective_source is EffectiveAttendanceSource.GATE
    assert rows["Guest F"].effective_source is EffectiveAttendanceSource.NONE


def test_c21_remaining_people_capacity_is_not_remaining_locked_entitlement() -> None:
    report = build_report()
    totals = report.totals

    assert totals.named_holders_invited == 7
    assert totals.companion_capacity == 7
    assert totals.maximum_total_people_capacity == 14
    assert totals.effective_admitted_companions == 5
    assert totals.effective_total_people_attended == 9
    assert totals.remaining_people_capacity == (
        totals.maximum_total_people_capacity - totals.effective_total_people_attended
    )
    assert totals.remaining_people_capacity == 5

    attended_rows = [
        row
        for row in report.attendance_roster
        if row.effective_status is EffectiveAttendanceStatus.ATTENDED
    ]
    invitation_allowance_minus_admitted = sum(
        row.allowed_companions - row.effective_admitted_companions for row in attended_rows
    )
    assert invitation_allowance_minus_admitted == 0
    assert totals.remaining_people_capacity != invitation_allowance_minus_admitted


@pytest.mark.parametrize(
    "future_field",
    [
        pytest.param("declared_companions", id="declared_companion_expectation"),
        pytest.param(
            "remaining_declared_companion_slots",
            id="remaining_locked_companion_slots",
        ),
        pytest.param("released_companion_allowance", id="released_companion_allowance"),
    ],
)
def test_c21_target_report_exposes_entitlement_metric(future_field: str) -> None:
    assert hasattr(build_report().totals, future_field)


def test_filters_are_exact_null_safe_and_do_not_change_event_totals() -> None:
    report = build_report(
        filters=ReportFilters(
            effective_status=EffectiveAttendanceStatus.ATTENDED,
            category="Family",
            table_number="1",
        )
    )

    assert [row.full_name for row in report.attendance_roster] == ["Guest B", "Guest C"]
    assert report.totals.named_holders_invited == 7

    none_category = build_report(
        filters=ReportFilters(effective_source=EffectiveAttendanceSource.NONE)
    )
    assert [row.full_name for row in none_category.attendance_roster] == ["Guest F", "Guest G"]


def test_zero_guest_percentage_and_counts_are_safe() -> None:
    dataset = reconciliation_dataset()
    dataset = ReportDatasetEvidence(dataset.event, dataset.raw_gate, ())
    report = ReportService(FakeReportRepository(dataset), maximum_rows=1).build_snapshot(
        dataset.event.id
    )
    assert report.totals.named_holders_invited == 0
    assert report.totals.named_holder_attendance_percentage == 0.0
    assert report.totals.maximum_total_people_capacity == 0


def test_report_row_limit_is_fail_closed() -> None:
    dataset = reconciliation_dataset()
    repository = FakeReportRepository(dataset)
    service = ReportService(repository, maximum_rows=3)
    with pytest.raises(
        ReportSizeLimitError,
        match=r"^report exceeds the configured row limit$",
    ):
        service.build_snapshot(dataset.event.id)


def test_json_csv_xlsx_share_values_and_neutralize_formulas() -> None:
    report = build_report()
    dangerous = guest_evidence(
        '=HYPERLINK("https://invalid.example")',
        allowed=2,
        category="+SUM(1,1)",
        table="-1\nnext",
        invitation_type="@command",
    )
    dataset = reconciliation_dataset()
    dataset = ReportDatasetEvidence(dataset.event, dataset.raw_gate, (*dataset.guests, dangerous))
    report = ReportService(FakeReportRepository(dataset), maximum_rows=5_000).build_snapshot(
        dataset.event.id
    )

    json_model = JsonReportResponse(
        summary=ReportSummaryResponse.model_validate(report),
        attendance_roster=list(report.attendance_roster),
        filters=ReportFiltersResponse.model_validate(report.filters),
    )
    json_data = json.loads(json_model.model_dump_json())
    csv_rows = list(csv.reader(io.StringIO(encode_csv(report).decode("utf-8"))))
    workbook = load_workbook(io.BytesIO(encode_xlsx(report)), read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ["Summary", "Attendance Roster"]
        assert workbook.vba_archive is None
        assert workbook._external_links == []
        summary_rows = list(workbook["Summary"].iter_rows(values_only=True))
        roster_rows = list(workbook["Attendance Roster"].iter_rows())

        json_attended = json_data["summary"]["totals"]["named_holders_effectively_attended"]
        csv_attended = next(
            int(row[2])
            for row in csv_rows
            if row[:2] == ["totals", "Named Holders Effectively Attended"]
        )
        xlsx_attended = next(
            row[2]
            for row in summary_rows
            if row[:2] == ("totals", "Named Holders Effectively Attended")
        )
        assert json_attended == csv_attended == xlsx_attended == 4

        csv_header_index = csv_rows.index(["Attendance Roster"]) + 1
        csv_header = csv_rows[csv_header_index]
        csv_dangerous = next(
            row for row in csv_rows[csv_header_index + 1 :] if "HYPERLINK" in row[0]
        )
        assert "Guest ID" not in csv_header
        assert csv_dangerous[0].startswith("'=")
        assert csv_dangerous[1].startswith("'@")
        assert csv_dangerous[2].startswith("'+")
        assert csv_dangerous[3].startswith("'-")
        assert "\n" not in csv_dangerous[3]

        xlsx_header = [cell.value for cell in roster_rows[0]]
        xlsx_dangerous = next(row for row in roster_rows[1:] if "HYPERLINK" in str(row[0].value))
        assert "Guest ID" not in xlsx_header
        assert all(cell.data_type != "f" for cell in xlsx_dangerous)
        assert xlsx_dangerous[0].value.startswith("'=")
        assert isinstance(roster_rows[1][4].value, int)
        assert isinstance(roster_rows[1][14].value, bool)
    finally:
        workbook.close()

    serialized = json_model.model_dump_json()
    for forbidden in (
        "phone_number",
        "email",
        "reason_note",
        "token_hash",
        "storage_key",
        "password_hash",
    ):
        assert forbidden not in serialized


def test_c26_csv_xlsx_preserve_companion_null_zero_and_operational_privacy() -> None:
    dataset = companion_reconciliation_dataset()
    report = ReportService(FakeReportRepository(dataset), maximum_rows=5_000).build_snapshot(
        dataset.event.id
    )
    csv_rows = list(csv.reader(io.StringIO(encode_csv(report).decode("utf-8"))))
    workbook = load_workbook(io.BytesIO(encode_xlsx(report)), read_only=True, data_only=False)
    try:
        csv_header_index = csv_rows.index(["Attendance Roster"]) + 1
        csv_header = csv_rows[csv_header_index]
        csv_by_name = {row[0]: row for row in csv_rows[csv_header_index + 1 :]}
        roster_rows = list(workbook["Attendance Roster"].iter_rows(values_only=True))
        xlsx_header = list(roster_rows[0])
        xlsx_by_name = {row[0]: row for row in roster_rows[1:]}

        assert csv_header == xlsx_header
        declared_index = csv_header.index("Declared Companions")
        admitted_index = csv_header.index("Effective Admitted Companions")
        raw_index = csv_header.index("Raw Companions Admitted")
        remaining_index = csv_header.index("Remaining Companion Slots")
        released_index = csv_header.index("Released Companion Allowance")
        assert csv_by_name["Guest A"][declared_index] == ""
        assert csv_by_name["Guest A"][remaining_index] == ""
        assert csv_by_name["Guest A"][released_index] == ""
        assert xlsx_by_name["Guest A"][declared_index] is None
        assert xlsx_by_name["Guest A"][remaining_index] is None
        assert xlsx_by_name["Guest A"][released_index] is None
        assert csv_by_name["Guest B"][declared_index] == "0"
        assert csv_by_name["Guest B"][admitted_index] == "0"
        assert csv_by_name["Guest B"][remaining_index] == "0"
        assert xlsx_by_name["Guest B"][declared_index] == 0
        assert xlsx_by_name["Guest B"][admitted_index] == 0
        assert xlsx_by_name["Guest B"][remaining_index] == 0
        assert xlsx_by_name["Guest F"][raw_index] == 2
        assert xlsx_by_name["Guest F"][admitted_index] == 1

        for forbidden in (
            "Phone Number",
            "Email",
            "QR",
            "Token",
            "Hash",
            "Correction Reason",
            "Correction Actor",
            "Entitlement ID",
            "Provenance",
            "Storage",
        ):
            assert forbidden not in csv_header
            assert forbidden not in xlsx_header
    finally:
        workbook.close()
