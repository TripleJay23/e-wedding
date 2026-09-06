"""Authenticated M9 API, privacy headers, pagination, and safe errors."""

import asyncio
import io
from datetime import UTC, datetime
from typing import Any
from uuid import uuid4

import httpx
import pytest
from openpyxl import load_workbook

from app.api.dependencies import (
    get_authentication_service,
    get_current_administrator,
    get_report_service,
)
from app.main import create_app
from app.models.administrator import Administrator, AdministratorRole
from app.models.attendance_corrections import (
    EffectiveAttendanceSource,
    EffectiveAttendanceStatus,
)
from app.models.event import EventStatus
from app.services.authentication import AuthenticationRejectedError
from app.services.reports import (
    AttendanceRosterRow,
    RawGateMetrics,
    ReportEvent,
    ReportFilters,
    ReportPersistenceError,
    ReportResourceNotFoundError,
    ReportSizeLimitError,
    ReportSnapshot,
    ReportTotals,
)

NOW = datetime(2026, 8, 28, 11, 0, tzinfo=UTC)


class FakeDatabaseRuntime:
    def close(self) -> None:
        pass


class RejectAuthenticationService:
    def current_administrator(self, token: str) -> Administrator:
        raise AuthenticationRejectedError("synthetic rejection")


class FakeReportService:
    def __init__(self) -> None:
        self.event_id = uuid4()
        self.calls: list[tuple[object, ReportFilters | None]] = []
        self.next_error: Exception | None = None

    def build_snapshot(self, event_id, *, filters=None) -> ReportSnapshot:
        if self.next_error is not None:
            error = self.next_error
            self.next_error = None
            raise error
        self.calls.append((event_id, filters))
        rows = tuple(
            AttendanceRosterRow(
                guest_id=uuid4(),
                full_name=name,
                invitation_type="Family",
                category="Friends",
                table_number="4",
                allowed_companions=2,
                declared_companions=2,
                effective_status=EffectiveAttendanceStatus.ATTENDED,
                effective_source=EffectiveAttendanceSource.GATE,
                effective_admitted_companions=1,
                raw_companions_admitted=1,
                remaining_companion_slots=1,
                released_companion_allowance=0,
                effective_party_size=2,
                effective_at=NOW,
                raw_gate_attended=True,
                raw_gate_checked_in_at=NOW,
                correction_revision=0,
                has_correction=False,
            )
            for name in ("Amina O'Neil", "Zoë Synthetic")
        )
        selected = filters or ReportFilters()
        return ReportSnapshot(
            event=ReportEvent(
                id=self.event_id,
                event_name="Synthetic M9 Wedding",
                bride_name="Jordan",
                groom_name="Jessica",
                event_datetime=NOW,
                venue="Synthetic Hall",
                status=EventStatus.COMPLETED,
            ),
            totals=ReportTotals(
                named_holders_invited=2,
                named_holders_effectively_attended=2,
                named_holders_effectively_not_attended=0,
                companion_capacity=4,
                effective_admitted_companions=2,
                declared_companions=4,
                raw_companions_admitted=2,
                remaining_declared_companion_slots=2,
                released_companion_allowance=0,
                uncommitted_companion_allowance=0,
                maximum_total_people_capacity=6,
                effective_total_people_attended=4,
                remaining_people_capacity=2,
                named_holder_attendance_percentage=100.0,
            ),
            raw_gate_metrics=RawGateMetrics(
                real_gate_admissions=2,
                duplicate_scan_attempts=1,
                unknown_credential_attempts=0,
                wrong_event_attempts=0,
                inactive_credential_attempts=0,
                superseded_credential_attempts=0,
                ineligible_guest_attempts=0,
                companion_limit_exceeded_attempts=0,
                actual_companions_admitted_at_gate=2,
            ),
            attendance_roster=rows,
            filters=selected,
        )


def make_administrator() -> Administrator:
    return Administrator(
        id=uuid4(),
        username="admin",
        password_hash="$argon2id$synthetic",
        role=AdministratorRole.SUPER_ADMIN,
        is_active=True,
        created_at=NOW,
        updated_at=NOW,
    )


def build_app(settings, *, authenticated: bool = True):
    service = FakeReportService()
    app = create_app(settings, database_runtime=FakeDatabaseRuntime())  # type: ignore[arg-type]
    app.dependency_overrides[get_report_service] = lambda: service
    if authenticated:
        app.dependency_overrides[get_current_administrator] = make_administrator
    else:
        app.dependency_overrides[get_authentication_service] = lambda: RejectAuthenticationService()
    return app, service


async def request(app: object, method: str, path: str, **kwargs: Any) -> httpx.Response:
    transport = httpx.ASGITransport(app=app)  # type: ignore[arg-type]
    async with (
        app.router.lifespan_context(app),  # type: ignore[attr-defined]
        httpx.AsyncClient(transport=transport, base_url="http://synthetic.test") as client,
    ):
        return await client.request(method, path, **kwargs)


def test_summary_and_roster_are_private_paginated_and_privacy_bounded(
    development_settings,
) -> None:
    app, service = build_app(development_settings)
    summary = asyncio.run(request(app, "GET", f"/api/v1/events/{service.event_id}/reports/summary"))
    roster = asyncio.run(
        request(
            app,
            "GET",
            f"/api/v1/events/{service.event_id}/reports/attendance"
            "?status=ATTENDED&source=GATE&category=Friends&table_number=4"
            "&invitation_type=Family&limit=1&offset=1",
        )
    )

    assert summary.status_code == roster.status_code == 200
    assert summary.headers["cache-control"] == "private, no-store"
    assert roster.headers["x-content-type-options"] == "nosniff"
    assert summary.json()["totals"]["effective_total_people_attended"] == 4
    assert roster.json()["total"] == 2
    assert roster.json()["limit"] == 1
    assert roster.json()["offset"] == 1
    assert roster.json()["items"][0]["full_name"] == "Zoë Synthetic"
    selected = service.calls[1][1]
    assert selected == ReportFilters(
        effective_status=EffectiveAttendanceStatus.ATTENDED,
        effective_source=EffectiveAttendanceSource.GATE,
        category="Friends",
        table_number="4",
        invitation_type="Family",
    )
    for response in (summary, roster):
        for forbidden in (
            "phone_number",
            "email",
            "reason_note",
            "token_hash",
            "storage_key",
            "password_hash",
        ):
            assert forbidden not in response.text


@pytest.mark.parametrize(
    ("output_format", "content_type"),
    [
        ("json", "application/json"),
        ("csv", "text/csv"),
        ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ],
)
def test_all_exports_are_real_protected_attachments(
    development_settings,
    output_format,
    content_type,
) -> None:
    app, service = build_app(development_settings)
    response = asyncio.run(
        request(
            app,
            "GET",
            f"/api/v1/events/{service.event_id}/reports/export/{output_format}",
        )
    )
    assert response.status_code == 200
    assert response.headers["cache-control"] == "private, no-store"
    assert response.headers["content-type"].startswith(content_type)
    assert response.headers["content-disposition"].endswith(f'.{output_format}"')
    if output_format == "json":
        assert response.json()["summary"]["totals"]["named_holders_invited"] == 2
    elif output_format == "csv":
        assert response.content.decode("utf-8").startswith("Section,Metric,Value")
    else:
        workbook = load_workbook(io.BytesIO(response.content), read_only=True)
        try:
            assert workbook.sheetnames == ["Summary", "Attendance Roster"]
        finally:
            workbook.close()


@pytest.mark.parametrize("authorization", [None, "Bearer invalid-token"])
def test_report_routes_require_current_persisted_administrator(
    development_settings,
    authorization,
) -> None:
    app, service = build_app(development_settings, authenticated=False)
    headers = {} if authorization is None else {"Authorization": authorization}
    response = asyncio.run(
        request(
            app,
            "GET",
            f"/api/v1/events/{service.event_id}/reports/summary",
            headers=headers,
        )
    )
    assert response.status_code == 401
    assert service.calls == []


@pytest.mark.parametrize(
    ("error", "status_code", "detail"),
    [
        (
            ReportResourceNotFoundError("report resource not found"),
            404,
            "report resource not found",
        ),
        (
            ReportSizeLimitError("report exceeds the configured row limit"),
            413,
            "report exceeds the configured row limit",
        ),
        (ReportPersistenceError("postgresql://private"), 503, "report service is unavailable"),
    ],
)
def test_report_domain_errors_are_bounded(
    development_settings,
    error,
    status_code,
    detail,
) -> None:
    app, service = build_app(development_settings)
    service.next_error = error
    response = asyncio.run(
        request(app, "GET", f"/api/v1/events/{service.event_id}/reports/summary")
    )
    assert response.status_code == status_code
    assert response.json() == {"detail": detail}
    assert "postgresql" not in response.text


def test_invalid_format_pagination_and_filter_use_safe_422_details(development_settings) -> None:
    app, service = build_app(development_settings)
    secret = "S" * 80
    paths = (
        f"/api/v1/events/{service.event_id}/reports/export/pdf",
        f"/api/v1/events/{service.event_id}/reports/attendance?limit=101",
        f"/api/v1/events/{service.event_id}/reports/attendance?category={secret}",
    )
    responses = [asyncio.run(request(app, "GET", path)) for path in paths]
    assert all(response.status_code == 422 for response in responses)
    assert all(secret not in response.text for response in responses)
    assert service.calls == []
