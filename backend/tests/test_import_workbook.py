"""Synthetic-only M4 XLSX security, structure, normalization, and validation tests."""

from __future__ import annotations

from datetime import date, time
from io import BytesIO
from math import inf, nan
from time import perf_counter
from types import SimpleNamespace
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook

from app.imports.workbook import (
    MAX_CANDIDATE_PROBE_ROWS,
    MAX_HEADER_SCAN_ROWS,
    WorkbookParser,
    WorkbookRejectedError,
    _normalize_cell,
)
from app.models.event import DEFAULT_GUEST_FIELD_CONFIGURATION, GuestFieldKey
from app.models.guest import Guest
from app.services.guest_validation import GuestFieldPolicy


def default_policies() -> list[GuestFieldPolicy]:
    return [
        GuestFieldPolicy(field_key, enabled, required)
        for field_key, enabled, required in DEFAULT_GUEST_FIELD_CONFIGURATION
    ]


def save_workbook(workbook: Workbook) -> bytes:
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def simple_workbook(rows: list[list[object]]) -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    for row in rows:
        worksheet.append(row)
    return save_workbook(workbook)


def parse(content: bytes, *, policies: list[GuestFieldPolicy] | None = None):
    return WorkbookParser(maximum_candidate_rows=5_000).parse(
        content,
        policies or default_policies(),
        2,
    )


def error_codes(parsed_row) -> set[str]:
    return {item["code"] for item in parsed_row.validation_errors}


def persistence_text_limit(field_key: GuestFieldKey) -> int:
    length = getattr(Guest.__table__.columns[field_key.value].type, "length", None)
    assert isinstance(length, int)
    return length


def contact_value_at_length(field_key: GuestFieldKey, length: int) -> str:
    if field_key is GuestFieldKey.PHONE_NUMBER:
        fixed_digits = "+1234567"
        assert length >= len(fixed_digits)
        return "+123" + (" " * (length - len(fixed_digits))) + "4567"
    if field_key is GuestFieldKey.EMAIL:
        suffix = "@example.test"
        assert length > len(suffix)
        return ("a" * (length - len(suffix))) + suffix
    raise AssertionError(f"unsupported contact field: {field_key}")


def workbook_with_excel_error(field_key: GuestFieldKey, error_value: str = "#N/A") -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(
        [
            "Full Name",
            "Phone Number",
            "Email",
            "Invitation Type",
            "Allowed Companions",
            "Table Number",
            "Category",
        ]
    )
    worksheet.append(
        [
            "Guest Alpha",
            "+255700000001",
            "alpha@example.test",
            "Family",
            1,
            "A1",
            "Friends",
        ]
    )
    field_columns = {
        GuestFieldKey.FULL_NAME: 1,
        GuestFieldKey.PHONE_NUMBER: 2,
        GuestFieldKey.EMAIL: 3,
        GuestFieldKey.INVITATION_TYPE: 4,
        GuestFieldKey.ALLOWED_COMPANIONS: 5,
        GuestFieldKey.TABLE_NUMBER: 6,
        GuestFieldKey.CATEGORY: 7,
    }
    error_cell = worksheet.cell(row=2, column=field_columns[field_key], value=error_value)
    assert error_cell.data_type == "e"
    return save_workbook(workbook)


def test_reordered_extra_columns_and_human_formatting_are_understood() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Guest List"
    worksheet.merge_cells("A1:H1")
    worksheet["A1"] = "Synthetic Wedding Guest Register"
    worksheet.merge_cells("A2:H2")
    worksheet["A2"] = "Please complete configured fields"
    worksheet.append([])
    worksheet.append(
        [
            "Unmanaged Notes",
            "Invitation Type",
            "Email Address",
            "Guest Name",
            "Table No.",
            "Phone",
            "Category",
            "Plus Ones",
        ]
    )
    worksheet.append(
        [
            "ignored",
            "Family",
            "SYNTHETIC.ONE@EXAMPLE.TEST",
            "Synthetic Guest One",
            "T-1",
            "+255 700 000 001",
            "Friends",
            2,
        ]
    )
    worksheet.append([])
    worksheet.merge_cells("A7:H7")
    worksheet["A7"] = "Synthetic Family Section"
    worksheet.append(
        [
            "Unmanaged Notes",
            "Invitation Type",
            "Email Address",
            "Guest Name",
            "Table No.",
            "Phone",
            "Category",
            "Plus Ones",
        ]
    )
    worksheet.append(
        [
            None,
            "Individual",
            None,
            "Synthetic Guest Two",
            None,
            "+255700000002",
            None,
            0,
        ]
    )
    worksheet.append(["Total Guests", None, None, None, None, None, None, "=COUNTA(D5:D9)"])

    result = parse(save_workbook(workbook))

    assert result.sheet_name == "Guest List"
    assert result.header_row_number == 4
    assert result.column_mapping == {
        "invitation_type": 2,
        "email": 3,
        "full_name": 4,
        "table_number": 5,
        "phone_number": 6,
        "category": 7,
        "allowed_companions": 8,
    }
    assert [row.source_row_number for row in result.rows] == [5, 9]
    assert result.rows[0].values["email"] == "synthetic.one@example.test"
    assert result.rows[0].validation_errors == []
    assert result.rows[1].validation_errors == []
    assert {item["reason"] for item in result.ignored_rows} >= {
        "PRESENTATION",
        "BLANK",
        "REPEATED_HEADER",
        "FOOTER",
    }


def test_unknown_only_row_is_preserved_as_invalid_not_discarded() -> None:
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type", "Unmanaged Notes"],
                ["Synthetic Valid", "Family", None],
                [None, None, "This row still represents unresolved workbook data"],
            ]
        )
    )

    assert len(result.rows) == 2
    assert result.rows[1].source_row_number == 3
    assert error_codes(result.rows[1]) == {"REQUIRED"}


def test_formula_guest_cells_fail_but_independently_structural_footer_is_ignored() -> None:
    result = parse(
        simple_workbook(
            [
                ["Unmanaged Notes", "Full Name", "Invitation Type", "Allowed Companions"],
                [None, "Synthetic Formula Guest", '=UPPER("family")', 0],
                ["Closing constellation reckoning", None, None, "=COUNTA(B2:B2)"],
            ]
        )
    )

    assert len(result.rows) == 1
    assert "FORMULA_NOT_ALLOWED" in error_codes(result.rows[0])
    assert any(
        item["row_number"] == 3 and item["reason"] == "FOOTER" for item in result.ignored_rows
    )


def test_terminal_malformed_guest_is_not_silently_discarded() -> None:
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type", "Allowed Companions"],
                ["Guest Alpha", "Family", 1],
                ["Guest Beta", "Individual", 0],
                ["Guest Gamma", "Family", 1],
                [None, "Unseen Household", "=SUM(C2:C4)"],
            ]
        )
    )

    candidate = next(row for row in result.rows if row.source_row_number == 5)
    assert {"FORMULA_NOT_ALLOWED", "REQUIRED"} <= error_codes(candidate)
    assert not any(item["row_number"] == 5 for item in result.ignored_rows)


def test_merged_guest_like_row_is_not_silently_presentation() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Full Name", "Invitation Type", "Allowed Companions"])
    worksheet.append(["Guest Alpha", "Family", 1])
    worksheet.merge_cells("A3:C3")
    worksheet["A3"] = "Guest Beta"

    result = parse(save_workbook(workbook))

    candidate = next(row for row in result.rows if row.source_row_number == 3)
    assert candidate.values["full_name"] == "Guest Beta"
    assert "REQUIRED" in error_codes(candidate)
    assert not any(item["row_number"] == 3 for item in result.ignored_rows)


def test_merged_presentation_heading_remains_structural() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.merge_cells("A1:C1")
    worksheet["A1"] = "Synthetic Guest Register"
    worksheet.append([])
    worksheet.append(["Full Name", "Invitation Type", "Allowed Companions"])
    worksheet.append(["Guest Alpha", "Family", 1])

    result = parse(save_workbook(workbook))

    assert [row.source_row_number for row in result.rows] == [4]
    assert {"row_number": 1, "reason": "PRESENTATION"} in result.ignored_rows


def test_post_header_merged_unmapped_separator_remains_presentation() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Unmanaged Notes", "Full Name", "Invitation Type"])
    worksheet.append([None, "Guest Alpha", "Family"])
    worksheet.merge_cells("A3:C3")
    worksheet["A3"] = "Synthetic Section Divider"

    result = parse(save_workbook(workbook))

    assert [row.source_row_number for row in result.rows] == [2]
    assert result.ignored_rows[-1] == {"row_number": 3, "reason": "PRESENTATION"}


def test_merged_guest_like_row_between_valid_guests_remains_candidate() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Full Name", "Invitation Type", "Allowed Companions"])
    worksheet.append(["Guest Alpha", "Family", 1])
    worksheet.merge_cells("A3:C3")
    worksheet["A3"] = "Guest Needs Correction"
    worksheet.append(["Guest Omega", "Individual", 0])

    result = parse(save_workbook(workbook))

    assert [row.source_row_number for row in result.rows] == [2, 3, 4]
    assert "REQUIRED" in error_codes(result.rows[1])


def test_merged_guest_row_with_additional_mapped_evidence_remains_candidate() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Full Name", "Invitation Type", "Allowed Companions", "Category"])
    worksheet.append(["Guest Alpha", "Family", 1, "Friends"])
    worksheet.merge_cells("A3:C3")
    worksheet["A3"] = "Guest Needs Correction"
    worksheet["D3"] = "Colleagues"

    result = parse(save_workbook(workbook))

    assert [row.source_row_number for row in result.rows] == [2, 3]
    assert result.rows[1].values["category"] == "Colleagues"
    assert "REQUIRED" in error_codes(result.rows[1])


@pytest.mark.parametrize(
    "field_key",
    [
        GuestFieldKey.FULL_NAME,
        GuestFieldKey.PHONE_NUMBER,
        GuestFieldKey.EMAIL,
        GuestFieldKey.INVITATION_TYPE,
        GuestFieldKey.ALLOWED_COMPANIONS,
        GuestFieldKey.TABLE_NUMBER,
        GuestFieldKey.CATEGORY,
    ],
)
def test_excel_error_cell_cannot_be_valid_guest_value(field_key: GuestFieldKey) -> None:
    result = parse(workbook_with_excel_error(field_key))
    candidate = result.rows[0]

    assert [row.source_row_number for row in result.rows] == [2]
    assert candidate.values[field_key.value] is None
    assert candidate.normalization_errors == [
        {
            "field": field_key.value,
            "code": "EXCEL_ERROR_VALUE_NOT_ALLOWED",
            "message": "Excel error cells are not allowed in guest fields",
        }
    ]
    assert candidate.normalization_errors[0] in candidate.validation_errors
    assert "#N/A" not in candidate.normalization_errors[0]["message"]


@pytest.mark.parametrize(
    "error_value",
    [
        "#N/A",
        "#VALUE!",
        "#REF!",
        "#DIV/0!",
        "#NAME?",
        "#NUM!",
        "#NULL!",
    ],
)
def test_known_excel_error_types_have_one_stable_normalization_error(error_value: str) -> None:
    result = parse(workbook_with_excel_error(GuestFieldKey.CATEGORY, error_value))

    assert result.rows[0].normalization_errors == [
        {
            "field": "category",
            "code": "EXCEL_ERROR_VALUE_NOT_ALLOWED",
            "message": "Excel error cells are not allowed in guest fields",
        }
    ]


def test_literal_excel_error_text_in_a_normal_text_cell_is_not_retyped_as_an_error() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Full Name", "Invitation Type"])
    worksheet.append(["#N/A", "Family"])
    worksheet["A2"].data_type = "s"

    result = parse(save_workbook(workbook))

    assert result.rows[0].values["full_name"] == "#N/A"
    assert result.rows[0].normalization_errors == []
    assert result.rows[0].validation_errors == []


@pytest.mark.parametrize("field_key", [GuestFieldKey.PHONE_NUMBER, GuestFieldKey.EMAIL])
def test_contact_value_at_persistence_limit_is_accepted(field_key: GuestFieldKey) -> None:
    limit = persistence_text_limit(field_key)
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type", field_key.value],
                ["Guest Alpha", "Family", contact_value_at_length(field_key, limit)],
            ]
        )
    )

    assert len(result.rows[0].values[field_key.value]) == limit
    assert result.rows[0].validation_errors == []


@pytest.mark.parametrize("field_key", [GuestFieldKey.PHONE_NUMBER, GuestFieldKey.EMAIL])
@pytest.mark.parametrize("extra_characters", [1, 100])
def test_contact_value_must_fit_persistence_contract(
    field_key: GuestFieldKey,
    extra_characters: int,
) -> None:
    limit = persistence_text_limit(field_key)
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type", field_key.value],
                [
                    "Guest Alpha",
                    "Family",
                    contact_value_at_length(field_key, limit + extra_characters),
                ],
            ]
        )
    )

    expected_code = "INVALID_PHONE" if field_key is GuestFieldKey.PHONE_NUMBER else "INVALID_EMAIL"
    assert (field_key.value, expected_code) in {
        (item["field"], item["code"]) for item in result.rows[0].validation_errors
    }


@pytest.mark.parametrize(
    ("table_number", "expected"),
    [
        (1, "1"),
        (1.0, "1"),
        (12, "12"),
        (12.0, "12"),
        ("12", "12"),
        ("A1", "A1"),
        ("Table 7", "Table 7"),
    ],
)
def test_table_number_losslessly_normalizes_integral_numbers_and_preserves_text(
    table_number: object,
    expected: str,
) -> None:
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type", "Table Number"],
                ["Synthetic Table Guest", "Family", table_number],
            ]
        )
    )

    assert result.rows[0].values["table_number"] == expected
    assert result.rows[0].validation_errors == []


@pytest.mark.parametrize("table_number", [12.5, True])
def test_table_number_rejects_ambiguous_or_unsafe_excel_scalars(table_number: object) -> None:
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type", "Table Number"],
                ["Synthetic Unsafe Table", "Family", table_number],
            ]
        )
    )

    assert error_codes(result.rows[0]) & {
        "LOSSY_VALUE_NOT_ALLOWED",
        "UNSUPPORTED_VALUE_TYPE",
    }


@pytest.mark.parametrize("table_number", [nan, inf, -inf])
def test_table_number_rejects_non_finite_numeric_scalars(table_number: float) -> None:
    synthetic_cell = SimpleNamespace(value=table_number, data_type="n")

    value, field_error = _normalize_cell(GuestFieldKey.TABLE_NUMBER, synthetic_cell)  # type: ignore[arg-type]

    assert value is None
    assert field_error is not None
    assert field_error["code"] == "LOSSY_VALUE_NOT_ALLOWED"


@pytest.mark.parametrize(
    ("table_number", "expected_code"),
    [
        (date(2027, 1, 2), "UNSUPPORTED_VALUE_TYPE"),
        (time(14, 30), "UNSUPPORTED_VALUE_TYPE"),
        (complex(12, 1), "LOSSY_VALUE_NOT_ALLOWED"),
        (object(), "LOSSY_VALUE_NOT_ALLOWED"),
    ],
)
def test_table_number_rejects_date_time_and_unsupported_values(
    table_number: object,
    expected_code: str,
) -> None:
    synthetic_cell = SimpleNamespace(value=table_number, data_type="n")

    value, field_error = _normalize_cell(GuestFieldKey.TABLE_NUMBER, synthetic_cell)  # type: ignore[arg-type]

    assert value is None
    assert field_error is not None
    assert field_error["code"] == expected_code


def test_table_number_formula_in_genuine_guest_row_is_rejected() -> None:
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type", "Table Number"],
                ["Synthetic Formula Table", "Family", "=SUM(6,6)"],
            ]
        )
    )

    assert result.rows[0].values["table_number"] is None
    assert "FORMULA_NOT_ALLOWED" in error_codes(result.rows[0])


def test_numeric_phone_remains_rejected_when_excel_formatting_is_unrecoverable() -> None:
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type", "Phone Number", "Table Number"],
                ["Synthetic Numeric Phone", "Family", 255700000012, 12],
            ]
        )
    )

    assert result.rows[0].values["table_number"] == "12"
    assert result.rows[0].values["phone_number"] is None
    assert "LOSSY_VALUE_NOT_ALLOWED" in error_codes(result.rows[0])


@pytest.mark.parametrize(
    ("label", "formula"),
    [
        ("Attendance constellation reckoning", "=SUM(D2:D2)"),
        ("Closing roster constellation", "=COUNT(D2:D2)"),
        ("Gathering census ledger", "=COUNTA(B2:B2)"),
    ],
)
def test_terminal_sparse_aggregate_footer_is_independent_of_label_vocabulary(
    label: str,
    formula: str,
) -> None:
    result = parse(
        simple_workbook(
            [
                [
                    "Unmanaged Notes",
                    "Guest Name",
                    "Invitation Type",
                    "Allowed Companions",
                ],
                [None, "Synthetic Guest", "Family", 1],
                [label, None, None, formula],
            ]
        )
    )

    assert [row.source_row_number for row in result.rows] == [2]
    assert any(
        item["row_number"] == 3 and item["reason"] == "FOOTER" for item in result.ignored_rows
    )


def test_summary_like_text_without_aggregate_formula_remains_candidate() -> None:
    result = parse(
        simple_workbook(
            [
                ["Unmanaged Notes", "Guest Name", "Invitation Type"],
                [None, "Synthetic Guest", "Family"],
                ["Total Guests", None, None, 1],
            ]
        )
    )

    assert [row.source_row_number for row in result.rows] == [2, 3]
    assert "REQUIRED" in error_codes(result.rows[1])


def test_terminal_summary_text_in_mapped_invitation_type_remains_candidate() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.merge_cells("A1:E1")
    worksheet["A1"] = "Synthetic Wedding Guest Register"
    worksheet.merge_cells("A2:E2")
    worksheet["A2"] = "Synthetic Import Preview"
    worksheet.append([])
    worksheet.append(
        [
            "Allowed Companions",
            "Full Name",
            "Unmanaged Notes",
            "Invitation Type",
            "Category",
        ]
    )
    worksheet.append([1, "Synthetic Guest One", None, "Family", "Friends"])
    worksheet.append([0, "Synthetic Guest Two", None, "Individual", None])
    worksheet.append(["=SUM(A5:A6)", None, None, "Grand Total", None])

    result = parse(save_workbook(workbook))

    assert [row.source_row_number for row in result.rows] == [5, 6, 7]
    assert [item["reason"] for item in result.ignored_rows].count("PRESENTATION") == 2
    assert [item["reason"] for item in result.ignored_rows].count("BLANK") == 1
    assert [item["reason"] for item in result.ignored_rows].count("FOOTER") == 0
    assert "FORMULA_NOT_ALLOWED" in error_codes(result.rows[-1])


@pytest.mark.parametrize(
    "summary_text",
    ["Uncatalogued constellation reckoning", "Approximate guest totals note"],
)
def test_unknown_summary_text_in_mapped_guest_column_remains_candidate(
    summary_text: str,
) -> None:
    result = parse(
        simple_workbook(
            [
                ["Allowed Companions", "Full Name", "Invitation Type"],
                [1, "Synthetic Guest", "Family"],
                ["=SUM(A2:A2)", None, summary_text],
            ]
        )
    )

    assert [row.source_row_number for row in result.rows] == [2, 3]
    assert "FORMULA_NOT_ALLOWED" in error_codes(result.rows[-1])


def test_non_terminal_malformed_guest_with_formula_remains_candidate() -> None:
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type", "Allowed Companions"],
                ["Guest Alpha", "Family", 1],
                [None, "Unseen Household", "=SUM(C2:C2)"],
                ["Guest Omega", "Individual", 0],
            ]
        )
    )

    assert [row.source_row_number for row in result.rows] == [2, 3, 4]
    assert {"FORMULA_NOT_ALLOWED", "REQUIRED"} <= error_codes(result.rows[1])


@pytest.mark.parametrize("table_number", [20, "A-20"])
def test_malformed_guest_after_blank_precedes_genuine_footer(
    table_number: int | str,
) -> None:
    result = parse(
        simple_workbook(
            [
                [
                    "Unmanaged Notes",
                    "Full Name",
                    "Invitation Type",
                    "Allowed Companions",
                    "Table Number",
                ],
                [None, "Guest Alpha", "Family", 1, table_number],
                [None, None, None, None, None],
                [None, None, "Unseen Household", "=SUM(D2:D2)", None],
                ["Synthetic aggregate label", None, None, "=SUM(D2:D4)", None],
            ]
        )
    )

    assert [row.source_row_number for row in result.rows] == [2, 4]
    assert {"FORMULA_NOT_ALLOWED", "REQUIRED"} <= error_codes(result.rows[1])
    assert {"row_number": 3, "reason": "BLANK"} in result.ignored_rows
    assert result.ignored_rows[-1] == {"row_number": 5, "reason": "FOOTER"}


def test_sparse_rows_with_only_canonical_evidence_remain_candidates() -> None:
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type", "Table Number", "Category"],
                ["Guest Alpha", "Family", None, None],
                [None, "Unseen Household", None, None],
                [None, None, "T-8", None],
                [None, None, None, "Colleagues"],
            ]
        )
    )

    assert [row.source_row_number for row in result.rows] == [2, 3, 4, 5]
    assert all("REQUIRED" in error_codes(row) for row in result.rows[1:])


@pytest.mark.parametrize(
    "formula",
    [
        "=UPPER(A2)",
        "=SUM(A1:A1)",
        "=SUM(B2:B2)",
        "=SUM(A2:B2)",
        "=SUM(A2:A3)",
        "=SUM(AAAAAA2:AAAAAA2)",
    ],
)
def test_terminal_unrelated_or_nonpreceding_formula_remains_candidate(formula: str) -> None:
    result = parse(
        simple_workbook(
            [
                ["Allowed Companions", "Full Name", "Invitation Type"],
                [1, "Synthetic Guest", "Family"],
                [formula, None, "Closing constellation reckoning"],
            ]
        )
    )

    assert [row.source_row_number for row in result.rows] == [2, 3]
    assert "FORMULA_NOT_ALLOWED" in error_codes(result.rows[1])


@pytest.mark.parametrize(
    ("guest_field_header", "guest_evidence"),
    [
        ("Phone Number", "+255700000099"),
        ("Email", "evidence@example.test"),
        ("Table Number", "T-9"),
        ("Category", "Friends"),
    ],
)
def test_terminal_summary_label_with_guest_field_evidence_remains_candidate(
    guest_field_header: str,
    guest_evidence: str,
) -> None:
    result = parse(
        simple_workbook(
            [
                [
                    "Allowed Companions",
                    "Full Name",
                    "Invitation Type",
                    guest_field_header,
                ],
                [1, "Synthetic Guest", "Family", None],
                ["=SUM(A2:A2)", None, "Grand Total", guest_evidence],
            ]
        )
    )

    assert [row.source_row_number for row in result.rows] == [2, 3]
    assert "FORMULA_NOT_ALLOWED" in error_codes(result.rows[1])


@pytest.mark.parametrize(
    "footer_row",
    [
        ["=SUM(A2:A2)", None, "Family", None],
        ["=SUM(A2:A2)", "Synthetic Formula Guest", None, None],
        ["=SUM(A2:A2)", "Total Anderson", None, None],
        ["=SUM(A2:A2)", None, "Grand Total", "T-9"],
        ["=SUM(A2:A2)", None, "Grand Total", "Meaningful guest category"],
    ],
)
def test_terminal_formula_rows_with_meaningful_guest_evidence_remain_candidates(
    footer_row: list[object],
) -> None:
    result = parse(
        simple_workbook(
            [
                ["Allowed Companions", "Full Name", "Invitation Type", "Category"],
                [1, "Synthetic Guest", "Family", "Friends"],
                footer_row,
            ]
        )
    )

    assert [row.source_row_number for row in result.rows] == [2, 3]
    assert "FORMULA_NOT_ALLOWED" in error_codes(result.rows[1])


@pytest.mark.parametrize(
    ("guest_field_header", "guest_evidence", "expected_error"),
    [
        ("Phone Number", "+25512", "INVALID_PHONE"),
        ("Email", "incomplete@", "INVALID_EMAIL"),
    ],
)
def test_terminal_malformed_but_guest_like_evidence_remains_candidate(
    guest_field_header: str,
    guest_evidence: str,
    expected_error: str,
) -> None:
    result = parse(
        simple_workbook(
            [
                [
                    "Allowed Companions",
                    "Full Name",
                    "Invitation Type",
                    guest_field_header,
                ],
                [1, "Synthetic Guest", "Family", None],
                ["=SUM(A2:A2)", None, None, guest_evidence],
            ]
        )
    )

    assert [row.source_row_number for row in result.rows] == [2, 3]
    assert {"FORMULA_NOT_ALLOWED", expected_error} <= error_codes(result.rows[1])


def test_terminal_malformed_full_name_remains_candidate() -> None:
    result = parse(
        simple_workbook(
            [
                ["Allowed Companions", "Full Name", "Invitation Type"],
                [1, "Synthetic Guest", "Family"],
                ["=SUM(A2:A2)", "X" * 161, None],
            ]
        )
    )

    assert [row.source_row_number for row in result.rows] == [2, 3]
    assert {"FORMULA_NOT_ALLOWED", "INVALID_TEXT"} <= error_codes(result.rows[1])


def test_non_terminal_footer_signature_remains_candidate() -> None:
    result = parse(
        simple_workbook(
            [
                ["Allowed Companions", "Full Name", "Invitation Type"],
                [1, "Synthetic Guest One", "Family"],
                ["=SUM(A2:A2)", None, "Grand Total"],
                [0, "Synthetic Guest Two", "Individual"],
            ]
        )
    )

    assert [row.source_row_number for row in result.rows] == [2, 3, 4]
    assert "FORMULA_NOT_ALLOWED" in error_codes(result.rows[1])


def test_incomplete_genuine_guest_at_end_remains_candidate() -> None:
    result = parse(
        simple_workbook(
            [
                ["Allowed Companions", "Full Name", "Invitation Type"],
                [1, "Synthetic Guest", "Family"],
                [0, "Synthetic Needs Correction", None],
            ]
        )
    )

    assert [row.source_row_number for row in result.rows] == [2, 3]
    assert "REQUIRED" in error_codes(result.rows[1])


def test_table_and_companion_fields_keep_explicit_semantics() -> None:
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type", "Allowed Companions", "Table Number"],
                ["Synthetic Semantic Guest", "Family", 2, 20],
            ]
        )
    )

    assert result.rows[0].values["allowed_companions"] == 2
    assert result.rows[0].values["table_number"] == "20"
    assert result.rows[0].validation_errors == []


def test_sparse_or_formula_guest_evidence_is_never_hidden_as_footer() -> None:
    result = parse(
        simple_workbook(
            [
                [
                    "Unmanaged Notes",
                    "Guest Name",
                    "Invitation Type",
                    "Allowed Companions",
                    "Category",
                ],
                [None, "Formula Guest", "Family", "=SUM(1,1)", None],
                [None, "Missing Type Guest", None, 0, None],
                ["Aggregate Before Guest", None, None, "=SUM(D2:D3)", None],
                [None, None, None, 0, "Incomplete guest evidence"],
            ]
        )
    )

    assert [row.source_row_number for row in result.rows] == [2, 3, 4, 5]
    assert "FORMULA_NOT_ALLOWED" in error_codes(result.rows[0])
    assert "REQUIRED" in error_codes(result.rows[1])
    assert "FORMULA_NOT_ALLOWED" in error_codes(result.rows[2])
    assert "REQUIRED" in error_codes(result.rows[3])


@pytest.mark.parametrize("companions", [0, 1, 2])
def test_companion_values_within_system_and_event_ceiling_are_valid(companions: int) -> None:
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type", "Allowed Companions"],
                [f"Synthetic Guest {companions}", "Family", companions],
            ]
        )
    )
    assert result.rows[0].validation_errors == []


@pytest.mark.parametrize("companions", [-1, 3, 1.5])
def test_invalid_companion_values_remain_visible(companions: float) -> None:
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type", "Allowed Companions"],
                ["Synthetic Invalid Companion", "Family", companions],
            ]
        )
    )
    assert error_codes(result.rows[0]) & {
        "INVALID_COMPANION_COUNT",
        "COMPANION_LIMIT_EXCEEDED",
    }


def test_invalid_phone_email_invitation_and_lossy_numeric_text_are_actionable() -> None:
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type", "Phone", "Email"],
                ["Synthetic Phone", "Family", "12-3", None],
                ["Synthetic Email", "Family", "+255700000010", "not-an-email"],
                ["Synthetic Type", "X" * 65, "+255700000011", None],
                ["Synthetic Numeric Phone", "Family", 255700000012, None],
            ]
        )
    )

    assert "INVALID_PHONE" in error_codes(result.rows[0])
    assert "INVALID_EMAIL" in error_codes(result.rows[1])
    assert "INVALID_INVITATION_TYPE" in error_codes(result.rows[2])
    assert "LOSSY_VALUE_NOT_ALLOWED" in error_codes(result.rows[3])


def test_duplicate_inside_workbook_marks_later_valid_identity() -> None:
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type", "Email"],
                ["Synthetic First", "Family", "same@example.test"],
                ["Synthetic Second", "Individual", "SAME@example.test"],
            ]
        )
    )
    assert result.rows[0].validation_errors == []
    assert error_codes(result.rows[1]) == {"DUPLICATE_IN_WORKBOOK"}


def test_optional_missing_columns_default_without_becoming_structural_failure() -> None:
    result = parse(
        simple_workbook(
            [
                ["Name", "Invite Type"],
                ["Synthetic Minimal", "Individual"],
            ]
        )
    )
    assert result.rows[0].values["allowed_companions"] == 0
    assert result.rows[0].values["email"] is None
    assert result.rows[0].validation_errors == []


def test_missing_required_configured_column_fails_closed() -> None:
    with pytest.raises(WorkbookRejectedError) as exc_info:
        parse(simple_workbook([["Full Name", "Phone"], ["Synthetic", "+255700000001"]]))
    assert exc_info.value.code == "MISSING_REQUIRED_COLUMNS"
    assert "invitation_type" in str(exc_info.value)


@pytest.mark.parametrize("first_value", [None, "Synthetic event instructions"])
def test_empty_or_irrelevant_first_sheet_does_not_receive_preference(
    first_value: str | None,
) -> None:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Cover"
    if first_value is not None:
        first.append([first_value])
    second = workbook.create_sheet("Guests")
    second.append(["Full Name", "Invitation Type"])
    second.append(["Synthetic Selected", "Family"])

    result = parse(save_workbook(workbook))

    assert result.sheet_name == "Guests"
    assert [row.values["full_name"] for row in result.rows] == ["Synthetic Selected"]


def test_plausible_lower_mapping_score_beats_empty_higher_score_sheet() -> None:
    workbook = Workbook()
    empty = workbook.active
    assert empty is not None
    empty.title = "Canonical Cover"
    empty.append(
        [
            "Full Name",
            "Invitation Type",
            "Phone",
            "Email",
            "Allowed Companions",
            "Table Number",
            "Category",
        ]
    )
    empty.merge_cells("A2:G2")
    empty["A2"] = "Synthetic presentation content"
    guests = workbook.create_sheet("Real Guests")
    guests.append(["Invitation Type", "Full Name"])
    guests.append(["Family", "Synthetic Plausible"])

    result = parse(save_workbook(workbook))

    assert result.sheet_name == "Real Guests"
    assert result.header_row_number == 1
    assert result.column_mapping == {"invitation_type": 1, "full_name": 2}


def test_hidden_high_score_sheet_never_competes_with_visible_plausible_sheet() -> None:
    workbook = Workbook()
    visible = workbook.active
    assert visible is not None
    visible.title = "Visible Guests"
    visible.append(["Full Name", "Invitation Type"])
    visible.append(["Synthetic Visible", "Family"])
    hidden = workbook.create_sheet("Hidden Canonical")
    hidden.sheet_state = "hidden"
    hidden.append(
        [
            "Full Name",
            "Invitation Type",
            "Phone",
            "Email",
            "Allowed Companions",
            "Table Number",
            "Category",
        ]
    )
    hidden.append(["Synthetic Hidden", "Family", None, None, 0, "H-1", "Synthetic"])

    result = parse(save_workbook(workbook))

    assert result.sheet_name == "Visible Guests"
    assert [row.values["full_name"] for row in result.rows] == ["Synthetic Visible"]


def test_two_plausible_sheets_use_bounded_data_evidence_for_clear_winner() -> None:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Short List"
    first.append(["Full Name", "Invitation Type", "Email"])
    first.append(["Synthetic One", "Family", None])
    second = workbook.create_sheet("Guest Register")
    second.append(["Full Name", "Invitation Type"])
    second.append(["Synthetic Two", "Family"])
    second.append(["Synthetic Three", None])

    result = parse(save_workbook(workbook))

    assert result.sheet_name == "Guest Register"
    assert [row.values["full_name"] for row in result.rows] == [
        "Synthetic Two",
        "Synthetic Three",
    ]
    assert "REQUIRED" in error_codes(result.rows[1])


def test_equal_guest_tables_on_two_visible_sheets_are_ambiguous() -> None:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.append(["Full Name", "Invitation Type"])
    first.append(["Synthetic One", "Family"])
    second = workbook.create_sheet("Second")
    second.append(["Full Name", "Invitation Type"])
    second.append(["Synthetic Two", "Family"])

    with pytest.raises(WorkbookRejectedError) as exc_info:
        parse(save_workbook(workbook))
    assert exc_info.value.code == "AMBIGUOUS_GUEST_TABLE"


def test_hidden_equal_sheet_does_not_make_visible_selection_ambiguous() -> None:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.append(["Full Name", "Invitation Type"])
    first.append(["Synthetic One", "Family"])
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden.append(["Full Name", "Invitation Type"])
    hidden.append(["Synthetic Two", "Family"])

    result = parse(save_workbook(workbook))
    assert [row.values["full_name"] for row in result.rows] == ["Synthetic One"]


def test_later_real_header_beats_same_sheet_empty_header_candidate() -> None:
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type"],
                [],
                ["Full Name", "Invitation Type"],
                ["Synthetic Selected", "Family"],
            ]
        )
    )

    assert result.header_row_number == 3
    assert [row.source_row_number for row in result.rows] == [4]


def test_two_independent_plausible_tables_on_one_sheet_are_ambiguous() -> None:
    with pytest.raises(WorkbookRejectedError) as exc_info:
        parse(
            simple_workbook(
                [
                    ["Full Name", "Invitation Type"],
                    ["Synthetic First", "Family"],
                    [],
                    ["Full Name", "Invitation Type", "Email"],
                    ["Synthetic Second", "Family", "second@example.test"],
                ]
            )
        )

    assert exc_info.value.code == "AMBIGUOUS_GUEST_TABLE"
    assert "Synthetic" not in str(exc_info.value)


def test_repeated_header_inside_one_data_region_is_not_a_competing_table() -> None:
    result = parse(
        simple_workbook(
            [
                ["Full Name", "Invitation Type"],
                ["Synthetic First", "Family"],
                ["Full Name", "Invitation Type"],
                ["Synthetic Second", "Individual"],
            ]
        )
    )

    assert result.header_row_number == 1
    assert [row.values["full_name"] for row in result.rows] == [
        "Synthetic First",
        "Synthetic Second",
    ]
    assert {item["reason"] for item in result.ignored_rows} == {"REPEATED_HEADER"}


def test_invalid_rows_with_mapped_occupancy_contribute_to_table_viability() -> None:
    workbook = Workbook()
    empty = workbook.active
    assert empty is not None
    empty.append(["Full Name", "Invitation Type", "Phone", "Email"])
    malformed = workbook.create_sheet("Malformed Guests")
    malformed.append(["Full Name", "Invitation Type"])
    malformed.append(["Synthetic Invalid", None])

    result = parse(save_workbook(workbook))

    assert result.sheet_name == "Malformed Guests"
    assert error_codes(result.rows[0]) == {"REQUIRED"}


def test_duplicate_canonical_header_does_not_gain_candidate_viability() -> None:
    workbook = Workbook()
    malformed = workbook.active
    assert malformed is not None
    malformed.append(["Full Name", "Name", "Invitation Type", "Email"])
    malformed.append(["Synthetic Wrong", "Synthetic Duplicate", "Family", "x@example.test"])
    valid = workbook.create_sheet("Valid")
    valid.append(["Full Name", "Invitation Type"])
    valid.append(["Synthetic Selected", "Family"])

    result = parse(save_workbook(workbook))

    assert result.sheet_name == "Valid"


def test_populated_optional_columns_cannot_replace_required_header_coverage() -> None:
    workbook = Workbook()
    optional_only = workbook.active
    assert optional_only is not None
    optional_only.append(["Phone", "Email", "Allowed Companions", "Table Number", "Category"])
    for index in range(5):
        optional_only.append(
            [f"+2557000000{index}", f"guest{index}@example.test", 1, index, "Synthetic"]
        )
    valid = workbook.create_sheet("Required Fields")
    valid.append(["Full Name", "Invitation Type"])
    valid.append(["Synthetic Selected", "Family"])

    result = parse(save_workbook(workbook))

    assert result.sheet_name == "Required Fields"


def test_unknown_columns_do_not_increase_canonical_selection_score() -> None:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.append(["Full Name", "Invitation Type", "Unmanaged One", "Unmanaged Two"])
    first.append(["Synthetic First", "Family", "x", "y"])
    second = workbook.create_sheet("Canonical Winner")
    second.append(["Full Name", "Invitation Type", "Email"])
    second.append(["Synthetic Second", "Family", "second@example.test"])

    result = parse(save_workbook(workbook))

    assert result.sheet_name == "Canonical Winner"


@pytest.mark.parametrize(
    ("header_row", "expected_header", "expected_code"),
    [
        (MAX_HEADER_SCAN_ROWS, MAX_HEADER_SCAN_ROWS, None),
        (MAX_HEADER_SCAN_ROWS + 1, None, "MISSING_REQUIRED_COLUMNS"),
    ],
)
def test_header_scan_boundary_remains_explicit_and_bounded(
    header_row: int,
    expected_header: int | None,
    expected_code: str | None,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.cell(row=header_row, column=1, value="Full Name")
    worksheet.cell(row=header_row, column=2, value="Invitation Type")
    worksheet.cell(row=header_row + 1, column=1, value="Synthetic Boundary")
    worksheet.cell(row=header_row + 1, column=2, value="Family")
    content = save_workbook(workbook)

    if expected_code is not None:
        with pytest.raises(WorkbookRejectedError) as exc_info:
            parse(content)
        assert exc_info.value.code == expected_code
    else:
        result = parse(content)
        assert result.header_row_number == expected_header


def test_candidate_probe_is_bounded_on_sparse_sheet() -> None:
    workbook = Workbook()
    sparse = workbook.active
    assert sparse is not None
    sparse.append(["Full Name", "Invitation Type", "Email"])
    sparse.cell(row=MAX_CANDIDATE_PROBE_ROWS + 2, column=1, value="Synthetic Late")
    sparse.cell(row=MAX_CANDIDATE_PROBE_ROWS + 2, column=2, value="Family")
    plausible = workbook.create_sheet("Bounded Winner")
    plausible.append(["Full Name", "Invitation Type"])
    plausible.append(["Synthetic In Probe", "Family"])

    result = parse(save_workbook(workbook))

    assert result.sheet_name == "Bounded Winner"


@pytest.mark.parametrize(
    ("content", "expected_code"),
    [
        (b"not a workbook", "UNSUPPORTED_WORKBOOK"),
        (b"PK malformed", "MALFORMED_WORKBOOK"),
    ],
)
def test_unsupported_and_corrupt_workbooks_fail_safely(
    content: bytes,
    expected_code: str,
) -> None:
    with pytest.raises(WorkbookRejectedError) as exc_info:
        parse(content)
    assert exc_info.value.code == expected_code


def test_xlsx_path_traversal_entry_is_rejected_before_parser() -> None:
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("_rels/.rels", "<Relationships/>")
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("xl/worksheets/sheet1.xml", "<worksheet/>")
        archive.writestr("../escape.xml", "private")

    with pytest.raises(WorkbookRejectedError) as exc_info:
        parse(stream.getvalue())
    assert exc_info.value.code == "UNSAFE_XLSX_PACKAGE"


def test_corrupt_xlsx_xml_is_converted_to_bounded_rejection() -> None:
    valid = simple_workbook([["Full Name", "Invitation Type"], ["Synthetic", "Family"]])
    source = ZipFile(BytesIO(valid))
    stream = BytesIO()
    with source, ZipFile(stream, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                content = b"<worksheet><broken>"
            target.writestr(item, content)

    with pytest.raises(WorkbookRejectedError) as exc_info:
        parse(stream.getvalue())
    assert exc_info.value.code == "MALFORMED_WORKBOOK"


def test_candidate_limit_is_enforced_after_structure_detection() -> None:
    content = simple_workbook(
        [
            ["Full Name", "Invitation Type"],
            ["Synthetic One", "Family"],
            ["Synthetic Two", "Family"],
        ]
    )
    parser = WorkbookParser(maximum_candidate_rows=1)
    with pytest.raises(WorkbookRejectedError) as exc_info:
        parser.parse(content, default_policies(), 2)
    assert exc_info.value.code == "CANDIDATE_LIMIT_EXCEEDED"


def test_one_thousand_row_synthetic_parse_is_bounded() -> None:
    rows: list[list[object]] = [["Full Name", "Invitation Type", "Phone"]]
    rows.extend(
        [f"Synthetic Guest {index}", "Individual", f"+255700{index:06d}"] for index in range(1_000)
    )
    content = simple_workbook(rows)

    started = perf_counter()
    result = parse(content)
    duration = perf_counter() - started

    assert len(result.rows) == 1_000
    assert duration < 10
