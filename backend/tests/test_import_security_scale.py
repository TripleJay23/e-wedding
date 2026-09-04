"""Synthetic XLSX package-security and event-scale acceptance tests."""

from __future__ import annotations

import socket
import struct
import warnings
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile

import pytest
from openpyxl import Workbook

from app.core.settings import Settings
from app.imports.workbook import (
    MAX_COMPRESSION_RATIO,
    MAX_ENTRY_BYTES,
    MAX_HEADER_SCAN_ROWS,
    MAX_WORKSHEET_COLUMNS,
    MAX_WORKSHEET_ROWS,
    MAX_ZIP_ENTRIES,
    WorkbookParser,
    WorkbookRejectedError,
    _inspect_xlsx_package,
)
from app.models.event import DEFAULT_GUEST_FIELD_CONFIGURATION
from app.services.guest_validation import GuestFieldPolicy


def policies() -> list[GuestFieldPolicy]:
    return [
        GuestFieldPolicy(field_key, enabled, required)
        for field_key, enabled, required in DEFAULT_GUEST_FIELD_CONFIGURATION
    ]


def save_workbook(workbook: Workbook) -> bytes:
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def parse(content: bytes, *, maximum_candidate_rows: int = 5_000):
    return WorkbookParser(maximum_candidate_rows=maximum_candidate_rows).parse(
        content,
        policies(),
        2,
    )


def minimal_package(extra_entries: list[tuple[str, bytes]] | None = None) -> bytes:
    entries = [
        ("[Content_Types].xml", b"<Types/>"),
        ("_rels/.rels", b"<Relationships/>"),
        ("xl/workbook.xml", b"<workbook/>"),
        ("xl/worksheets/sheet1.xml", b"<worksheet/>"),
    ]
    entries.extend(extra_entries or [])
    stream = BytesIO()
    with ZipFile(stream, "w", compression=ZIP_STORED) as archive:
        for name, content in entries:
            archive.writestr(name, content)
    return stream.getvalue()


def patch_central_directory(
    content: bytes,
    member_name: str,
    *,
    flag_bits: int | None = None,
    compressed_size: int | None = None,
    file_size: int | None = None,
) -> bytes:
    """Change bounded ZIP metadata without allocating the represented payload."""
    data = bytearray(content)
    offset = 0
    while (offset := data.find(b"PK\x01\x02", offset)) >= 0:
        filename_length, extra_length, comment_length = struct.unpack_from(
            "<HHH", data, offset + 28
        )
        filename_start = offset + 46
        filename_end = filename_start + filename_length
        if data[filename_start:filename_end].decode("utf-8") == member_name:
            if flag_bits is not None:
                struct.pack_into("<H", data, offset + 8, flag_bits)
            if compressed_size is not None:
                struct.pack_into("<I", data, offset + 20, compressed_size)
            if file_size is not None:
                struct.pack_into("<I", data, offset + 24, file_size)
            return bytes(data)
        offset = filename_end + extra_length + comment_length
    raise AssertionError(f"central-directory member not found: {member_name}")


def synthetic_guest_workbook(candidate_rows: int) -> bytes:
    workbook = Workbook(write_only=True)
    worksheet = workbook.create_sheet("Synthetic Scale")
    worksheet.append(
        [
            "Full Name",
            "Invitation Type",
            "Allowed Companions",
            "Phone Number",
            "Email",
            "Table Number",
            "Category",
        ]
    )
    for index in range(candidate_rows):
        worksheet.append(
            [
                f"Synthetic Guest {index:05d}",
                "Family" if index % 2 == 0 else "Individual",
                index % 3,
                f"+2557{index:08d}",
                f"guest-{index:05d}@example.test",
                str(index % 500) if index % 2 == 0 else f"A-{index % 500}",
                f"Synthetic Category {index % 10}",
            ]
        )
    return save_workbook(workbook)


def test_authoritative_import_limits_match_the_bounded_contract() -> None:
    settings = Settings(
        APP_ENV="development",
        DATABASE_URL="postgresql+psycopg://synthetic:synthetic@127.0.0.1/synthetic",
        JWT_SECRET="synthetic-jwt-signing-secret-at-least-32-bytes",
        JWT_ISSUER="synthetic-issuer",
        JWT_AUDIENCE="synthetic-audience",
    )

    assert settings.MAX_IMPORT_WORKBOOK_BYTES == 10 * 1024 * 1024
    assert settings.MAX_IMPORT_CANDIDATE_ROWS == 5_000
    assert MAX_WORKSHEET_ROWS == 10_000
    assert MAX_WORKSHEET_COLUMNS == 100
    assert MAX_HEADER_SCAN_ROWS == 100


def test_zip_entry_count_guard_uses_metadata_only() -> None:
    extras = [(f"synthetic/{index}.xml", b"") for index in range(MAX_ZIP_ENTRIES - 3)]

    with pytest.raises(WorkbookRejectedError) as exc_info:
        _inspect_xlsx_package(minimal_package(extras))

    assert exc_info.value.code == "UNSAFE_XLSX_PACKAGE"
    assert "synthetic/" not in str(exc_info.value)


@pytest.mark.parametrize(
    ("file_size", "compressed_size", "extra_sizes"),
    [
        (MAX_ENTRY_BYTES + 1, 1_000_000, []),
        (MAX_COMPRESSION_RATIO + 1, 1, []),
        (
            MAX_ENTRY_BYTES,
            1_000_000,
            [
                ("synthetic/second.bin", MAX_ENTRY_BYTES),
                ("synthetic/third.bin", 10 * 1024 * 1024 + 1),
            ],
        ),
    ],
)
def test_zip_expansion_guards_use_small_synthetic_metadata(
    file_size: int,
    compressed_size: int,
    extra_sizes: list[tuple[str, int]],
) -> None:
    package = minimal_package(
        [("synthetic/payload.bin", b"x")] + [(name, b"x") for name, _ in extra_sizes]
    )
    package = patch_central_directory(
        package,
        "synthetic/payload.bin",
        compressed_size=compressed_size,
        file_size=file_size,
    )
    for name, size in extra_sizes:
        package = patch_central_directory(
            package,
            name,
            compressed_size=1_000_000,
            file_size=size,
        )

    with pytest.raises(WorkbookRejectedError) as exc_info:
        _inspect_xlsx_package(package)

    assert exc_info.value.code == "UNSAFE_XLSX_PACKAGE"
    assert "payload.bin" not in str(exc_info.value)


def test_duplicate_zip_member_is_rejected_without_disclosing_its_name() -> None:
    stream = BytesIO()
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        with ZipFile(stream, "w", compression=ZIP_STORED) as archive:
            for name, content in [
                ("[Content_Types].xml", b"<Types/>"),
                ("_rels/.rels", b"<Relationships/>"),
                ("xl/workbook.xml", b"<workbook/>"),
                ("xl/worksheets/sheet1.xml", b"<worksheet/>"),
                ("xl/worksheets/sheet1.xml", b"private synthetic marker"),
            ]:
                archive.writestr(name, content)

    with pytest.raises(WorkbookRejectedError) as exc_info:
        _inspect_xlsx_package(stream.getvalue())

    assert exc_info.value.code == "UNSAFE_XLSX_PACKAGE"
    assert "sheet1.xml" not in str(exc_info.value)
    assert "private synthetic marker" not in str(exc_info.value)


def test_encrypted_zip_member_flag_is_rejected() -> None:
    package = minimal_package([("synthetic/encrypted.bin", b"x")])
    package = patch_central_directory(
        package,
        "synthetic/encrypted.bin",
        flag_bits=0x1,
    )

    with pytest.raises(WorkbookRejectedError) as exc_info:
        _inspect_xlsx_package(package)

    assert exc_info.value.code == "UNSAFE_XLSX_PACKAGE"


def test_missing_required_openxml_part_is_a_bounded_malformed_package() -> None:
    package = minimal_package()
    source = ZipFile(BytesIO(package))
    stream = BytesIO()
    with source, ZipFile(stream, "w", compression=ZIP_STORED) as target:
        for item in source.infolist():
            if item.filename != "xl/workbook.xml":
                target.writestr(item, source.read(item.filename))

    with pytest.raises(WorkbookRejectedError) as exc_info:
        _inspect_xlsx_package(stream.getvalue())

    assert exc_info.value.code == "MALFORMED_WORKBOOK"
    assert "workbook.xml" not in str(exc_info.value)


def test_non_active_custom_openxml_part_is_safely_inert_to_package_inspection() -> None:
    _inspect_xlsx_package(minimal_package([("docProps/custom.xml", b"<Properties/>")]))


@pytest.mark.parametrize(
    "unsafe_name",
    [
        "../escape.xml",
        "../../escape.xml",
        "/absolute/escape.xml",
        "C:/absolute/escape.xml",
        "C:drive-relative.xml",
        r"..\escape.xml",
        r"C:\absolute\escape.xml",
    ],
)
def test_package_member_path_tricks_are_rejected(unsafe_name: str) -> None:
    with pytest.raises(WorkbookRejectedError) as exc_info:
        _inspect_xlsx_package(minimal_package([(unsafe_name, b"synthetic")]))

    assert exc_info.value.code == "UNSAFE_XLSX_PACKAGE"
    assert "escape" not in str(exc_info.value)


@pytest.mark.parametrize(
    "active_part",
    [
        "xl/vbaProject.bin",
        "xl/activeX/activeX1.bin",
        "xl/embeddings/oleObject1.bin",
        "xl/externalLinks/externalLink1.xml",
        "customUI/customUI.xml",
    ],
)
def test_active_or_external_package_parts_are_explicitly_rejected(active_part: str) -> None:
    with pytest.raises(WorkbookRejectedError) as exc_info:
        _inspect_xlsx_package(minimal_package([(active_part, b"synthetic")]))

    assert exc_info.value.code == "UNSAFE_XLSX_PACKAGE"
    assert active_part not in str(exc_info.value)


@pytest.mark.parametrize(
    "formula",
    [
        "=1+1",
        "=SUM(A1:A1)",
        "='[external.xlsx]Sheet1'!A1",
        '=HYPERLINK("https://synthetic.invalid/", "Synthetic")',
        '="Synthetic text output"',
    ],
)
def test_representative_guest_formulas_are_inert_and_invalid(formula: str) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Full Name", "Invitation Type", "Category"])
    worksheet.append(["Synthetic Guest", "Family", formula])

    result = parse(save_workbook(workbook))

    assert len(result.rows) == 1
    assert result.rows[0].values["category"] is None
    assert {item["code"] for item in result.rows[0].validation_errors} == {"FORMULA_NOT_ALLOWED"}


def test_external_hyperlink_relationship_is_retained_as_inert_metadata(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Full Name", "Invitation Type"])
    worksheet.append(["Synthetic Guest", "Family"])
    worksheet["A2"].hyperlink = "https://synthetic.invalid/private"
    content = save_workbook(workbook)
    network_calls: list[object] = []

    def reject_network(*args: object, **kwargs: object) -> None:
        network_calls.append((args, kwargs))
        raise AssertionError("workbook parsing attempted network access")

    monkeypatch.setattr(socket.socket, "connect", reject_network)
    result = parse(content)

    assert len(result.rows) == 1
    assert result.rows[0].values["full_name"] == "Synthetic Guest"
    assert network_calls == []


@pytest.mark.parametrize("candidate_rows", [100, 1_000, 4_999, 5_000])
def test_synthetic_event_scale_is_deterministic(candidate_rows: int) -> None:
    result = parse(synthetic_guest_workbook(candidate_rows))

    assert len(result.rows) == candidate_rows
    assert all(not row.validation_errors for row in result.rows)
    assert len(
        {fingerprint for row in result.rows for fingerprint in row.identity_fingerprints}
    ) == (candidate_rows * 2)


def test_candidate_limit_rejects_5_001_without_partial_result() -> None:
    with pytest.raises(WorkbookRejectedError) as exc_info:
        parse(synthetic_guest_workbook(5_001))

    assert exc_info.value.code == "CANDIDATE_LIMIT_EXCEEDED"


def test_worksheet_row_boundary_accepts_maximum_and_rejects_next_row() -> None:
    accepted = Workbook()
    accepted_sheet = accepted.active
    assert accepted_sheet is not None
    accepted_sheet.append(["Full Name", "Invitation Type"])
    accepted_sheet.append(["Synthetic First", "Family"])
    accepted_sheet.cell(MAX_WORKSHEET_ROWS, 1, "Synthetic Last")
    accepted_sheet.cell(MAX_WORKSHEET_ROWS, 2, "Individual")
    accepted_result = parse(save_workbook(accepted))
    assert [row.source_row_number for row in accepted_result.rows] == [2, MAX_WORKSHEET_ROWS]

    rejected = Workbook()
    rejected_sheet = rejected.active
    assert rejected_sheet is not None
    rejected_sheet.append(["Full Name", "Invitation Type"])
    rejected_sheet.append(["Synthetic First", "Family"])
    rejected_sheet.cell(MAX_WORKSHEET_ROWS + 1, 1, "Synthetic Beyond")
    with pytest.raises(WorkbookRejectedError) as exc_info:
        parse(save_workbook(rejected))
    assert exc_info.value.code == "WORKSHEET_LIMIT_EXCEEDED"


def test_worksheet_column_boundary_accepts_maximum_and_rejects_next_column() -> None:
    accepted = Workbook()
    accepted_sheet = accepted.active
    assert accepted_sheet is not None
    accepted_sheet.append(["Full Name", "Invitation Type"])
    accepted_sheet.append(["Synthetic Guest", "Family"])
    accepted_sheet.cell(2, MAX_WORKSHEET_COLUMNS, "Synthetic bounded extra")
    assert len(parse(save_workbook(accepted)).rows) == 1

    rejected = Workbook()
    rejected_sheet = rejected.active
    assert rejected_sheet is not None
    rejected_sheet.append(["Full Name", "Invitation Type"])
    rejected_sheet.append(["Synthetic Guest", "Family"])
    rejected_sheet.cell(2, MAX_WORKSHEET_COLUMNS + 1, "Synthetic beyond")
    with pytest.raises(WorkbookRejectedError) as exc_info:
        parse(save_workbook(rejected))
    assert exc_info.value.code == "WORKSHEET_LIMIT_EXCEEDED"


def test_empty_sparse_dimension_claim_uses_actual_meaningful_cells() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Full Name", "Invitation Type"])
    worksheet.append(["Synthetic Guest", "Family"])
    content = save_workbook(workbook)
    source = ZipFile(BytesIO(content))
    stream = BytesIO()
    with source, ZipFile(stream, "w", compression=ZIP_DEFLATED) as target:
        for item in source.infolist():
            member = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                member = member.replace(b'ref="A1:B2"', b'ref="A1:CV10001"')
            target.writestr(item, member)

    result = parse(stream.getvalue())

    assert len(result.rows) == 1
    assert result.rows[0].source_row_number == 2
