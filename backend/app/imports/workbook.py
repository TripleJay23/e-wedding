"""Bounded XLSX package inspection, structure detection, and row normalization."""

from __future__ import annotations

import math
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, time
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.utils.cell import column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet

from app.models.event import GuestFieldKey
from app.services.guest_validation import (
    GuestFieldError,
    GuestFieldPolicy,
    error,
    guest_identity_fingerprints,
    validate_guest_values,
)

MAX_ZIP_ENTRIES = 2_000
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ENTRY_BYTES = 20 * 1024 * 1024
MAX_COMPRESSION_RATIO = 250
MAX_WORKSHEET_ROWS = 10_000
MAX_WORKSHEET_COLUMNS = 100
MAX_HEADER_SCAN_ROWS = 100
MAX_CANDIDATE_PROBE_ROWS = 100

_REQUIRED_XLSX_PARTS = {
    "[Content_Types].xml",
    "_rels/.rels",
    "xl/workbook.xml",
}
_FORBIDDEN_PREFIXES = (
    "xl/externalLinks/",
    "xl/embeddings/",
    "xl/activeX/",
    "customUI/",
)
_FORBIDDEN_PARTS = {"xl/vbaProject.bin"}
_WINDOWS_DRIVE_MEMBER = re.compile(r"^[A-Za-z]:")
_HEADER_PUNCTUATION = re.compile(r"[^a-z0-9]+")
_PHONE_LIKE = re.compile(r"^\+?[0-9 ()\-.]+$")
_COMPANION_INTEGER_TEXT = re.compile(r"^(?P<integer>[0-9]+)(?:\.0+)?$")
_AGGREGATE_FORMULA = re.compile(
    r"^\s*=\s*(SUM|COUNT|COUNTA)\s*\(\s*\$?([A-Z]+)\$?(\d+)\s*:\s*"
    r"\$?([A-Z]+)\$?(\d+)\s*\)\s*$",
    re.IGNORECASE,
)
_COUNTABLE_FIELDS = frozenset(GuestFieldKey)

HEADER_ALIASES: dict[GuestFieldKey, frozenset[str]] = {
    GuestFieldKey.FULL_NAME: frozenset({"full name", "guest", "guest name", "name"}),
    GuestFieldKey.PHONE_NUMBER: frozenset(
        {
            "contact number",
            "mobile",
            "mobile number",
            "phone",
            "phone no",
            "phone number",
            "telephone",
        }
    ),
    GuestFieldKey.EMAIL: frozenset({"email", "email address", "e mail"}),
    GuestFieldKey.INVITATION_TYPE: frozenset(
        {"guest type", "invitation", "invitation type", "invite type"}
    ),
    GuestFieldKey.ALLOWED_COMPANIONS: frozenset(
        {
            "allowed companions",
            "companions",
            "number of companions",
            "plus one",
            "plus ones",
        }
    ),
    GuestFieldKey.TABLE_NUMBER: frozenset({"table", "table no", "table number", "table number no"}),
    GuestFieldKey.CATEGORY: frozenset({"category", "group", "guest category"}),
}


class WorkbookRejectedError(ValueError):
    """A bounded client-safe XLSX rejection."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True, slots=True)
class ParsedGuestRow:
    """One actual guest row normalized without retaining unrelated cell data."""

    source_row_number: int
    values: dict[str, str | int | None]
    normalization_errors: list[GuestFieldError]
    validation_errors: list[GuestFieldError]
    identity_fingerprints: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class ParsedWorkbook:
    """Privacy-minimal deterministic workbook structure and candidate rows."""

    sheet_name: str
    header_row_number: int
    column_mapping: dict[str, int]
    ignored_rows: tuple[dict[str, int | str], ...]
    rows: tuple[ParsedGuestRow, ...]


@dataclass(frozen=True, slots=True)
class _HeaderCandidate:
    worksheet: Worksheet
    row_number: int
    mapping: dict[GuestFieldKey, int]

    @property
    def score(self) -> int:
        return len(self.mapping)


@dataclass(frozen=True, slots=True)
class _CandidateAssessment:
    """Privacy-free bounded evidence for one header and its following region."""

    candidate: _HeaderCandidate
    probe_start_row: int
    probe_end_row: int
    candidate_row_count: int
    plausible_candidate_count: int
    structural_row_count: int
    mapped_occupancy_count: int

    @property
    def has_meaningful_guest_data(self) -> bool:
        return self.plausible_candidate_count > 0

    @property
    def selection_score(self) -> tuple[int, int, int, int]:
        return (
            int(self.has_meaningful_guest_data),
            self.plausible_candidate_count,
            self.mapped_occupancy_count,
            self.candidate.score,
        )


@dataclass(slots=True)
class _GuestTableContext:
    """Bounded semantic context accumulated from preceding candidate rows."""

    header_row_number: int
    mapped_candidate_row_numbers: list[int]
    invitation_types: set[str]
    has_strong_guest_evidence: bool = False

    def start_reordered_section(self, header_row_number: int) -> None:
        """Reset field-position evidence when a repeated header changes layout."""
        self.header_row_number = header_row_number
        self.mapped_candidate_row_numbers.clear()
        self.has_strong_guest_evidence = False

    def observe(
        self,
        row_number: int,
        values: dict[str, str | int | None],
        policies: list[GuestFieldPolicy],
        event_companion_maximum: int,
        *,
        has_mapped_occupancy: bool,
    ) -> None:
        if has_mapped_occupancy:
            self.mapped_candidate_row_numbers.append(row_number)
        evidence = _guest_evidence(
            values,
            policies,
            event_companion_maximum,
            known_invitation_types=self.invitation_types,
        )
        if evidence:
            self.has_strong_guest_evidence = True

        invitation_type = values.get(GuestFieldKey.INVITATION_TYPE.value)
        if evidence and _field_value_is_valid(
            GuestFieldKey.INVITATION_TYPE,
            invitation_type,
            event_companion_maximum,
        ):
            assert isinstance(invitation_type, str)
            self.invitation_types.add(_semantic_text(invitation_type))


class WorkbookParser:
    """Parse one bounded XLSX upload under an immutable event-field policy."""

    def __init__(self, *, maximum_candidate_rows: int) -> None:
        self._maximum_candidate_rows = maximum_candidate_rows

    def parse(
        self,
        workbook_bytes: bytes,
        policies: list[GuestFieldPolicy],
        event_companion_maximum: int,
    ) -> ParsedWorkbook:
        """Inspect, understand, normalize, and validate a workbook in memory."""
        _inspect_xlsx_package(workbook_bytes)
        try:
            workbook = load_workbook(
                BytesIO(workbook_bytes),
                read_only=False,
                data_only=False,
                keep_links=False,
            )
        except Exception as exc:
            raise WorkbookRejectedError(
                "MALFORMED_WORKBOOK", "workbook could not be read safely"
            ) from exc

        try:
            visible_sheets = [
                sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"
            ]
            if not visible_sheets:
                raise WorkbookRejectedError(
                    "NO_VISIBLE_WORKSHEET", "workbook has no visible worksheet"
                )
            for sheet in visible_sheets:
                if sheet.max_row > MAX_WORKSHEET_ROWS or sheet.max_column > MAX_WORKSHEET_COLUMNS:
                    raise WorkbookRejectedError(
                        "WORKSHEET_LIMIT_EXCEEDED", "worksheet dimensions exceed import limits"
                    )

            candidate = self._detect_header(
                visible_sheets,
                policies,
                event_companion_maximum,
            )
            return self._parse_sheet(candidate, policies, event_companion_maximum)
        finally:
            workbook.close()

    def _detect_header(
        self,
        worksheets: list[Worksheet],
        policies: list[GuestFieldPolicy],
        event_companion_maximum: int,
    ) -> _HeaderCandidate:
        required = {policy.field_key for policy in policies if policy.enabled and policy.required}
        enabled = {policy.field_key for policy in policies if policy.enabled}
        aliases = _enabled_alias_lookup(enabled)
        per_sheet: list[_CandidateAssessment] = []
        best_partial: set[GuestFieldKey] = set()

        for worksheet in worksheets:
            sheet_candidates: list[_HeaderCandidate] = []
            scan_end_row = min(worksheet.max_row, MAX_HEADER_SCAN_ROWS)
            maximum_column = worksheet.max_column
            for row_number, cells in enumerate(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=scan_end_row,
                    max_col=maximum_column,
                ),
                start=1,
            ):
                mapping, duplicated = _map_header_row(cells, aliases)
                if len(mapping) > len(best_partial):
                    best_partial = set(mapping)
                if duplicated or not required.issubset(mapping) or len(mapping) < 2:
                    continue
                sheet_candidates.append(_HeaderCandidate(worksheet, row_number, mapping))
            if sheet_candidates:
                assessments = [
                    _assess_header_candidate(
                        item,
                        policies,
                        event_companion_maximum,
                        stop_before_row=(
                            sheet_candidates[index + 1].row_number
                            if index + 1 < len(sheet_candidates)
                            else None
                        ),
                    )
                    for index, item in enumerate(sheet_candidates)
                ]
                assessments = _collapse_repeated_header_assessments(assessments, required)
                meaningful = [item for item in assessments if item.has_meaningful_guest_data]
                if len(meaningful) > 1:
                    raise WorkbookRejectedError(
                        "AMBIGUOUS_GUEST_TABLE",
                        "multiple plausible guest tables were detected on one visible worksheet",
                    )
                contenders = meaningful or assessments
                best_score = max(item.selection_score for item in contenders)
                finalists = [item for item in contenders if item.selection_score == best_score]
                if len(finalists) != 1:
                    raise WorkbookRejectedError(
                        "AMBIGUOUS_GUEST_TABLE",
                        "multiple equal guest table candidates were detected "
                        "on one visible worksheet",
                    )
                per_sheet.append(finalists[0])

        if not per_sheet:
            missing = sorted(field.value for field in required - best_partial)
            if missing:
                raise WorkbookRejectedError(
                    "MISSING_REQUIRED_COLUMNS",
                    f"required configured columns are missing: {', '.join(missing)}",
                )
            raise WorkbookRejectedError(
                "NO_GUEST_TABLE", "no unambiguous configured guest table was detected"
            )

        best_score = max(item.selection_score for item in per_sheet)
        finalists = [item for item in per_sheet if item.selection_score == best_score]
        if len(finalists) != 1:
            raise WorkbookRejectedError(
                "AMBIGUOUS_GUEST_TABLE", "multiple visible worksheets contain equal guest tables"
            )
        return finalists[0].candidate

    def _parse_sheet(
        self,
        candidate: _HeaderCandidate,
        policies: list[GuestFieldPolicy],
        event_companion_maximum: int,
    ) -> ParsedWorkbook:
        worksheet = candidate.worksheet
        maximum_column = worksheet.max_column
        ignored: list[dict[str, int | str]] = []
        if candidate.row_number > 1:
            for row_number, cells in enumerate(
                worksheet.iter_rows(
                    min_row=1,
                    max_row=candidate.row_number - 1,
                    max_col=maximum_column,
                ),
                start=1,
            ):
                reason = "BLANK" if _row_is_blank(cells) else "PRESENTATION"
                ignored.append({"row_number": row_number, "reason": reason})

        parsed_rows: list[ParsedGuestRow] = []
        seen_valid_fingerprints: set[str] = set()
        aliases = _enabled_alias_lookup(set(candidate.mapping))
        required = {policy.field_key for policy in policies if policy.enabled and policy.required}
        selected_fields = set(candidate.mapping)
        active_mapping = candidate.mapping
        merged_ranges = tuple(worksheet.merged_cells.ranges)
        context = _GuestTableContext(candidate.row_number, [], set())

        for row_number, cells in enumerate(
            worksheet.iter_rows(
                min_row=candidate.row_number + 1,
                max_row=worksheet.max_row,
                max_col=maximum_column,
            ),
            start=candidate.row_number + 1,
        ):
            repeated_mapping = _safe_repeated_header_mapping(
                cells,
                aliases=aliases,
                selected_fields=selected_fields,
                required_fields=required,
            )
            if repeated_mapping is not None:
                ignored.append({"row_number": row_number, "reason": "REPEATED_HEADER"})
                if repeated_mapping != active_mapping:
                    context.start_reordered_section(row_number)
                active_mapping = repeated_mapping
                continue

            reason = _ignored_row_reason(
                cells,
                row_number=row_number,
                selected_mapping=active_mapping,
                merged_ranges=merged_ranges,
                context=context,
            )
            if reason is not None:
                ignored.append({"row_number": row_number, "reason": reason})
                continue

            values, normalization_errors = _normalize_candidate(cells, active_mapping, policies)
            validation_errors = validate_guest_values(
                values,
                policies,
                event_companion_maximum,
                prior_errors=normalization_errors,
            )
            fingerprints = guest_identity_fingerprints(values)
            if not validation_errors and set(fingerprints) & seen_valid_fingerprints:
                validation_errors.append(
                    error(None, "DUPLICATE_IN_WORKBOOK", "row duplicates an earlier valid row")
                )
            if not validation_errors:
                seen_valid_fingerprints.update(fingerprints)

            parsed_rows.append(
                ParsedGuestRow(
                    source_row_number=row_number,
                    values=values,
                    normalization_errors=normalization_errors,
                    validation_errors=validation_errors,
                    identity_fingerprints=fingerprints,
                )
            )
            context.observe(
                row_number,
                values,
                policies,
                event_companion_maximum,
                has_mapped_occupancy=_mapped_occupancy(cells, active_mapping) > 0,
            )
            if len(parsed_rows) > self._maximum_candidate_rows:
                raise WorkbookRejectedError(
                    "CANDIDATE_LIMIT_EXCEEDED", "workbook contains too many guest rows"
                )

        if not parsed_rows:
            raise WorkbookRejectedError("NO_GUEST_ROWS", "guest table contains no candidate rows")

        return ParsedWorkbook(
            sheet_name=worksheet.title,
            header_row_number=candidate.row_number,
            column_mapping={field.value: column for field, column in candidate.mapping.items()},
            ignored_rows=tuple(ignored),
            rows=tuple(parsed_rows),
        )


def _assess_header_candidate(
    candidate: _HeaderCandidate,
    policies: list[GuestFieldPolicy],
    event_companion_maximum: int,
    *,
    stop_before_row: int | None,
) -> _CandidateAssessment:
    """Probe one bounded region without using row validity as selection evidence."""
    worksheet = candidate.worksheet
    probe_start = candidate.row_number + 1
    probe_end = min(worksheet.max_row, candidate.row_number + MAX_CANDIDATE_PROBE_ROWS)
    if stop_before_row is not None:
        probe_end = min(probe_end, stop_before_row - 1)

    candidate_rows = 0
    plausible_rows = 0
    structural_rows = 0
    mapped_occupancy = 0
    aliases = _enabled_alias_lookup(set(candidate.mapping))
    required = {policy.field_key for policy in policies if policy.enabled and policy.required}
    selected_fields = set(candidate.mapping)
    active_mapping = candidate.mapping
    merged_ranges = tuple(worksheet.merged_cells.ranges)
    context = _GuestTableContext(candidate.row_number, [], set())

    maximum_column = worksheet.max_column
    for row_number, cells in enumerate(
        worksheet.iter_rows(
            min_row=probe_start,
            max_row=probe_end,
            max_col=maximum_column,
        ),
        start=probe_start,
    ):
        repeated_mapping = _safe_repeated_header_mapping(
            cells,
            aliases=aliases,
            selected_fields=selected_fields,
            required_fields=required,
        )
        if repeated_mapping is not None:
            structural_rows += 1
            if repeated_mapping != active_mapping:
                context.start_reordered_section(row_number)
            active_mapping = repeated_mapping
            continue

        reason = _ignored_row_reason(
            cells,
            row_number=row_number,
            selected_mapping=active_mapping,
            merged_ranges=merged_ranges,
            context=context,
        )
        if reason is not None:
            structural_rows += 1
            continue

        candidate_rows += 1
        occupied = _mapped_occupancy(cells, active_mapping)
        if occupied:
            plausible_rows += 1
            mapped_occupancy += occupied

        values, _ = _normalize_candidate(cells, active_mapping, policies)
        context.observe(
            row_number,
            values,
            policies,
            event_companion_maximum,
            has_mapped_occupancy=occupied > 0,
        )

    return _CandidateAssessment(
        candidate=candidate,
        probe_start_row=probe_start,
        probe_end_row=max(candidate.row_number, probe_end),
        candidate_row_count=candidate_rows,
        plausible_candidate_count=plausible_rows,
        structural_row_count=structural_rows,
        mapped_occupancy_count=mapped_occupancy,
    )


def _collapse_repeated_header_assessments(
    assessments: list[_CandidateAssessment],
    required_fields: set[GuestFieldKey],
) -> list[_CandidateAssessment]:
    """Fold a deterministic later header into an established guest-data region."""
    collapsed: list[_CandidateAssessment] = []
    for assessment in assessments:
        if (
            collapsed
            and collapsed[-1].has_meaningful_guest_data
            and required_fields.issubset(assessment.candidate.mapping)
            and set(assessment.candidate.mapping).issubset(collapsed[-1].candidate.mapping)
        ):
            preceding = collapsed[-1]
            collapsed[-1] = _CandidateAssessment(
                candidate=preceding.candidate,
                probe_start_row=preceding.probe_start_row,
                probe_end_row=assessment.probe_end_row,
                candidate_row_count=(
                    preceding.candidate_row_count + assessment.candidate_row_count
                ),
                plausible_candidate_count=(
                    preceding.plausible_candidate_count + assessment.plausible_candidate_count
                ),
                structural_row_count=(
                    preceding.structural_row_count + assessment.structural_row_count + 1
                ),
                mapped_occupancy_count=(
                    preceding.mapped_occupancy_count + assessment.mapped_occupancy_count
                ),
            )
            continue
        collapsed.append(assessment)
    return collapsed


def _inspect_xlsx_package(workbook_bytes: bytes) -> None:
    if not workbook_bytes.startswith(b"PK"):
        raise WorkbookRejectedError("UNSUPPORTED_WORKBOOK", "upload is not an XLSX package")
    try:
        with ZipFile(BytesIO(workbook_bytes)) as archive:
            entries = archive.infolist()
            if not entries or len(entries) > MAX_ZIP_ENTRIES:
                raise WorkbookRejectedError(
                    "UNSAFE_XLSX_PACKAGE", "XLSX package entry count is unsafe"
                )

            names: set[str] = set()
            total_uncompressed = 0
            for entry in entries:
                normalized_name = entry.filename.replace("\\", "/")
                path = PurePosixPath(normalized_name)
                if (
                    normalized_name in names
                    or path.is_absolute()
                    or _WINDOWS_DRIVE_MEMBER.match(normalized_name) is not None
                    or ".." in path.parts
                    or any(part == "" for part in path.parts)
                    or entry.flag_bits & 0x1
                ):
                    raise WorkbookRejectedError(
                        "UNSAFE_XLSX_PACKAGE", "XLSX package contains an unsafe entry"
                    )
                names.add(normalized_name)
                total_uncompressed += entry.file_size
                ratio = entry.file_size / max(entry.compress_size, 1)
                if (
                    entry.file_size > MAX_ENTRY_BYTES
                    or total_uncompressed > MAX_UNCOMPRESSED_BYTES
                    or ratio > MAX_COMPRESSION_RATIO
                ):
                    raise WorkbookRejectedError(
                        "UNSAFE_XLSX_PACKAGE", "XLSX package expansion exceeds safety limits"
                    )

            if not _REQUIRED_XLSX_PARTS.issubset(names) or not any(
                name.startswith("xl/worksheets/sheet") and name.endswith(".xml") for name in names
            ):
                raise WorkbookRejectedError(
                    "MALFORMED_WORKBOOK", "XLSX package is missing required workbook parts"
                )
            if names & _FORBIDDEN_PARTS or any(
                name.startswith(prefix) for name in names for prefix in _FORBIDDEN_PREFIXES
            ):
                raise WorkbookRejectedError(
                    "UNSAFE_XLSX_PACKAGE", "XLSX package contains unsupported active content"
                )
    except BadZipFile as exc:
        raise WorkbookRejectedError(
            "MALFORMED_WORKBOOK", "workbook could not be read safely"
        ) from exc


def _enabled_alias_lookup(enabled: set[GuestFieldKey]) -> dict[str, GuestFieldKey]:
    lookup: dict[str, GuestFieldKey] = {}
    for field_key in enabled:
        aliases = set(HEADER_ALIASES[field_key]) | {
            normalize_header(field_key.value),
            normalize_header(field_key.value.replace("_", " ")),
        }
        for alias in aliases:
            existing = lookup.get(alias)
            if existing is not None and existing is not field_key:
                raise RuntimeError("configured header aliases are ambiguous")
            lookup[alias] = field_key
    return lookup


def normalize_header(value: object) -> str:
    if not isinstance(value, str):
        return ""
    normalized = unicodedata.normalize("NFKC", value).casefold().replace("&", " and ")
    return " ".join(_HEADER_PUNCTUATION.sub(" ", normalized).split())


def _map_header_row(
    cells: tuple[Cell, ...],
    aliases: dict[str, GuestFieldKey],
) -> tuple[dict[GuestFieldKey, int], bool]:
    mapping: dict[GuestFieldKey, int] = {}
    duplicated = False
    for cell in cells:
        field_key = aliases.get(normalize_header(cell.value))
        if field_key is None:
            continue
        if field_key in mapping:
            duplicated = True
        else:
            mapping[field_key] = cell.column
    return mapping, duplicated


def _ignored_row_reason(
    cells: tuple[Cell, ...],
    *,
    row_number: int,
    selected_mapping: dict[GuestFieldKey, int],
    merged_ranges: tuple[Any, ...],
    context: _GuestTableContext,
) -> str | None:
    populated = [cell for cell in cells if cell.value is not None and cell.value != ""]
    if not populated:
        return "BLANK"

    is_single_merged_value = len(populated) == 1 and any(
        merged.min_row <= row_number <= merged.max_row
        and merged.min_col <= populated[0].column <= merged.max_col
        and merged.max_col > merged.min_col
        for merged in merged_ranges
    )
    if is_single_merged_value and populated[0].column not in selected_mapping.values():
        # A post-header value in a canonical column is plausible guest content even
        # when merge geometry leaves the rest of the row empty.
        return "PRESENTATION"

    if _is_contextual_aggregate_footer(
        cells,
        populated,
        row_number=row_number,
        selected_mapping=selected_mapping,
        context=context,
    ):
        return "FOOTER"
    return None


def _row_is_blank(cells: tuple[Cell, ...]) -> bool:
    return all(cell.value is None or cell.value == "" for cell in cells)


def _is_contextual_aggregate_footer(
    cells: tuple[Cell, ...],
    populated: list[Cell],
    *,
    row_number: int,
    selected_mapping: dict[GuestFieldKey, int],
    context: _GuestTableContext,
) -> bool:
    """Recognize a formula footer from independent structural and semantic evidence."""
    if (
        len(populated) != 2
        or not context.has_strong_guest_evidence
        or not context.mapped_candidate_row_numbers
    ):
        return False

    aggregate_column = selected_mapping.get(GuestFieldKey.ALLOWED_COMPANIONS)
    if aggregate_column is None:
        return False
    aggregate_cells = [cell for cell in populated if cell.column == aggregate_column]
    if len(aggregate_cells) != 1 or not _is_supported_preceding_aggregate(
        aggregate_cells[0],
        footer_row_number=row_number,
        header_row_number=context.header_row_number,
        candidate_row_numbers=context.mapped_candidate_row_numbers,
        selected_mapping=selected_mapping,
    ):
        return False

    summary_cells = [cell for cell in populated if cell is not aggregate_cells[0]]
    if (
        len(summary_cells) != 1
        or not isinstance(summary_cells[0].value, str)
        or _cell_is_formula(summary_cells[0])
    ):
        return False
    # Text in a canonical guest column can be an incomplete or invalid guest
    # field. Only an independently structural, unmapped label proves a footer.
    return summary_cells[0].column not in selected_mapping.values()


def _safe_repeated_header_mapping(
    cells: tuple[Cell, ...],
    *,
    aliases: dict[str, GuestFieldKey],
    selected_fields: set[GuestFieldKey],
    required_fields: set[GuestFieldKey],
) -> dict[GuestFieldKey, int] | None:
    """Return a deterministic repeated header confined to the selected field set."""
    mapping, duplicated = _map_header_row(cells, aliases)
    if (
        duplicated
        or len(mapping) < 2
        or not required_fields.issubset(mapping)
        or not set(mapping).issubset(selected_fields)
    ):
        return None
    return mapping


def _mapped_occupancy(
    cells: tuple[Cell, ...],
    mapping: dict[GuestFieldKey, int],
) -> int:
    return sum(
        1
        for column in mapping.values()
        if cells[column - 1].value is not None and cells[column - 1].value != ""
    )


def _is_supported_preceding_aggregate(
    cell: Cell,
    *,
    footer_row_number: int,
    header_row_number: int,
    candidate_row_numbers: list[int],
    selected_mapping: dict[GuestFieldKey, int],
) -> bool:
    """Accept one non-executed aggregate over the preceding guest-table range."""
    if not _cell_is_formula(cell) or not isinstance(cell.value, str):
        return False
    match = _AGGREGATE_FORMULA.fullmatch(cell.value)
    if match is None:
        return False

    function, first_column, first_row, last_column, last_row = match.groups()
    if first_column.casefold() != last_column.casefold():
        return False
    try:
        range_column = column_index_from_string(first_column)
    except ValueError:
        return False
    reference_fields = {
        field_key for field_key, column in selected_mapping.items() if column == range_column
    }
    if len(reference_fields) != 1:
        return False
    reference_field = next(iter(reference_fields))
    if function.casefold() == "sum" and reference_field is not GuestFieldKey.ALLOWED_COMPANIONS:
        return False
    if function.casefold() in {"count", "counta"} and reference_field not in _COUNTABLE_FIELDS:
        return False

    range_start = int(first_row)
    range_end = int(last_row)
    return (
        header_row_number < range_start <= min(candidate_row_numbers)
        and max(candidate_row_numbers) <= range_end < footer_row_number
    )


def _guest_evidence(
    values: dict[str, str | int | None],
    policies: list[GuestFieldPolicy],
    event_companion_maximum: int,
    *,
    known_invitation_types: set[str],
) -> bool:
    """Return whether normalized values contain meaningful guest evidence."""
    enabled = {policy.field_key for policy in policies if policy.enabled}
    full_name = values.get(GuestFieldKey.FULL_NAME.value)
    if GuestFieldKey.FULL_NAME in enabled and isinstance(full_name, str) and full_name:
        return True

    for field_key in (GuestFieldKey.PHONE_NUMBER, GuestFieldKey.EMAIL):
        value = values.get(field_key.value)
        if field_key not in enabled or not isinstance(value, str):
            continue
        if _field_value_is_valid(field_key, value, event_companion_maximum):
            return True
        if field_key is GuestFieldKey.PHONE_NUMBER and _PHONE_LIKE.fullmatch(value):
            return True
        if field_key is GuestFieldKey.EMAIL and "@" in value:
            return True

    invitation_type = values.get(GuestFieldKey.INVITATION_TYPE.value)
    invitation_evidence = (
        GuestFieldKey.INVITATION_TYPE in enabled
        and _field_value_is_valid(
            GuestFieldKey.INVITATION_TYPE,
            invitation_type,
            event_companion_maximum,
        )
        and isinstance(invitation_type, str)
        and _semantic_text(invitation_type) in known_invitation_types
    )
    if invitation_evidence:
        return True

    ancillary_fields = (GuestFieldKey.TABLE_NUMBER, GuestFieldKey.CATEGORY)
    ancillary_evidence = any(
        field_key in enabled
        and _field_value_is_valid(
            field_key,
            values.get(field_key.value),
            event_companion_maximum,
        )
        for field_key in ancillary_fields
    )
    return ancillary_evidence


def _field_value_is_valid(
    field_key: GuestFieldKey,
    value: str | int | None,
    event_companion_maximum: int,
) -> bool:
    if value is None or value == "":
        return False
    field_policy = GuestFieldPolicy(field_key, True, False)
    candidate: dict[str, str | int | None] = {
        GuestFieldKey.ALLOWED_COMPANIONS.value: 0,
        field_key.value: value,
    }
    return not validate_guest_values(candidate, [field_policy], event_companion_maximum)


def _semantic_text(value: str) -> str:
    return " ".join(value.casefold().split())


def _cell_is_formula(cell: Cell) -> bool:
    return cell.data_type == "f" or (isinstance(cell.value, str) and cell.value.startswith("="))


def _normalize_candidate(
    cells: tuple[Cell, ...],
    mapping: dict[GuestFieldKey, int],
    policies: list[GuestFieldPolicy],
) -> tuple[dict[str, str | int | None], list[GuestFieldError]]:
    values: dict[str, str | int | None] = {}
    errors: list[GuestFieldError] = []
    cells_by_column = {cell.column: cell for cell in cells}

    for policy in policies:
        field_key = policy.field_key
        if not policy.enabled:
            values[field_key.value] = 0 if field_key is GuestFieldKey.ALLOWED_COMPANIONS else None
            continue
        column = mapping.get(field_key)
        if column is None:
            values[field_key.value] = 0 if field_key is GuestFieldKey.ALLOWED_COMPANIONS else None
            continue
        cell = cells_by_column[column]
        value, field_error = _normalize_cell(field_key, cell)
        values[field_key.value] = value
        if field_error is not None:
            errors.append(field_error)

    return values, errors


def _normalize_cell(
    field_key: GuestFieldKey,
    cell: Cell,
) -> tuple[str | int | None, GuestFieldError | None]:
    if cell.data_type == "e":
        return None, error(
            field_key,
            "EXCEL_ERROR_VALUE_NOT_ALLOWED",
            "Excel error cells are not allowed in guest fields",
        )
    if _cell_is_formula(cell):
        return None, error(
            field_key, "FORMULA_NOT_ALLOWED", "formulas are not allowed in guest rows"
        )

    value = cell.value
    if value is None or value == "":
        default = 0 if field_key is GuestFieldKey.ALLOWED_COMPANIONS else None
        return default, None
    if isinstance(value, (bool, date, datetime, time)):
        return None, error(field_key, "UNSUPPORTED_VALUE_TYPE", "cell value type is unsupported")

    if field_key is GuestFieldKey.ALLOWED_COMPANIONS:
        if isinstance(value, int):
            return value, None
        if isinstance(value, float) and math.isfinite(value) and value.is_integer():
            return int(value), None
        if isinstance(value, str):
            normalized = value.strip()
            if not normalized:
                return 0, None
            match = _COMPANION_INTEGER_TEXT.fullmatch(normalized)
            if match is not None:
                try:
                    return int(match.group("integer")), None
                except ValueError:
                    pass
        return None, error(
            field_key,
            "INVALID_COMPANION_COUNT",
            "allowed companions must be an integer",
        )

    if field_key is GuestFieldKey.TABLE_NUMBER and isinstance(value, (int, float)):
        if isinstance(value, int):
            return str(value), None
        if math.isfinite(value) and value.is_integer():
            return str(int(value)), None
        return None, error(
            field_key,
            "LOSSY_VALUE_NOT_ALLOWED",
            "table number must be losslessly representable as text",
        )

    if not isinstance(value, str):
        return None, error(
            field_key,
            "LOSSY_VALUE_NOT_ALLOWED",
            "text fields must use lossless text cells",
        )
    normalized = value.strip()
    if not normalized:
        return None, None
    if field_key is GuestFieldKey.EMAIL:
        normalized = normalized.casefold()
    return normalized, None
