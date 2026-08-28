"""Deterministic CSV and XLSX encoders for one canonical report snapshot."""

import csv
import io
import unicodedata
from collections.abc import Callable
from datetime import UTC, datetime
from typing import Any

from openpyxl import Workbook

from app.services.reports import AttendanceRosterRow, ReportSnapshot

SUMMARY_FIELDS: tuple[tuple[str, str, Callable[[ReportSnapshot], Any]], ...] = (
    ("event", "Event ID", lambda report: str(report.event.id)),
    ("event", "Event Name", lambda report: report.event.event_name),
    ("event", "Bride Name", lambda report: report.event.bride_name),
    ("event", "Groom Name", lambda report: report.event.groom_name),
    ("event", "Event Date/Time", lambda report: report.event.event_datetime),
    ("event", "Venue", lambda report: report.event.venue),
    ("event", "Event Status", lambda report: report.event.status.value),
    (
        "totals",
        "Named Holders Invited",
        lambda report: report.totals.named_holders_invited,
    ),
    (
        "totals",
        "Named Holders Effectively Attended",
        lambda report: report.totals.named_holders_effectively_attended,
    ),
    (
        "totals",
        "Named Holders Effectively Not Attended",
        lambda report: report.totals.named_holders_effectively_not_attended,
    ),
    ("totals", "Companion Capacity", lambda report: report.totals.companion_capacity),
    (
        "totals",
        "Effective Admitted Companions",
        lambda report: report.totals.effective_admitted_companions,
    ),
    (
        "totals",
        "Maximum Total People Capacity",
        lambda report: report.totals.maximum_total_people_capacity,
    ),
    (
        "totals",
        "Effective Total People Attended",
        lambda report: report.totals.effective_total_people_attended,
    ),
    (
        "totals",
        "Remaining People Capacity",
        lambda report: report.totals.remaining_people_capacity,
    ),
    (
        "totals",
        "Named Holder Attendance Percentage",
        lambda report: report.totals.named_holder_attendance_percentage,
    ),
    (
        "raw_gate_metrics",
        "Real Gate Admissions",
        lambda report: report.raw_gate_metrics.real_gate_admissions,
    ),
    (
        "raw_gate_metrics",
        "Duplicate Scan Attempts",
        lambda report: report.raw_gate_metrics.duplicate_scan_attempts,
    ),
    (
        "raw_gate_metrics",
        "Unknown Credential Attempts",
        lambda report: report.raw_gate_metrics.unknown_credential_attempts,
    ),
    (
        "raw_gate_metrics",
        "Wrong Event Attempts",
        lambda report: report.raw_gate_metrics.wrong_event_attempts,
    ),
    (
        "raw_gate_metrics",
        "Inactive Credential Attempts",
        lambda report: report.raw_gate_metrics.inactive_credential_attempts,
    ),
    (
        "raw_gate_metrics",
        "Superseded Credential Attempts",
        lambda report: report.raw_gate_metrics.superseded_credential_attempts,
    ),
    (
        "raw_gate_metrics",
        "Ineligible Guest Attempts",
        lambda report: report.raw_gate_metrics.ineligible_guest_attempts,
    ),
    (
        "raw_gate_metrics",
        "Companion Limit Exceeded Attempts",
        lambda report: report.raw_gate_metrics.companion_limit_exceeded_attempts,
    ),
    (
        "raw_gate_metrics",
        "Actual Companions Admitted at Gate",
        lambda report: report.raw_gate_metrics.actual_companions_admitted_at_gate,
    ),
)

ROSTER_FIELDS: tuple[tuple[str, Callable[[AttendanceRosterRow], Any]], ...] = (
    ("Full Name", lambda row: row.full_name),
    ("Invitation Type", lambda row: row.invitation_type),
    ("Category", lambda row: row.category),
    ("Table Number", lambda row: row.table_number),
    ("Allowed Companions", lambda row: row.allowed_companions),
    ("Effective Attendance Status", lambda row: row.effective_status.value),
    ("Effective Attendance Source", lambda row: row.effective_source.value),
    ("Effective Admitted Companions", lambda row: row.effective_admitted_companions),
    ("Effective Party Size", lambda row: row.effective_party_size),
    ("Effective Timestamp", lambda row: row.effective_at),
    ("Raw Gate Attended", lambda row: row.raw_gate_attended),
    ("Raw Gate Timestamp", lambda row: row.raw_gate_checked_in_at),
    ("Correction Revision", lambda row: row.correction_revision),
    ("Has Correction", lambda row: row.has_correction),
)


def sanitize_spreadsheet_text(value: str) -> str:
    """Neutralize formulas and control characters without changing stored data."""

    cleaned = "".join(
        " " if unicodedata.category(character) == "Cc" else character for character in value
    )
    if cleaned.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + cleaned
    return cleaned


def encode_csv(report: ReportSnapshot) -> bytes:
    """Create a standards-compliant UTF-8 CSV with summary and roster sections."""

    output = io.StringIO(newline="")
    writer = csv.writer(output, lineterminator="\r\n")
    writer.writerow(("Section", "Metric", "Value"))
    for section, label, getter in SUMMARY_FIELDS:
        writer.writerow((section, label, _csv_value(getter(report))))
    writer.writerow(())
    writer.writerow(("Attendance Roster",))
    writer.writerow(tuple(label for label, _ in ROSTER_FIELDS))
    for row in report.attendance_roster:
        writer.writerow(tuple(_csv_value(getter(row)) for _, getter in ROSTER_FIELDS))
    return output.getvalue().encode("utf-8")


def encode_xlsx(report: ReportSnapshot) -> bytes:
    """Create a real write-only XLSX with native numbers, booleans, and dates."""

    workbook = Workbook(write_only=True)
    summary = workbook.create_sheet("Summary")
    summary.append(("Section", "Metric", "Value"))
    for section, label, getter in SUMMARY_FIELDS:
        summary.append((section, label, _xlsx_value(getter(report))))

    roster = workbook.create_sheet("Attendance Roster")
    roster.append(tuple(label for label, _ in ROSTER_FIELDS))
    for row in report.attendance_roster:
        roster.append(tuple(_xlsx_value(getter(row)) for _, getter in ROSTER_FIELDS))

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _csv_value(value: Any) -> str | int | float:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.astimezone(UTC).isoformat()
    if isinstance(value, str):
        return sanitize_spreadsheet_text(value)
    return value


def _xlsx_value(value: Any) -> str | int | float | bool | datetime:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.astimezone(UTC).replace(tzinfo=None)
    if isinstance(value, str):
        return sanitize_spreadsheet_text(value)
    return value
